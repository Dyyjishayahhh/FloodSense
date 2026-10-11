from collections import Counter
from pathlib import Path
import openpyxl

ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
SOURCE = ROOT / "1. FLOODSENSE-MAIN" / "QUESTIONNAIRES_SURVEY" / "FLOODSENSE_INITIAL_SURVEY_RESULTS_AND_SUMMARY" / "FloodSense Resident Survey Questionnaire (Responses).xlsx"

wb = openpyxl.load_workbook(SOURCE, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(2, ws.max_row + 1)]

def norm(v):
    if v is None:
        return ""
    return " ".join(str(v).strip().lower().split())

for excluded, label in [({0,5}, "excluding timestamp and name"), ({0,5,31}, "excluding timestamp, name, comment")]:
    signatures = [tuple(norm(v) for i, v in enumerate(row) if i not in excluded) for row in rows]
    groups = [count for count in Counter(signatures).values() if count > 1]
    print(label, "duplicate group sizes:", groups)

timestamps = [row[0] for row in rows]
for i in range(1, len(timestamps)):
    delta = (timestamps[i] - timestamps[i-1]).total_seconds()
    if delta < 10:
        print(f"rows {i+1}-{i+2}: {delta:.3f} seconds apart")

comments = [norm(row[31]) for row in rows if norm(row[31])]
print("comment repeats:")
for comment, count in Counter(comments).most_common():
    if count > 1:
        print(count, comment[:180])
