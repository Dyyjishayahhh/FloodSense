from pathlib import Path
from copy import deepcopy
import re
import openpyxl
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense')
SOURCE = ROOT / '1. FLOODSENSE-MAIN/INTERVIEWS/FLOODSENSE_INITIAL_SURVEY_RESULTS_AND_SUMMARY/FloodSense Resident Survey Questionnaire (Responses).xlsx'
OUT = ROOT / '1. FLOODSENSE-MAIN/QUESTIONNAIRES_SURVEY/Initial_Survey_(2026-10-14)_ThesisA1stsem.docx'
sheet = openpyxl.load_workbook(SOURCE, data_only=True).worksheets[0]

def clean(s):
    return str(s).replace('\ufffd', "'").replace('\u2019', "'").replace('\u2013', '-').replace('\u2014', '-')

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.6)
sec.bottom_margin = Inches(.6)
sec.left_margin = Inches(.65)
sec.right_margin = Inches(.65)
sec.header_distance = Inches(.25)
sec.footer_distance = Inches(.25)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.03
doc.settings.odd_and_even_pages_header_footer = False
sec.different_first_page_header_footer = False
for name, size in [('Title', 18), ('Heading 1', 14), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
    if name == 'Title':
        style.paragraph_format.space_before = Pt(0)
for style in doc.styles:
    if style.type == 1:
        pr = style.element.find(qn('w:pPr'))
        if pr is not None:
            for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)

header = sec.header.paragraphs[0]
header.text = ''
for other_header in [sec.first_page_header, sec.even_page_header]:
    for p in other_header.paragraphs: p.clear()
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('FloodSense Initial Survey  |  Respondent code: __________  |  Page ').font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def para(text, size=None, italic=False, bold=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    return p

def section(en, fil='', page=False):
    p = doc.add_paragraph(en, 'Heading 1')
    if page: p.paragraph_format.page_break_before = True
    if fil: para(fil, size=10.5, italic=True, keep=True)

def options(items):
    if items == ['Yes / Oo', 'No / Hindi']:
        para('\u2610  Yes / Oo        \u2610  No / Hindi', 10.5)
        return
    if items and items[0].startswith('18-24 years old'):
        for i in range(0, len(items), 2):
            p = para('\u2610  ' + items[i] + '\t\u2610  ' + items[i+1], 10.5, keep=i < 4)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.0))
        return
    if items[:5] == ['Very easy / Napakadali', 'Easy / Madali', 'Neither easy nor difficult / Katamtaman', 'Difficult / Mahirap', 'Very difficult / Napakahirap']:
        for a, b in [(items[0], items[1]), (items[2], None), (items[3], items[4])]:
            p = para('\u2610  ' + a + ('\t\u2610  ' + b if b else ''), 10.5, keep=True)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(3.65))
        items = items[5:]
    for i, item in enumerate(items):
        p = para('\u2610  ' + item, size=10.5)
        p.paragraph_format.left_indent = Inches(.15)
        p.paragraph_format.first_line_indent = Inches(-.15)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = i < len(items) - 1

def question(number, col, items=None, instruction='Select one. / Pumili ng isa.', optional=False):
    text = clean(sheet.cell(1, col).value)
    if col == 2:
        text = text.split('\n\n', 1)[1]
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    p = para(f'{number}. {lines[0]}', size=11, bold=True, keep=True)
    p.paragraph_format.space_before = Pt(9)
    for line in lines[1:]: para(line, size=10.5, italic=True, keep=True)
    if instruction:
        if 'Select all that apply.' in lines[0]:
            instruction = instruction.replace('Select all that apply. / Piliin ang lahat ng naaangkop. ', '').replace('Select all that apply. / Piliin ang lahat ng naaangkop.', '')
        if 'Select up to three.' in lines[0]: instruction = ''
        if instruction: para(instruction, size=9.5, keep=True)
    if items is not None: options(items)

YES_NO = ['Yes / Oo', 'No / Hindi']
FREQ = ['Never / Hindi kailanman', 'Rarely / Bihira', 'Sometimes / Paminsan-minsan', 'Often / Madalas', 'Always / Palagi', 'Not sure / Hindi tiyak']
EASE = ['Very easy / Napakadali', 'Easy / Madali', 'Neither easy nor difficult / Katamtaman', 'Difficult / Mahirap', 'Very difficult / Napakahirap']
MULTI = 'Select all that apply. / Piliin ang lahat ng naaangkop.'
EXCLUSIVE = 'If you select a no-problem, none, or not-sure option, select it alone. / Kung wala o hindi tiyak ang sagot, iyon lamang ang piliin.'

doc.add_paragraph('FloodSense Initial Resident Survey Questionnaire', 'Title')
para('Face to face survey / Harapang survey', 11, bold=True)
para('Survey date / Petsa: __________________    Interviewer code: __________', 10)
para('We are studying the flood-information and preparedness needs of Bacoor City residents. Please answer based on your own experience. There are no right or wrong answers.', 10.5)
para('Pinag-aaralan namin ang pangangailangan ng mga residente ng Bacoor City sa impormasyon tungkol sa baha at paghahanda. Sagutin ayon sa sariling karanasan. Walang tama o maling sagot.', 10.5, italic=True)
section('Data Privacy Notice and Consent', 'Pahayag sa Pagkapribado at Pahintulot')
para('Under Republic Act No. 10173, the Data Privacy Act of 2012, the FloodSense research team will use your survey answers, age group, barangay, and optional name or initials for this academic study. Participation is voluntary; you may omit optional answers or stop without penalty. Records will be kept confidential, accessible only to the research team and academic advisers, and reported as combined results. Records will be kept only for the study and required academic recordkeeping, then securely disposed of. For privacy questions or requests concerning access, correction, erasure, or withdrawal of consent, contact the research team through the survey administrator, subject to applicable law.', 10)
para('Gagamitin ng FloodSense research team ang inyong mga sagot, age group, barangay, at opsyonal na pangalan o initials para sa akademikong pag-aaral na ito, alinsunod sa Republic Act No. 10173. Kusang-loob ang pagsali; maaaring laktawan ang opsyonal na tanong o huminto nang walang parusa. Kumpidensiyal ang mga rekord, para lamang sa research team at academic advisers, at pinagsama-samang resulta ang ilalahad. Itatago ang rekord para sa pag-aaral at kinakailangang academic recordkeeping, saka ligtas na sisirain. Para sa privacy concerns o kahilingan tungkol sa datos o pagbawi ng pahintulot, makipag-ugnayan sa survey administrator, alinsunod sa batas.', 10, italic=True)
para('Research team contact / Contact ng research team: ______________________________', 9.5)
question(1, 2, ['Yes, I have read and understood the notice, and I voluntarily agree to participate. / Oo, nabasa at naunawaan ko ang notice, at kusang-loob akong pumapayag na lumahok.', 'No, I do not agree to participate. / Hindi ako pumapayag na lumahok.'])
para('If No, end the survey. / Kung Hindi, tapusin ang survey.', 10, bold=True)

section('Eligibility and Respondent Information', 'Pagiging Kwalipikado at Impormasyon ng Respondent', page=True)
para('End the survey if any eligibility answer is No or cannot be confirmed. / Tapusin kung Hindi ang sagot o hindi makumpirma ang pagiging kwalipikado.', 10)
question(2, 3, YES_NO)
question(3, 4, YES_NO)
question(4, 5, YES_NO)
section('Demographic Information', 'Impormasyong Demograpiko')
question(5, 6, instruction='Optional. Leave blank if you prefer. / Opsyonal. Maaaring iwang blangko.')
para('________________________________________________________________________')
question(6, 7, instruction='Write your barangay only; do not write your exact home address. / Barangay lamang; huwag isulat ang eksaktong address.')
para('Barangay: ______________________________________________________________')
question(7, 8, ['18-24 years old / taong gulang', '25-34 years old / taong gulang', '35-44 years old / taong gulang', '45-54 years old / taong gulang', '55-64 years old / taong gulang', '65 years old or older / 65 taong gulang o mas matanda'])

section('Flood and Heavy Rainfall Experience', 'Karanasan sa Baha at Malakas na Ulan', page=True)
question(8, 9, [
    'Floodwater entered our house / Pumasok ang baha sa aming bahay',
    'Floodwater affected our street or community / Binaha ang aming kalye o komunidad',
    'My usual travel route became difficult or impassable / Naging mahirap o hindi madaanan ang karaniwang ruta',
    'My work or studies were disrupted / Naantala ang trabaho o pag-aaral',
    'We evacuated or temporarily relocated / Lumikas o pansamantalang lumipat',
    'Property, utilities, or essential services were affected / Naapektuhan ang ari-arian, utilities, o essential services',
], MULTI)
question(9, 10, ['1 event / Isang beses', '2-3 events / Dalawa hanggang tatlong beses', '4-5 events / Apat hanggang limang beses', 'More than 5 events / Mahigit limang beses', 'Not sure / Hindi tiyak'])

section('Current Flood Information Sources', 'Kasalukuyang Pinagmumulan ng Impormasyon Tungkol sa Baha', page=True)
question(10, 11, [
    'Barangay announcements or public-address system / Anunsyo o public-address system ng barangay',
    'Bacoor City or BDRRMO official social-media page',
    'PAGASA or another government agency / PAGASA o ibang ahensya ng pamahalaan',
    'Television or radio news / Balita sa telebisyon o radyo',
    'Online news outlet / Online news', 'SMS or text message / SMS o text message',
    'Facebook groups or other social-media posts',
    'Family, friends, or neighbors / Pamilya, kaibigan, o kapitbahay',
    "Community organization or homeowners' association / Organisasyon ng komunidad o homeowners' association",
    'Government website or online bulletin / Website o online bulletin ng pamahalaan',
    'I usually receive no flood information / Karaniwan akong walang natatanggap na flood information',
], MULTI + ' ' + EXCLUSIVE)
question(11, 12, EASE + ['I usually cannot obtain the information / Karaniwan ay hindi ko nakukuha ang impormasyon', 'Not sure / Hindi tiyak'])

section('Problems with Existing Information', 'Mga Problema sa Kasalukuyang Impormasyon', page=True)
matrix_header = clean(sheet.cell(1, 13).value).split(' [')[0]
lines = [s for s in matrix_header.splitlines() if s.strip()]
para('12. ' + lines[0], 11, bold=True, keep=True)
for s in lines[1:]: para(s, 10.5, italic=True, keep=True)
para('Mark ONE box in EACH row. / Lagyan ng tsek ang ISANG kahon sa BAWAT hanay.', 10.5, bold=True, keep=True)
para('1 Never / Hindi kailanman     2 Rarely / Bihira     3 Sometimes / Paminsan-minsan\n4 Often / Madalas     5 Always / Palagi     NS Not sure / Hindi tiyak', 10.5, keep=True)
t = doc.add_table(rows=1, cols=7)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [3.65, .56, .56, .56, .56, .56, .65]
for i, text in enumerate(['Problem / Problema', '1', '2', '3', '4', '5', 'NS']):
    t.rows[0].cells[i].text = text
for idx, c in enumerate(range(13, 21)):
    row = t.add_row()
    label = clean(sheet.cell(1, c).value).rsplit(' [', 1)[1].rstrip(']')
    row.cells[0].text = f'{chr(97+idx)}. ' + label.replace(' / ', '\n')
    for cell in row.cells[1:]: cell.text = '\u2610'
for row in t.rows:
    pr = row._tr.get_or_add_trPr()
    pr.append(OxmlElement('w:cantSplit'))
    for i, cell in enumerate(row.cells):
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.06
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.size = Pt(10.5 if i == 0 else 13)
        cp = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
        cp.append(borders)
for cell in t.rows[0].cells:
    for p in cell.paragraphs:
        for r in p.runs: r.bold = True
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'F2F2F2'); cell._tc.get_or_add_tcPr().append(sh)
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))

section('Effects on Decisions and Local Information', 'Epekto sa Desisyon at Lokal na Impormasyon', page=True)
question(13, 21, FREQ)
question(14, 22, [
    'Preparing emergency supplies / Paghahanda ng emergency supplies',
    'Moving important belongings to a safer place / Paglilipat ng mahahalagang gamit',
    'Deciding whether it is safe to travel / Pagpapasya kung ligtas bumiyahe',
    'Choosing a road or area to avoid / Pagpili ng daan o lugar na iiwasan',
    'Deciding whether or when to evacuate / Pagpapasya kung lilikas o kailan lilikas',
    'Identifying an evacuation center / Pagtukoy ng evacuation center',
    'Deciding whether to attend work or school / Pagpapasya tungkol sa trabaho o pag-aaral',
    'Helping children, older persons, persons with disabilities, or other household members / Pagtulong sa mga kasambahay na nangangailangan ng tulong',
    'None of the above / Wala sa mga nabanggit',
], MULTI + ' ' + EXCLUSIVE)

section('Location Specific Information and Priorities', 'Impormasyon para sa Lugar at mga Prayoridad', page=True)
question(15, 23, EASE + ['I do not know how to determine it / Hindi ko alam kung paano ito matutukoy', 'Not sure / Hindi tiyak'])
question(16, 24, EASE + ['I do not know where to find it / Hindi ko alam kung saan ito makikita', 'I have not needed this information / Hindi ko pa kinailangan ang impormasyong ito', 'Not sure / Hindi tiyak'])
question(17, 25, [
    'Flood susceptibility of my barangay or selected location / Flood susceptibility ng aking barangay o selected location',
    'Possible effects of rainfall intensity and duration / Posibleng epekto ng lakas at tagal ng ulan',
    'Current road or travel conditions / Kasalukuyang kondisyon ng daan o biyahe',
    'Pre-event household preparedness actions / Mga paghahanda ng household bago ang baha',
    'Evacuation-center references / Evacuation-center information',
    'Map-based visualization of the selected location and susceptibility result / Map visualization ng napiling lokasyon at susceptibility result',
    'Verified evacuation-center names and map locations / Verified evacuation-center names at map locations',
], 'Select up to three. / Pumili ng hanggang tatlo.')

section('Preparedness Guidance and Evacuation Details', 'Gabay sa Paghahanda at Detalye ng Evacuation Center', page=True)
question(18, 26, [
    'The guidance is too general / Masyadong general ang guidance',
    'It is not specific to my location or situation / Hindi specific sa aking lugar o sitwasyon',
    'The recommended actions are difficult to understand / Mahirap unawain ang recommended actions',
    'It is unclear which action should be done first / Hindi malinaw kung ano ang unang gagawin',
    'The source of the guidance is unclear / Hindi malinaw ang source',
    'The guidance is difficult to access / Mahirap itong ma-access',
    'I do not usually receive preparedness guidance / Karaniwan akong walang natatanggap na guidance',
    'I have not experienced a problem with available guidance / Wala akong naranasang problema',
    'Not sure / Hindi tiyak',
], MULTI + ' ' + EXCLUSIVE)
question(19, 27, ['Name of the center / Pangalan ng center', 'Address or map location / Address o map location', 'Approximate straight-line distance from a selected location / Approximate straight-line distance', 'Contact information / Impormasyon sa pakikipag-ugnayan', 'Information source or responsible office / Source o responsible office', 'Date the information was last verified / Petsa ng huling verification', 'I do not need evacuation-center information / Hindi ko kailangan ang impormasyong ito', 'Not sure / Hindi tiyak'], MULTI + ' ' + EXCLUSIVE)

section('Technology Access and Preferences', 'Access sa Teknolohiya at mga Kagustuhan', page=True)
question(20, 28, ['Yes, my own device / Oo, sarili kong device', 'Yes, a shared device / Oo, shared device', 'No / Hindi', 'Not sure whether it is Android / Hindi tiyak kung Android'])
question(21, 29, ['Always reliable / Palaging reliable', 'Usually reliable / Kadalasang reliable', 'Sometimes reliable / Paminsan-minsang reliable', 'Rarely reliable / Bihirang reliable', 'No internet access / Walang internet access', 'Not sure / Hindi tiyak'])
question(22, 30, ['Manual barangay selection / Manual na pagpili ng barangay', 'Temporary map pin that is not saved as location history / Temporary map pin na hindi ise-save bilang location history', 'One-time GPS access after an explanation and permission request / One-time GPS access matapos ang paliwanag at permission request', 'I am not comfortable providing or selecting a location / Hindi ako komportableng magbigay o pumili ng lokasyon', 'Not sure / Hindi tiyak'], MULTI + ' ' + EXCLUSIVE)
para('The proposed system does not require continuous GPS tracking or background location monitoring. / Hindi nangangailangan ang proposed system ng tuloy-tuloy na GPS tracking o background location monitoring.', 9.5)
question(23, 31, ['English', 'Filipino', 'Both English and Filipino / Parehong English at Filipino'])

section('Additional Information Needs', 'Karagdagang Pangangailangan sa Impormasyon', page=True)
question(24, 32, instruction='Optional. Write your answer, or tell it to the interviewer. / Opsyonal. Isulat ang sagot o sabihin sa tagapanayam.')
for _ in range(8):
    p = para('________________________________________________________________________')
    p.paragraph_format.space_after = Pt(13)
para('Thank you for sharing your experience. / Salamat sa pagbabahagi ng inyong karanasan.', 11, bold=True)
para('Interviewer: Check for one answer in each single-choice question and each matrix row, and no more than three answers in Question 17. Clarify conflicting selections neutrally; do not suggest an answer. Leave unanswered items blank.', 10)
para('Tagapanayam: Suriin ang tig-isang sagot sa single-choice questions at bawat matrix row, at hanggang tatlong sagot sa Tanong 17. Linawin ang magkasalungat na sagot nang walang pagkiling. Iwang blangko ang hindi sinagutan.', 10, italic=True)
para('Privacy notice legal reference: Republic Act No. 10173, National Privacy Commission, https://privacy.gov.ph/data-privacy-act-/', 8)

# Explicitly populate each footer variant for consistent printing in Word.
for other_footer in [sec.first_page_footer, sec.even_page_footer]:
    for node in list(other_footer._element): other_footer._element.remove(node)
    other_footer._element.append(deepcopy(footer._p))

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
