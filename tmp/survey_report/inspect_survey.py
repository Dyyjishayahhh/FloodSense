from collections import Counter
from pathlib import Path
import openpyxl

ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
SOURCE = ROOT / "1. FLOODSENSE-MAIN" / "QUESTIONNAIRES_SURVEY" / "FLOODSENSE_INITIAL_SURVEY_RESULTS_AND_SUMMARY" / "FloodSense Resident Survey Questionnaire (Responses).xlsx"

wb = openpyxl.load_workbook(SOURCE, data_only=True)
ws = wb[wb.sheetnames[0]]
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
rows = [
    [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1)
    if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1))
]

for idx, header in enumerate(headers):
    short = str(header).split("\n")[0]
    print(f"\n[{idx}] {short}")
    values = [str(row[idx]).strip() for row in rows if row[idx] not in (None, "")]
    if idx == 5:
        print(f"nonblank={len(values)} unique={len(set(values))} [names suppressed]")
    elif idx == 31:
        print(f"nonblank={len(values)}")
        for n, value in enumerate(values, 1):
            print(f"  {n}. {value}")
    else:
        for value, count in Counter(values).most_common():
            print(f"  {count:>2} | {value}")
