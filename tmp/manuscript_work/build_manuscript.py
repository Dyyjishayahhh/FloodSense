from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path("1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Revised_Manuscript_Chapters_1_to_3.docx")
FONT = "Arial"


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run)


def set_page_start(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def body(doc, text, first_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    set_run_font(r, italic=italic)
    return p


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, bold=False)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, bold=True)
    return p


def chapter_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(text.upper())
    set_run_font(r, bold=False)
    return p


def numbered_item(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{number}. {text}")
    set_run_font(r)
    return p


def definition(doc, term, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(f"{term}. ")
    set_run_font(r1, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(f"[{text}]")
    set_run_font(r, italic=True)
    return p


def reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Inches(1.0)
sec.bottom_margin = Inches(1.0)
sec.left_margin = Inches(1.5)
sec.right_margin = Inches(1.0)
sec.header_distance = Inches(0.45)
sec.footer_distance = Inches(0.5)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)

for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    style = doc.styles[name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)

# Title page
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("FLOODSENSE: AN ANDROID-BASED EXPERT SYSTEM FOR FLOOD\nSUSCEPTIBILITY ASSESSMENT AND PRE-EVENT PREPAREDNESS\nIN BACOOR CITY")
set_run_font(r, bold=True)

for text in [
    "Undergraduate Thesis",
    "Submitted to the Faculty of the",
    "Department of Computer Studies",
    "Cavite State University - Bacoor City Campus",
    "City of Bacoor, Cavite",
    "",
    "In partial fulfillment of the requirements for the degree",
    "Bachelor of Science in Computer Science",
    "",
    "REYMART V. GOC-ONG",
    "MARTIN LORENZ D. JUANITES",
    "ROMAR T. PULAO",
    "",
    "May 2027",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, bold=text.isupper() and bool(text))

main_sec = doc.add_section(WD_SECTION.NEW_PAGE)
main_sec.page_width = Cm(21.0)
main_sec.page_height = Cm(29.7)
main_sec.top_margin = Inches(1.0)
main_sec.bottom_margin = Inches(1.0)
main_sec.left_margin = Inches(1.5)
main_sec.right_margin = Inches(1.0)
main_sec.header_distance = Inches(0.45)
main_sec.header.is_linked_to_previous = False
page_number(main_sec.header.paragraphs[0])
set_page_start(main_sec, 1)

# CHAPTER 1
chapter_title(doc, "INTRODUCTION")
intro_paragraphs = [
    "Flooding remains one of the most disruptive hazards affecting communities in the Philippines. Recurrent heavy rainfall can interrupt transportation, damage homes and livelihoods, displace families, and create uncertainty about when and where protective action is needed. In Cavite, the Imus River Basin flood-mitigation project includes structures intended to reduce flooding in low-lying portions of Bacoor and nearby areas, confirming that flood exposure is a continuing local concern (Japan International Cooperation Agency [JICA], 2021). The persistence of the hazard makes timely, understandable, and location-specific preparedness information important to residents before an event develops.",
    "Official flood forecasting and warning depend on hydrometeorological observations, technical interpretation, and authorized dissemination. The Philippine Atmospheric, Geophysical and Astronomical Services Administration reported that its flood bulletins communicate current weather conditions, forecast rainfall, water-level trends, possible impacts, and warning messages, while corresponding actions remain coordinated with disaster risk reduction and management offices (PAGASA, 2024). FloodSense does not reproduce that operational function. Instead, it is designed as a scenario-based support tool that allows a resident to examine how a hypothetical rainfall intensity and duration may relate to the susceptibility of a supported location using validated knowledge and geographic information.",
    "The researchers' first requirements survey involved 50 Bacoor City residents from 12 barangays. Ninety-two percent of the respondents had experienced flooding, 58% described previous flooding as severe or waist-deep and higher, only 20% considered themselves well informed about flood-prone areas, and 64% reported making an unsafe or incorrect decision because sufficient area-specific information was unavailable (FloodSense Research Team, 2026a). These findings indicate that receiving a general warning is not always equivalent to understanding the susceptibility of a particular place or knowing what pre-event actions are appropriate.",
    "A supplementary requirements survey involving 40 respondents from 14 barangays produced a similar pattern. Only 22.5% considered themselves well informed, 65% were dissatisfied or very dissatisfied with the flood information available to them, 72.5% reported an unsafe or incorrect decision, and 57.5% experienced difficulty determining the risk of a particular barangay. Respondents gave strong support to a plain-language result, a severity-based action checklist, and distance-ranked evacuation-center information (FloodSense Research Team, 2026b). The two surveys are treated as separate requirements-gathering datasets because they used different respondent groups; their sample sizes are not combined as though they were one survey.",
    "The first and second interviews with the Bacoor Disaster Risk Reduction and Management Office established that historical records, hazard information, preparedness guidance, and evacuation-center records require institutional custody and validation. The second interview further emphasized that one barangay can contain locations with different susceptibility conditions and that an application should not assign a single conclusion to an entire barangay without sufficient spatial support. Coordinates, sitios, zones, affected areas, and verified boundaries are therefore important when the available data permit more specific assessment (FloodSense Research Team, 2026c; 2026d).",
    "The City Planning and Development Coordinator Office identified geographic materials relevant to the study, including elevation, land-use, flood-hazard, and barangay-boundary information (FloodSense Research Team, 2026e). These materials are not automatically equivalent to an operational FloodSense assessment. Each dataset must be reviewed for source, date, geographic resolution, processing history, limitations, and permission before it can be used. HazardHunterPH similarly cautions that scenario-based maps should be treated as references and complemented by current observations, field assessment, and other hazard information (GeoRisk Philippines, 2026).",
    "FloodSense addresses the identified information gap through an Android-based Rule-Based Expert System for flood susceptibility assessment and pre-event preparedness in Bacoor City. A resident selects or confirms a hypothetical rainfall intensity, rainfall duration, and supported location. The backend applies a fixed and documented inference procedure to eligible, versioned knowledge and returns a susceptibility state with an explanation. If available knowledge is insufficient, conflicting, provisional, or outside the supported area, the system returns a limitation state instead of fabricating a classification.",
    "The application presents the assessment through an interactive geographic map and a guided sequence of screens. The sequence introduces the system's purpose and limitations, explains privacy and location use, collects the scenario and location, confirms the assessment request, displays the result and its basis, and presents separately sourced preparedness guidance and verified resource information. Research on interactive flood-risk communication shows that understandable legends, plain language, user-friendly maps, and contextual information can help non-specialists interpret spatial risk (Oubennaceur et al., 2021).",
    "The susceptibility result and the Decision Support System are intentionally separated. The Expert System derives a classification from validated facts and rules, while the Decision Support System presents sourced preparedness content appropriate to the result and the user's stated situation. Recent reviews of explainable decision support emphasize that users should be able to understand the basis of a computerized recommendation and that rule-based explanations can make decision paths transparent (Kostopoulos et al., 2024). FloodSense therefore presents matched conditions, source and version information, limitations, and plain-language explanations without claiming authority to issue an evacuation order.",
    "A separate Web Administration System supports the maintenance of verified geographic records, approved scenario parameters, source metadata, preparedness guidance, evacuation-center records, and publication status. Ordinary administrators are not permitted to create or alter the raw inference method, conflict-resolution behavior, or expert rules. Changes to the research algorithm or rule knowledge require controlled research review, expert validation, versioning, testing, and explicit authorization. This boundary responds directly to the thesis adviser's instruction that operational administrators should not manipulate a research algorithm they did not develop (FloodSense Research Team, 2026f).",
    "FloodSense is therefore positioned as a complementary preparedness application rather than a live monitoring or forecasting platform. It does not continuously poll rainfall stations, run an hourly background timer, issue autonomous alerts, determine road passability, report evacuation-center occupancy, or replace PAGASA and BDRRMO advisories. The current scope favors a defensible scenario-based assessment because water-level and automated-rainfall-gauge datasets have not yet been methodologically connected to the inference model. This limitation prevents the system from presenting unvalidated numerical relationships as scientific conclusions.",
]
for t in intro_paragraphs:
    body(doc, t)

section_heading(doc, "Statement of the Problem")
body(doc, "The study aimed to develop an Android-based Expert System that provides explainable, location-specific flood susceptibility assessment and pre-event preparedness support for residents of Bacoor City. The research problems were derived from two existing requirements surveys, interviews with the BDRRMO and CPDCO, the thesis adviser consultation, current system records, and recent literature. The evidence was organized according to the five cause categories of Measurement, Manpower, Machinery, Material, and Method to ensure that the problems describe observed conditions rather than assumptions.")
body(doc, "Under Measurement, residents lack a consistent way to translate rainfall intensity, duration, geographic attributes, and historical information into an understandable susceptibility state for a supported location. The first survey found that only 20% of respondents considered themselves well informed and that 64% had made an unsafe or incorrect decision because of insufficient information. The supplementary survey likewise found that only 22.5% felt well informed, while 57.5% had difficulty determining the risk of a particular barangay. How can an Android-based Expert System use validated rainfall-scenario and geographic facts to provide residents with an understandable and explainable flood susceptibility assessment for a supported location?")
body(doc, "Under Manpower and Material, the BDRRMO and CPDCO hold or identify institutional knowledge, historical records, geographic materials, preparedness guidance, and evacuation-center information, but these sources require authorization, validation, updating, and careful interpretation before public use. The BDRRMO also explained that susceptibility can vary within a barangay and that evacuation-center records may become outdated. How can FloodSense organize source-backed institutional knowledge, geographic information, preparedness content, and verified resource records while preserving provenance, review status, geographic limitations, and the authority of the responsible offices?")
body(doc, "Under Machinery and Method, residents currently depend on fragmented channels and may need to perform several actions before reaching information relevant to their situation. The thesis adviser requested a more automatic, guided flow, location support, map response to the chosen scenario, preparedness guidance, and first-launch privacy information. At the same time, the adviser restricted ordinary administrators from changing expert rules, and the research team retained a scenario-based boundary instead of background monitoring. How can the mobile and web components provide a guided, low-friction, privacy-aware assessment process while keeping the inference method fixed, protecting expert knowledge from unauthorized editing, and avoiding unsupported real-time claims?")
body(doc, "Taken together, the identified problems lead to the overarching question: How can FloodSense integrate validated rule-based inference, hypothetical rainfall scenarios, supported geographic selection, explainable map visualization, protected data governance, preparedness guidance, and verified resource information to strengthen flood susceptibility awareness and pre-event preparedness in Bacoor City without replacing official forecasts, warnings, or emergency decisions?")

section_heading(doc, "Objectives of the Study")
body(doc, "Generally, the study aimed to develop FloodSense, an Android-based Rule-Based Expert System for explainable flood susceptibility assessment and pre-event preparedness in Bacoor City.")
body(doc, "Specifically, it aimed to:")
objectives = [
    "identify the flood-information, location-awareness, preparedness, governance, and usability needs of Bacoor City residents and relevant local government offices using the available requirements surveys, interviews, consultation records, and the ongoing validated data-gathering process;",
    "analyze authorized flood-related knowledge, geographic records, rainfall-scenario references, historical information, preparedness guidance, and evacuation-center data to define supported facts, limitations, provenance requirements, and the fixed inference procedure;",
    "develop an Android application and supporting web services that provide first-launch privacy and limitation information, a guided rainfall-scenario and location flow, explainable susceptibility results, an interactive map, sourced preparedness guidance, and verified resource details;",
    "develop a role-controlled Web Administration System for maintaining authorized reference data, approved parameters, sources, DSS guidance, geographic records, and evacuation-center information without allowing ordinary administrators to edit the raw expert rules or inference algorithm;",
    "test and evaluate the functional suitability, performance efficiency, compatibility, interaction capability, reliability, security, maintainability, flexibility, and safety of FloodSense using appropriate software tests and a validated user-evaluation instrument based on ISO/IEC 25010:2023.",
]
for i, item in enumerate(objectives, 1):
    numbered_item(doc, i, item)

section_heading(doc, "Theoretical Framework")
body(doc, "The theoretical framework explains how FloodSense converts validated research and institutional inputs into an explainable scenario-based output. The framework separates the source and governance layer, fixed Expert System inference, spatial processing, resident interaction, Decision Support System guidance, and verified resource presentation. This separation prevents a content edit from silently changing the susceptibility conclusion and makes each output traceable to its controlling source and version.")
placeholder(doc, "Insert the revised FloodSense theoretical framework here")
subheading(doc, "Authentication, Onboarding, and Privacy Module")
body(doc, "This module presents the application's purpose, scenario-based nature, limitations, privacy information, and location-use explanation before precise location is requested. Location access, if included in the final approved design, is user-triggered and foreground-only. Manual map or area selection remains the safe fallback. Coordinates are retained only for the active request unless a separately justified and approved purpose is documented.")
subheading(doc, "Scenario and Location Input Module")
body(doc, "This module allows the resident to choose or confirm hypothetical rainfall intensity and duration and to select a supported location. The values are planning inputs and do not describe current weather. A temporary pin may be resolved against approved geographic polygons. A result is not produced for an unsupported point or incomplete scenario.")
subheading(doc, "Protected Knowledge and Inference Module")
body(doc, "This module contains the fixed forward-chaining procedure and eligible, versioned research knowledge. It evaluates facts for the selected scenario and location and returns Low, Moderate, High, Very High, Uncertain, Insufficient Data, or Outside Supported Area as appropriate. Raw expert rules, rule conditions, priority handling, and algorithm behavior are not operational-admin controls.")
subheading(doc, "Spatial Processing and Dynamic Map Module")
body(doc, "GeoDjango and PostGIS manage supported polygons, points, containment queries, and geographic relationships. The Android map displays the result through both color and text. The basemap supplies geographic context only; it is not the source of susceptibility knowledge. A color change is produced by the evaluated scenario and published data, not by an administrator manually painting a conclusion.")
subheading(doc, "Explainable Result and Decision Support Module")
body(doc, "The result screen presents the supported area, scenario, susceptibility state, matched basis, source version, limitations, and warnings. A separate DSS presents validated pre-event guidance. The guidance can be maintained by authorized personnel when it is sourced, reviewed, and published, but guidance edits cannot alter the susceptibility class.")
subheading(doc, "Administration, Provenance, and Resource Module")
body(doc, "The management portal maintains approved parameters, geographic records, sources, guidance, and verified evacuation-center records using draft, review, approval, activation, and rollback states. Technical rule changes follow a separate research and expert-validation process. Every public record must identify its source, effective date, status, and limitations.")

section_heading(doc, "Significance of the Study")
definition(doc, "For Bacoor City Residents", "FloodSense provides an additional way to explore the susceptibility of a supported location under a hypothetical rainfall scenario, understand why a result was produced, review pre-event preparedness guidance, and identify verified resource information. It is intended to reduce uncertainty without presenting itself as an official warning or guarantee of safety.")
definition(doc, "For the Bacoor Disaster Risk Reduction and Management Office", "The study demonstrates a governed channel for publishing approved preparedness content, source information, and verified resource records while preserving the office's authority over official advisories and emergency directions. It also documents why arbitrary administrator editing of expert rules is inappropriate.")
definition(doc, "For the City Planning and Development Coordinator Office", "The study illustrates how verified geographic and planning information may support community-oriented spatial assessment when source, scale, processing, and limitations are preserved.")
definition(doc, "For the Department of Computer Studies", "The project provides a documented case of combining rule-based reasoning, spatial databases, mobile development, administrative governance, explainability, and user-centered requirements within an undergraduate computer science study.")
definition(doc, "For Future Researchers", "The study offers a reproducible boundary for scenario-based flood susceptibility tools and identifies unresolved areas, including water-level integration, automated-rainfall-gauge data, expert validation, spatial granularity, and larger-scale evaluation.")

section_heading(doc, "Time and Place of the Study")
body(doc, "The formal research and development period extends from September 2026 to May 2027. Requirements gathering began earlier in May 2026 through resident surveys and key informant interviews in Bacoor City. The BDRRMO and CPDCO interviews were conducted at the Bacoor Government Center, while software development, academic consultation, integration, testing, and manuscript preparation were conducted primarily through the Department of Computer Studies at Cavite State University - Bacoor City Campus.")
body(doc, "Development proceeded iteratively as requirements, interface designs, data structures, rule behavior, spatial processing, mobile screens, and administrative functions were reviewed and refined. The current evaluation survey is still being conducted. Dates, final participant counts, response rates, and evaluation findings will be added only after the responses have been completed, validated, and analyzed.")

section_heading(doc, "Scope and Limitation of the Study")
scope_paragraphs = [
    "The study covers an Android application for scenario-based flood susceptibility assessment and pre-event preparedness in Bacoor City. A resident selects or confirms a hypothetical rainfall intensity, duration, and supported location. The fixed Expert System evaluates eligible facts and returns an explainable state. The DSS then displays separately sourced preparedness guidance. FloodSense is not an official forecast, warning, evacuation order, or emergency-dispatch system.",
    "The geographic scope is limited to Bacoor City and to areas for which the researchers possess usable, authorized, sufficiently detailed, and validated records. A barangay boundary identifies an administrative area but does not, by itself, establish one susceptibility conclusion for every place inside it. Where available data do not support finer spatial differences, the application must disclose that limitation rather than imply precision.",
    "The mobile flow includes first-launch limitations and privacy information, rainfall-scenario selection, supported location selection or temporary pin placement, assessment confirmation, map and result presentation, explanation, preparedness guidance, and verified resource details. Device GPS remains optional pending final adviser approval and must not be implemented as continuous background tracking. Manual selection remains available.",
    "The dynamic map may display supported polygons using green, yellow, orange, and red with accompanying Low, Moderate, High, and Very High text. It does not display live flood extent, current water depth, road closures, traffic, river level, satellite rainfall, or evacuation-center occupancy. OpenStreetMap, when used, functions only as a basemap.",
    "Evacuation-center information is limited to records obtained from an authorized custodian and reviewed for currency, coordinates, capacity or facilities when available, verification date, and operational limitations. Distance from a selected point does not prove route safety, accessibility, available capacity, or current opening status.",
    "The custom Web Administration System permits role-controlled management of approved datasets, scenario references, geographic records, sources, DSS guidance, evacuation-center records, publication states, and audits. It does not permit ordinary administrators to create or edit raw expert rules, conflict-resolution logic, or the inference algorithm. Rule changes require research ownership, expert validation, versioning, tests, and explicit authorization.",
    "Water-level and automated-rainfall-gauge datasets are outside the active inference method until their fields, units, temporal resolution, spatial coverage, missing values, datum, license, and methodological contribution have been examined and validated. CSV or Excel import and export are also outside the committed implementation until the required direction, schema, validation, approval, and rollback procedure are confirmed.",
    "The application requires an internet connection for current server-managed assessment, map, guidance, and resource records. A fully offline version, iOS release, live meteorological integration, machine-learning retraining, autonomous alerts, continuous timers, and background location services are outside the present scope.",
    "The existing resident surveys and interviews support requirements analysis but do not constitute the final system evaluation. Findings from the evaluation survey currently being conducted cannot be reported until collection and validation are complete. The final results will be limited by the approved sampling procedure, respondent composition, instrument validity, prototype maturity, and geographic coverage.",
]
for t in scope_paragraphs:
    body(doc, t)

section_heading(doc, "Definition of Terms")
terms = [
    ("Android Application", "The resident-facing FloodSense software that provides the guided scenario, supported location selection, assessment result, interactive map, preparedness guidance, and verified resource information."),
    ("Decision Support System", "The component that presents sourced pre-event preparedness guidance using the latest valid susceptibility result and approved situational inputs without changing the classification or making the resident's final decision."),
    ("Expert System", "A knowledge-based software component that applies a fixed inference procedure to eligible facts and validated rules to derive an explainable result."),
    ("Flood Susceptibility", "The relative potential of a supported location to experience rainfall-related flooding under a defined scenario. In FloodSense, it is not a prediction of the exact time, depth, or extent of an actual flood."),
    ("Forward Chaining", "The inference approach that begins with available scenario and geographic facts and evaluates applicable rules until a conclusion or limitation state is reached."),
    ("Geographic Information System", "A system for storing, processing, querying, and presenting information connected to geographic locations."),
    ("Knowledge Base", "The controlled collection of validated facts, rules, conditions, priorities, explanations, versions, and sources used by the Expert System."),
    ("Location Resolution", "The process of determining which approved geographic feature contains or corresponds to a temporary point selected by the resident."),
    ("Pre-event Preparedness", "Actions completed before a flood emergency, such as preparing essential supplies, planning family communication, protecting important documents, and monitoring official advisories."),
    ("Provenance", "Information that records where a dataset, parameter, rule, guidance item, geographic feature, or resource record came from and how it was reviewed, changed, and approved."),
    ("Scenario-based Assessment", "An assessment initiated from hypothetical rainfall and location inputs chosen or confirmed by the resident while the application is open."),
    ("Susceptibility State", "The system output of Low, Moderate, High, Very High, Uncertain, Insufficient Data, or Outside Supported Area."),
    ("Verified Resource", "An evacuation-center or preparedness record obtained from an authorized custodian and reviewed for source, date, status, coordinates, and stated limitations."),
    ("Web Administration System", "The role-controlled browser interface used to maintain approved data, sources, guidance, geographic records, resource records, and publication states without exposing raw rule editing to ordinary administrators."),
]
for term, text in terms:
    definition(doc, term, text)

# CHAPTER 2
doc.add_page_break()
chapter_title(doc, "REVIEW OF RELATED LITERATURE")
body(doc, "This chapter reviews recent literature and studies published from 2021 through 2026 that support the design boundaries of FloodSense. The discussion focuses on flood susceptibility, geographic information, risk communication, explainable decision support, mobile interaction, data governance, and iterative software development. The review distinguishes susceptibility assessment from operational forecasting and identifies the research gap addressed by a scenario-based, locally governed, explainable Android application for Bacoor City.")

section_heading(doc, "Related Foreign Literature")
subheading(doc, "Flood Susceptibility and Spatial Factors")
body(doc, "Flood susceptibility describes the relative tendency of an area to experience flooding based on relevant physical, hydrological, environmental, and human-related factors. It differs from a real-time forecast because it characterizes potential under specified conditions rather than predicting the exact occurrence of an event. Debnath et al. (2024) demonstrated that geospatial and expert-informed models can integrate multiple flood-conditioning factors and classify areas into ordered susceptibility zones. Their work also identified data scarcity, spatial resolution, validation, and the availability of long-term water-level and discharge records as continuing limitations.")
body(doc, "Nguyen et al. (2024) combined GIS and an analytical hierarchy process to evaluate urban flood susceptibility using nine factors, expert consultation, and historical flood points. The resulting map classified areas into five levels and was validated using receiver operating characteristic analysis. The study supports the use of explicit factor definitions, expert review, spatial normalization, and independent validation. FloodSense adopts the need for traceability and expert validation but does not copy the AHP method or weights because Bacoor-specific parameters have not yet been approved.")
subheading(doc, "Interactive Mapping and Risk Communication")
body(doc, "Spatial information is most useful to the public when map symbols, legends, labels, and explanations are understandable. Oubennaceur et al. (2021) found that interactive story maps can combine hazard maps, explanatory text, and user-oriented interaction to communicate flood risk to non-specialists. The authors emphasized clear legends, simple categories, careful color selection, and limited technical terminology. These principles support FloodSense's use of text-supported color classes, result explanations, and a guided interface rather than a map that expects residents to interpret raw technical layers.")
body(doc, "Risk communication must also lead users toward appropriate protective understanding. Lin (2023) reported relationships among flood-risk information seeking, mobile-app use, coping beliefs, collective efficacy, and community action. This suggests that a digital platform should not stop at displaying a classification. It should connect the result to comprehensible, source-backed preparedness information while avoiding statements that exceed the authority or evidence of the system.")
body(doc, "Inclusive warning communication also requires more than simply sending a technical message. The United Nations Office for Disaster Risk Reduction (UNDRR, 2025) described the role of mobile technology in expanding access to early-warning communication for groups that may otherwise be missed. Although FloodSense is not an official warning service, this evidence supports readable mobile presentation, clear limitations, and alternative input methods that do not make one device permission the only route to preparedness information.")
subheading(doc, "Decision Support for Flood Preparedness")
body(doc, "Alabbad et al. (2022) developed a web-based flood mitigation decision-support framework that combined maps, property characteristics, scenarios, guidance, and decision-tree analysis. The study illustrates how a DSS can organize alternatives and present context-specific information instead of making an unexplained decision for the user. FloodSense applies this principle to pre-event household preparedness, keeping the DSS logically separate from susceptibility classification.")
body(doc, "Mattos et al. (2022) connected flood modeling, forecast inputs, map visualization, and a web application in an operational flood-alert context. Their architecture demonstrates the value of integrating data, models, and public communication, but it also shows the extensive infrastructure required for actual forecasting. Because FloodSense does not yet have validated continuous hydrometeorological feeds or a calibrated forecasting model, the study reinforces the decision to retain scenario-based assessment rather than imitate a real-time warning service.")
subheading(doc, "Explainable Rule-based Decision Support")
body(doc, "Kostopoulos et al. (2024) reviewed explainable decision-support approaches and identified visual, rule-based, case-based, natural-language, and knowledge-based explanations. Production rules and IF-THEN explanations are particularly relevant when the system must show a traceable decision path. The review also notes that domain knowledge and human-understandable explanations are central to user trust. FloodSense consequently exposes the facts, matched basis, version, and limitations of a result without exposing rule editing to ordinary administrators.")
body(doc, "Galanti et al. (2023) likewise showed that explanations must be intelligible to the people who act on a decision-support output. Although their predictive-process context differs from flood susceptibility, their user evaluation reinforces an important design requirement: accuracy metrics alone do not establish that an explanation is understandable. FloodSense must therefore test the wording, sequence, and visual presentation of explanations with intended users.")
subheading(doc, "Human-centered Mobile Flood Applications")
body(doc, "Alsabhan and Dudin (2023) examined human-computer-interaction concerns in mobile flood forecasting and warning applications. Their work emphasized the difficulty of presenting complex hydrological information on mobile devices and the need to consider accessibility, interaction, and the conditions under which users make decisions. FloodSense addresses the interface problem through short guided steps, visible limitations, plain-language scenario labels, and consistent result presentation while avoiding the real-time claims of the studied systems.")
subheading(doc, "Iterative Requirements and Software Quality")
body(doc, "Hoy and Xu (2023) found that agile requirements engineering can respond to changing client needs through short cycles, user-oriented requirements, use cases, test cases, and concise acceptance criteria, while still requiring disciplined communication and documentation. This supports the iterative process already reflected in FloodSense's consultation records, mockups, implementation guides, and tests. The project therefore uses an Agile iterative SDLC rather than presenting Design Thinking as the complete software-development lifecycle.")
body(doc, "ISO/IEC 25010:2023 defines a product quality model for specifying, measuring, and evaluating ICT products. Its characteristics provide a current basis for evaluating FloodSense beyond whether individual features run. Functional suitability, performance efficiency, compatibility, interaction capability, reliability, security, maintainability, flexibility, and safety will be considered when the final evaluation instrument and respondent plan are approved (International Organization for Standardization [ISO], 2023).")

section_heading(doc, "Related Local Literature")
subheading(doc, "Flood Conditions and Institutional Responsibility in Cavite")
body(doc, "JICA (2021) reported the inauguration of flood-mitigation facilities in the Imus River Basin, including structures intended to reduce flooding in low-lying portions of Bacoor and Imus. The project establishes the continuing relevance of basin-level flood management to Bacoor. It also shows that a resident-facing application must coexist with engineering interventions and government operations rather than imply that information technology alone controls the hazard.")
subheading(doc, "Official Flood Information and Warning Boundaries")
body(doc, "PAGASA's 2024 annual report describes flood bulletins as products of near-real-time monitoring, forecast rainfall, water-level trends, possible impacts, and warning protocols coordinated with local disaster offices (PAGASA, 2024). FloodSense does not have those operational inputs or institutional authority. Accordingly, its interface must label rainfall inputs as hypothetical and direct residents to official channels during actual events.")
subheading(doc, "Scenario-based Hazard Information")
body(doc, "GeoRisk Philippines (2026) presents HazardHunterPH as a platform for examining hazard information while cautioning that map outputs are references rather than absolute predictions. The platform's disclaimer stresses the need to combine maps with real-time data and field assessment. This boundary directly supports FloodSense's limitation states, source labels, and refusal to treat a boundary or basemap as sufficient proof of susceptibility.")
subheading(doc, "Local Knowledge and Preparedness Requirements")
body(doc, "The BDRRMO interviews documented the need to preserve authoritative data, avoid barangay-wide generalization when conditions vary inside the boundary, validate evacuation-center records, and maintain preparedness guidance from authorized sources (FloodSense Research Team, 2026c; 2026d). The CPDCO interview identified available geographic materials but did not authorize the researchers to treat every identified layer as a validated system input (FloodSense Research Team, 2026e). These records establish the governance and data-validation requirements that distinguish FloodSense from a generic map application.")

section_heading(doc, "Related Foreign Studies")
body(doc, "Recent foreign studies demonstrate several possible flood-information architectures. Oubennaceur et al. (2021) focused on interactive communication of mapped flood risk; Alabbad et al. (2022) developed a decision-support framework for mitigation alternatives; Mattos et al. (2022) connected models and a web application for flood alerts; Noori and Bonakdari (2023) used expert-informed GIS modeling; Alsabhan and Dudin (2023) examined a human-centered mobile warning application; Debnath et al. (2024) compared geospatial and expert-informed susceptibility models; and Nguyen et al. (2024) developed and validated an urban susceptibility map. Together, these studies show that public-facing flood systems require defensible data, a clearly defined method, validation, understandable maps, and communication appropriate to the authority of the system.")
body(doc, "The studies also clarify what FloodSense does not yet possess. Operational alert applications depend on current sensor or forecast data and calibrated hydrological or hydraulic models. Susceptibility studies depend on selected factors, weights, spatial layers, and validation observations. Because those components cannot be transferred to Bacoor without expert and methodological approval, FloodSense uses only approved local knowledge and returns Insufficient Data when a reliable conclusion cannot be supported.")

section_heading(doc, "Related Local Studies")
body(doc, "Johnson et al. (2021) modeled urban expansion and flood exposure across the Philippines using open geospatial data. Their findings indicate that future urban growth may increase the number of people and developed land exposed to flooding, supporting the need for spatially explicit preparedness information. However, their national-scale exposure modeling cannot be treated as a location-level Bacoor assessment without local validation.")
body(doc, "Nagumo et al. (2022) created a three-dimensional flood-hazard map for a flood-prone Philippine area and emphasized that visualization can help residents understand local flood characteristics where geographic information and disaster records are limited. The study supports accessible map presentation but also demonstrates that visualization must be linked to a documented model and appropriate local data.")
body(doc, "Recent Philippine GIS studies have continued to integrate rainfall, elevation, slope, soil, flood height, and government data in flood-risk mapping. These studies demonstrate the usefulness of structured spatial criteria and validation but do not provide Bacoor-specific thresholds for FloodSense. The current study therefore treats externally developed factor weights and formulas as related evidence, not as values that an administrator may import or invent.")

section_heading(doc, "Synthesis")
body(doc, "The reviewed literature establishes five principles relevant to FloodSense. First, susceptibility is spatial and depends on clearly defined factors and geographic resolution. Second, interactive maps require simple legends, explanations, and contextual limitations. Third, a DSS should organize evidence and guidance while leaving the final decision to people and authorized institutions. Fourth, rule-based explanations can improve transparency, but rules and methods must remain governed and validated. Fifth, operational forecasting requires data and models beyond those presently approved for FloodSense.")
body(doc, "Existing systems commonly emphasize either technical mapping, real-time monitoring, public warning, or general risk communication. FloodSense addresses a narrower local gap: it connects a user-confirmed rainfall scenario and supported Bacoor location to a deterministic, explainable susceptibility result and then to separately sourced pre-event preparedness guidance. Its contribution is not the invention of a new forecast model. It is the governed integration of knowledge, scenario interaction, spatial resolution, explanation, guidance, verified resources, and explicit limitation states in an Android and web architecture.")
body(doc, "The local requirements evidence reinforces this gap. Residents reported limited area-specific awareness, dissatisfaction with available information, unsafe decisions, difficulty determining location-specific risk, and strong interest in actionable guidance and evacuation-center information. The BDRRMO and CPDCO records identify both potential sources and necessary restrictions. These findings justify a system that prioritizes clarity and traceability while refusing unsupported claims of precision or real-time authority.")

# Related systems table
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
r = p.add_run("Table 1. Comparative synthesis of selected related studies and systems")
set_run_font(r)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["Study or system", "Primary approach", "Relevance and remaining gap"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pr = cell.paragraphs[0]
    pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = pr.add_run(text)
    set_run_font(rr, size=9, bold=True)
rows = [
    ("Oubennaceur et al. (2021)", "Interactive flood-risk story map", "Supports plain-language spatial communication; does not provide Bacoor-specific inference."),
    ("Alabbad et al. (2022)", "Web DSS and decision tree", "Supports structured mitigation guidance; differs from household pre-event susceptibility assessment."),
    ("Mattos et al. (2022)", "Model-based flood alert web app", "Shows real-time infrastructure requirements that FloodSense currently excludes."),
    ("Alsabhan and Dudin (2023)", "HCI-centered mobile flood app", "Supports guided mobile interaction; operates in a forecasting context."),
    ("Debnath et al. (2024)", "GIS and expert-informed susceptibility models", "Supports multiple spatial factors and validation; parameters cannot be transferred without local review."),
    ("Nguyen et al. (2024)", "GIS-AHP urban susceptibility map", "Supports expert weighting and validation; uses a method not yet approved for FloodSense."),
    ("HazardHunterPH (2026)", "National multi-hazard reference platform", "Provides hazard-reference access; explicitly warns that maps are not absolute predictions."),
    ("FloodSense", "Scenario-based rule inference, map, DSS, and governance", "Targets explainable Bacoor assessment; still requires validated local data and final user evaluation."),
]
for row in rows:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        pr = cells[i].paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rr = pr.add_run(text)
        set_run_font(rr, size=9)

# CHAPTER 3
doc.add_page_break()
chapter_title(doc, "METHODOLOGY")
body(doc, "This chapter presents the materials, research design, development method, participants, instruments, data-gathering procedures, governance controls, and planned analyses for FloodSense. It distinguishes completed requirements-gathering activities from the evaluation survey that is still in progress. No diagram is inserted in this chapter; designated markers show where the researchers may later place the revised figures.")

section_heading(doc, "Materials")
body(doc, "FloodSense is implemented through a Flutter and Dart Android application, a Python Django and Django REST Framework backend, GeoDjango spatial services, and a PostgreSQL database with PostGIS. Flutter's documented application-architecture guidance supports separation of interface and data responsibilities in maintainable applications (Flutter, 2026). The mobile interface uses flutter_map for interactive map presentation, while the custom management portal uses Django templates and controlled administrative workflows. Development records also identify Visual Studio Code, Git, Figma or equivalent interface-design tools, and document-processing software as project resources.")
body(doc, "The software repository contains unit, widget, API, spatial-query, permission, workflow, and integration tests. Development uses fictional or provisional records until a source has been reviewed and approved. Approved research data and provisional development fixtures are maintained separately to prevent demonstration content from being presented as official Bacoor information.")
placeholder(doc, "PENDING RESEARCH INFORMATION: Insert the final development-computer and Android test-device specifications after the team verifies the actual equipment used")

section_heading(doc, "Method")
body(doc, "The study uses a developmental research approach supported by an Agile iterative software-development lifecycle. The choice reflects the team's actual process: requirements were gathered, converted into bounded features, implemented in small increments, tested, reviewed during consultations, and revised when evidence or adviser feedback changed the design. Agile requirements engineering supports short feedback cycles, testable requirements, and adaptation when user or stakeholder needs evolve (Hoy & Xu, 2023). Design Thinking activities were used during early problem discovery and interface exploration, but they are not presented as the complete SDLC.")
placeholder(doc, "Insert the revised Agile iterative development methodology diagram here")
subheading(doc, "Requirements and Evidence Review")
body(doc, "The researchers reviewed survey summaries, interview transcripts, adviser-consultation records, concept documents, feature specifications, geographic-data notes, existing interface designs, and implementation records. Each requested feature was classified as supported, unresolved, deferred, or outside the current scope. Requirements that depended on unexamined datasets or unspecified scientific relationships were not converted into algorithm behavior.")
subheading(doc, "Planning and Analysis")
body(doc, "The team defined the scenario-based boundary, user roles, assessment inputs, limitation states, output explanations, data provenance, and authorization rules. The five-M fishbone categories were used to connect observed causes to the research problems. User stories and acceptance criteria were prepared for the resident flow, map interaction, location handling, assessment, guidance, resources, and administration.")
subheading(doc, "Design")
body(doc, "The design stage organized the Android application, web services, spatial database, Expert System, DSS, and management portal into components with distinct responsibilities. Interface designs were revised toward a guided sequence instead of one long form. First-launch privacy and limitation information, readable susceptibility states, map legends, explanation content, and safe fallbacks were incorporated into the design.")
placeholder(doc, "Insert revised system and process diagrams here after completing the adviser-required diagram corrections")
subheading(doc, "Implementation")
body(doc, "Implementation proceeds through small, testable increments. The Android application sends validated requests to versioned backend endpoints. The backend resolves supported locations, assembles scenario facts, evaluates the fixed rule set, stores or retrieves source-backed guidance and resources, and returns structured responses. Ordinary administrative routes exclude raw rule editing. Provisional data remain visibly separated from approved records.")
subheading(doc, "Testing")
body(doc, "Each increment is tested at the unit and integration levels before it is included in the end-to-end resident or administrator flow. Tests cover request validation, deterministic rule execution, conflict and insufficient-data states, geographic containment, stale response handling, permissions, publication states, map consistency, guidance separation, and safe error messages. The final product evaluation will use the approved ISO/IEC 25010:2023 instrument after the current survey and evaluator plan are completed.")
subheading(doc, "Review and Iteration")
body(doc, "Consultation findings and test failures are converted into documented revisions. A change that affects the research method, scientific parameters, privacy behavior, or system boundary requires corresponding updates to the manuscript, specifications, diagrams, implementation, and tests. This process prevents the code and manuscript from describing different systems.")

section_heading(doc, "Research Design")
body(doc, "The requirements component uses descriptive research to summarize resident responses and qualitative key informant evidence. The development component uses developmental research to design, implement, and verify the software artifact. The evaluation component will use descriptive quantitative measures and structured qualitative feedback after the ongoing instrument administration is complete. The study does not claim an experimental design or causal effect on flood outcomes.")

section_heading(doc, "Research Locale")
body(doc, "The study focuses on the City of Bacoor, Cavite. Requirements evidence was obtained from Bacoor residents and relevant city offices. Software development and academic review were primarily conducted through Cavite State University - Bacoor City Campus. System outputs are limited to supported Bacoor locations represented by authorized and validated records.")

section_heading(doc, "Participants of the Study")
body(doc, "Completed requirements gathering included a first survey of 50 residents from 12 barangays and a supplementary survey of 40 respondents from 14 barangays. These samples are reported separately because they were not administered as one combined survey. Key informants included representatives of the BDRRMO and CPDCO whose institutional responsibilities were relevant to flood information, preparedness, and geographic data.")
placeholder(doc, "PENDING RESEARCH DATA: The new survey and final system evaluation are still being conducted. Insert the approved sampling method, target population, final respondent classifications, frequency, percentage, inclusion criteria, and response rate after validation")

section_heading(doc, "Research Instruments")
body(doc, "The completed requirements instruments consisted of structured resident questionnaires and semi-structured interview guides. The questionnaires gathered flood experience, access to information, perceived clarity and timeliness, location-specific awareness, previous decision difficulties, desired guidance, and interest in evacuation-resource information. Interview guides examined institutional workflows, data availability, geographic resolution, preparedness content, system governance, and validation requirements.")
body(doc, "The final evaluation instrument is intended to measure applicable ISO/IEC 25010:2023 product-quality characteristics and gather task-based usability observations. The instrument must be reviewed for wording, relevance, scale interpretation, respondent suitability, and alignment with the implemented prototype before administration.")
placeholder(doc, "PENDING RESEARCH DATA: Insert the final validated evaluation questionnaire, validation procedure, reliability evidence if required, and scoring interpretation after adviser approval")

section_heading(doc, "Data Gathering Procedure")
body(doc, "For requirements gathering, the researchers distributed resident questionnaires and conducted key informant interviews. Responses were summarized using frequencies, percentages, means where appropriate, and thematic grouping of open-ended answers. Interview transcripts were reviewed to identify authoritative sources, geographic and operational constraints, preparedness needs, data-custody concerns, and system boundaries.")
body(doc, "For the ongoing evaluation, participants will receive an explanation of the study, applicable consent and privacy information, and a structured set of representative tasks. They will use the implemented prototype only after the test data and system state have been checked. Responses will be encoded, reviewed for completeness, and analyzed only after the collection period closes.")
placeholder(doc, "PENDING RESEARCH DATA: Insert the actual survey dates, administration mode, completed-response count, exclusions, and data-cleaning procedure after the ongoing survey is finished")

section_heading(doc, "Knowledge Acquisition and Rule Validation")
body(doc, "Knowledge acquisition begins with authorized records, interviews, official guidance, and the research method. Candidate facts and rules are documented with their source, owner, effective date, geographic applicability, units, assumptions, and limitations. The researchers convert approved relationships into structured conditions and expected conclusions. Each rule is reviewed using representative, boundary, conflict, and insufficient-data cases before activation.")
body(doc, "Ordinary administrators cannot create, edit, or publish raw rules. A future rule change requires a research-owned proposal, subject-matter review, test cases, version assignment, approval, activation, and rollback capability. DSS guidance follows a separate content workflow and cannot modify the susceptibility result.")

section_heading(doc, "Data and Geographic Validation")
body(doc, "Every dataset is inventoried by source, custodian, date, format, coordinate reference system, geographic coverage, resolution, attributes, units, missing values, processing history, license, and restrictions. Polygon validity and point coordinates are checked before use. A barangay boundary can support location identification but cannot independently supply a susceptibility class. Processed data must remain traceable to the unaltered authoritative source.")
body(doc, "Water-level and automated-rainfall-gauge files will not be added to the inference method merely because they are available. The team must first document their measurement definitions, time interval, station coverage, datum, missingness, quality controls, and proposed scientific relationship to susceptibility. Any formula or threshold requires methodological and expert validation.")

section_heading(doc, "System Testing")
body(doc, "Unit tests verify models, serializers, rule conditions, conflict handling, provenance policies, controllers, and parsers. Integration tests verify API contracts, database operations, spatial queries, role permissions, publication workflows, and mobile-backend communication. Widget and usability-oriented tests verify guided navigation, readable states, error recovery, map interaction, manual fallback, first-launch disclosure, and the separation of guidance from classification.")
body(doc, "End-to-end validation will compare the same scenario, location, and active rule version across the map, detailed assessment, explanation, and DSS context. The system must return the same supported conclusion or the same limitation state. Tests must also confirm that an ordinary administrator cannot modify the protected inference method through either the visible portal or a direct backend request.")

section_heading(doc, "Statistical Treatment of Data")
body(doc, "The completed requirements surveys were summarized using frequency and percentage for categorical responses and arithmetic or weighted means for scaled responses where the instrument supported those calculations. Qualitative answers were grouped into recurring themes while retaining contradictory or minority responses when they affected requirements or limitations.")
numbered_item(doc, 1, "Percentage. Percentage may be computed as P = (f / n) x 100, where P is the percentage, f is the frequency of a response, and n is the number of valid responses for the item.")
numbered_item(doc, 2, "Weighted mean. For an approved scaled item, the weighted mean may be computed as x-bar = sum(fx) / N, where f is the frequency, x is the assigned scale value, and N is the total number of valid responses.")
numbered_item(doc, 3, "Qualitative synthesis. Interview and open-ended responses will be organized into evidence-based themes, compared across sources, and used to explain requirements, contradictions, and limitations.")
placeholder(doc, "PENDING RESEARCH DATA: Do not insert a final Likert interpretation table, sample-size computation, reliability statistic, or evaluation result until the approved questionnaire, completed responses, and adviser-approved analysis plan are available")

section_heading(doc, "Ethical, Privacy, and Governance Considerations")
body(doc, "Participation in surveys and evaluation activities must be voluntary and based on an appropriate explanation of purpose, use, and confidentiality. The manuscript will report aggregated results and will not publish unnecessary personal identifiers. Precise location, if approved for the application, is requested only after a purpose explanation and is used temporarily for the active assessment unless a separate retention purpose is authorized.")
body(doc, "Research and institutional records are handled according to their custody, permission, and publication restrictions. The application distinguishes provisional demonstration data from approved information. Public outputs display source and limitation information, while administrative changes are authenticated, authorized, reviewed, versioned, and auditable.")

section_heading(doc, "Pending Methodology Components")
body(doc, "The following components remain intentionally incomplete because the necessary evidence is not yet available: the final new-survey sample and results; the final system-evaluation participants; the validated evaluation instrument; the confirmed GPS requirement; the decision between automatic evaluation and final user confirmation; the role of water-level and station data; the exact CSV or Excel import-export requirement; and the final ownership of approval and publication actions. These items must be completed through adviser confirmation, expert validation, or finished data collection rather than assumption.")

# REFERENCES
references_title = chapter_title(doc, "REFERENCES")
references_title.paragraph_format.page_break_before = True
section_heading(doc, "Websites and Institutional Publications")
web_refs = [
    "Flutter. (2026). Architecting Flutter apps. https://docs.flutter.dev/app-architecture",
    "GeoRisk Philippines. (2026). HazardHunterPH: Hazard assessment at your fingertips. https://hazardhunter.georisk.gov.ph/map",
    "International Organization for Standardization. (2023). ISO/IEC 25010:2023 systems and software engineering - Systems and software Quality Requirements and Evaluation (SQuaRE) - Product quality model. https://www.iso.org/standard/78176.html",
    "Japan International Cooperation Agency. (2021, October 1). Flood risk protection project in Cavite inaugurated. https://www.jica.go.jp/english/overseas/philippine/information/press/2021/211001.html",
    "Philippine Atmospheric, Geophysical and Astronomical Services Administration. (2024). Annual report 2024. https://pubfiles.pagasa.dost.gov.ph/pagasaweb/files/transparency/Fiscal%20Year%202024%20Annual%20Report.pdf",
    "United Nations Office for Disaster Risk Reduction. (2025). Mobile technology expanding inclusive early warning communication. https://www.undrr.org/resource/case-study/mobile-technology-expanding-inclusive-early-warning-communication",
]
for ref in web_refs:
    reference(doc, ref)

section_heading(doc, "Journals and Conference Papers")
journal_refs = [
    "Alabbad, Y., Yildirim, E., & Demir, I. (2022). Flood mitigation data analytics and decision support framework: Iowa Middle Cedar Watershed case study. Science of the Total Environment, 814, 152768. https://doi.org/10.1016/j.scitotenv.2021.152768",
    "Alsabhan, W., & Dudin, B. (2023). Real-time flood forecasting and warning: A comprehensive approach toward HCI-centric mobile app development. Multimodal Technologies and Interaction, 7(5), 44. https://doi.org/10.3390/mti7050044",
    "Debnath, J., Sahariah, D., Nath, N., Saikia, A., Lahon, D., Islam, M. N., Hashimoto, S., Meraj, G., Kumar, P., Singh, S. K., Kanga, S., & Chand, K. (2024). Modelling on assessment of flood risk susceptibility at the Jia Bharali River basin in Eastern Himalayas by integrating multicollinearity tests and geospatial techniques. Modeling Earth Systems and Environment, 10, 2393-2419. https://doi.org/10.1007/s40808-023-01912-1",
    "Galanti, R., Coma-Puig, B., de Leoni, M., Carmona, J., & Navarin, N. (2023). An explainable decision support system for predictive process analytics. Engineering Applications of Artificial Intelligence, 120, 105904. https://doi.org/10.1016/j.engappai.2023.105904",
    "Hoy, Z., & Xu, M. (2023). Agile software requirements engineering challenges-solutions: A conceptual framework from systematic literature review. Information, 14(6), 322. https://doi.org/10.3390/info14060322",
    "Johnson, B. A., Estoque, R. C., Li, X., Kumar, P., Dasgupta, R., Avtar, R., & Magcale-Macandog, D. B. (2021). High-resolution urban change modeling and flood exposure estimation at a national scale using open geospatial data: A case study of the Philippines. Computers, Environment and Urban Systems, 90, 101704. https://doi.org/10.1016/j.compenvurbsys.2021.101704",
    "Kostopoulos, G., Davrazos, G., & Kotsiantis, S. (2024). Explainable artificial intelligence-based decision support systems: A recent review. Electronics, 13(14), 2842. https://doi.org/10.3390/electronics13142842",
    "Lin, C. A. (2023). Flood risk management via risk communication, cognitive appraisal, collective efficacy, and community action. Sustainability, 15(19), 14191. https://doi.org/10.3390/su151914191",
    "Mattos, T. S., Oliveira, P. T. S., Bruno, L. S., Carvalho, G. A., Pereira, R. B., Crivellaro, L. L., Lucas, M. C., & Roy, T. (2022). Towards reducing flood risk disasters in a tropical urban basin by the development of flood alert web application. Environmental Modelling & Software, 151, 105367. https://doi.org/10.1016/j.envsoft.2022.105367",
    "Nagumo, N., Ohara, M., Fujikane, M., Inoue, T., Hiramatsu, Y., & Jaranilla-Sanchez, P. A. (2022). Creation of a 3D flood hazard map for a flood-prone area in the Republic of the Philippines and dissemination of the mapping technology. E-journal GEO, 17(1), 123-136. https://doi.org/10.4157/ejgeo.17.123",
    "Nguyen, H. N., Fukuda, H., & Nguyen, M. N. (2024). Assessment of the susceptibility of urban flooding using GIS with an analytical hierarchy process in Hanoi, Vietnam. Sustainability, 16(10), 3934. https://doi.org/10.3390/su16103934",
    "Noori, A., & Bonakdari, H. (2023). A GIS-based fuzzy hierarchical modeling for flood susceptibility mapping: A case study in Ontario, Eastern Canada. Environmental Sciences Proceedings, 25(1), 62. https://doi.org/10.3390/ECWS-7-14242",
    "Oubennaceur, K., Chokmani, K., El Alem, A., & Gauthier, Y. (2021). Flood risk communication using ArcGIS StoryMaps. Hydrology, 8(4), 152. https://doi.org/10.3390/hydrology8040152",
]
for ref in journal_refs:
    reference(doc, ref)

section_heading(doc, "Research Records")
research_refs = [
    "FloodSense Research Team. (2026a). FloodSense first resident requirements survey summary [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026b). FloodSense supplementary resident requirements survey analysis [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026c). First key informant interview transcript with the Bacoor Disaster Risk Reduction and Management Office [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026d). Second key informant interview transcript with the Bacoor Disaster Risk Reduction and Management Office [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026e). Key informant interview transcript with the Bacoor City Planning and Development Coordinator Office [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026f). Thesis adviser consultation transcript and current system decision record [Unpublished research record]. Cavite State University - Bacoor City Campus.",
]
for ref in research_refs:
    reference(doc, ref)

# Keep title page unnumbered.
doc.sections[0].header.is_linked_to_previous = False
doc.sections[0].header.paragraphs[0].clear()

# Apply table widths and repeat header.
widths = [Inches(1.55), Inches(1.55), Inches(3.25)]
for row in table.rows:
    for idx, cell in enumerate(row.cells):
        cell.width = widths[idx]
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for edge in ("top", "left", "bottom", "right"):
            node = tc_mar.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                tc_mar.append(node)
            node.set(qn("w:w"), "80")
            node.set(qn("w:type"), "dxa")
tr_pr = table.rows[0]._tr.get_or_add_trPr()
tbl_header = OxmlElement("w:tblHeader")
tbl_header.set(qn("w:val"), "true")
tr_pr.append(tbl_header)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
