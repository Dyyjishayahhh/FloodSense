from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn

path = Path(r"1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Revised_Manuscript_Chapters_1_to_3.docx")
doc = Document(path)
text = "\n".join(p.text for p in doc.paragraphs)
sec = doc.sections[1] if len(doc.sections) > 1 else doc.sections[0]
years = sorted(set(int(y) for y in re.findall(r"\((20\d{2})", text[text.find("REFERENCES"):])) )

print("file_bytes", path.stat().st_size)
print("sections", len(doc.sections))
print("paragraphs", len(doc.paragraphs))
print("tables", len(doc.tables))
print("inline_shapes", len(doc.inline_shapes))
print("page_cm", round(sec.page_width.cm, 2), round(sec.page_height.cm, 2))
print("margins_in", round(sec.left_margin.inches, 2), round(sec.right_margin.inches, 2), round(sec.top_margin.inches, 2), round(sec.bottom_margin.inches, 2))
print("normal_font", doc.styles["Normal"].font.name, doc.styles["Normal"].font.size.pt)
print("references_start", text.find("REFERENCES"))
print("reference_years", years)
print("pre2021_years", [y for y in years if y < 2021])
print("pending_notes", text.count("PENDING RESEARCH"))
print("diagram_markers", text.lower().count("diagram here"))
print("old_rule_edit_claims", len(re.findall(r"administrator(?:s)? (?:can|may) (?:create|edit|modify).*rule", text, re.I)))
print("has_sop", "Statement of the Problem" in text)
print("has_chapters", all(h in text for h in ["INTRODUCTION", "REVIEW OF RELATED LITERATURE", "METHODOLOGY", "REFERENCES"]))
print("title", next((p.text for p in doc.paragraphs if p.text.startswith("FLOODSENSE:")), "MISSING"))
for i, section in enumerate(doc.sections):
    pg = section._sectPr.find(qn("w:pgNumType"))
    print("section", i, "page_start", pg.get(qn("w:start")) if pg is not None else None)
