from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX = Path(
    r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense\1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11\FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"
)


def text_hash(values: list[str]) -> str:
    return sha256("\n".join(values).encode("utf-8")).hexdigest()


def find_paragraph(doc: Document, exact: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.strip() == exact]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph {exact!r}; found {len(matches)}")
    return matches[0]


def paragraph_index(doc: Document, exact: str) -> int:
    matches = [i for i, p in enumerate(doc.paragraphs) if p.text.strip() == exact]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph index for {exact!r}; found {len(matches)}")
    return matches[0]


def replace_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def clone_paragraph_before(anchor: Paragraph, template: Paragraph, text: str) -> Paragraph:
    p = OxmlElement("w:p")
    if template._p.pPr is not None:
        p.append(deepcopy(template._p.pPr))
    anchor._p.addprevious(p)
    paragraph = Paragraph(p, anchor._parent)
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    return paragraph


def remove_between(doc: Document, start_text: str, end_text: str) -> tuple[Paragraph, Paragraph]:
    start = find_paragraph(doc, start_text)
    end = find_paragraph(doc, end_text)
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        if node.tag.endswith("}p"):
            node.getparent().remove(node)
        node = nxt
    return start, end


def set_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    replace_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_row(table, row_index: int, values: list[str]) -> None:
    row = table.rows[row_index]
    if len(row.cells) != len(values):
        raise RuntimeError("Column count mismatch")
    for cell, value in zip(row.cells, values):
        set_cell(cell, value)


def resize_rows(table, count: int) -> None:
    while len(table.rows) > count:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < count:
        new_tr = deepcopy(table.rows[-1]._tr)
        table._tbl.append(new_tr)
        for cell in table.rows[-1].cells:
            set_cell(cell, "")


def mark_table_rows(table) -> None:
    for idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if idx == 0 and tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        if tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


doc = Document(DOCX)

# Integrity checkpoints for content outside the authorized Chapter 1 scope and Chapter 2 review/reference areas.
before_intro = text_hash([p.text for p in doc.paragraphs[:111]])
before_definitions = text_hash([p.text for p in doc.paragraphs[122:144]])
before_chapter3 = text_hash([p.text for p in doc.paragraphs[308:344]])
before_appendices = text_hash([p.text for p in doc.paragraphs[376:]])

# Chapter 1: one continuous Scope and Limitations discussion, as directed in the consultation.
scope_heading = find_paragraph(doc, "Scope and Limitation of the Study")
definitions_heading = find_paragraph(doc, "Definition of Terms")
scope_paragraph = doc.paragraphs[paragraph_index(doc, "Scope and Limitation of the Study") + 1]
scope_text = (
    "FloodSense is a scenario-based, pre-event flood-susceptibility assessment and preparedness-support system intended for Bacoor City residents and authorized administrators. A resident authenticates, completes the required legal and onboarding steps, selects or confirms a hypothetical rainfall intensity and duration, selects or confirms a supported location, and explicitly requests an assessment. The deterministic Rule-Based Expert System returns an explained Low, Moderate, High, or Very High demonstration class, while a separate Decision Support System presents source-backed preparedness guidance without changing that class. The resident application supports manual area selection, map-pin adjustment, optional one-time foreground location, barangay confirmation or correction, interactive geographic display, account and preference controls, and nearest-center lookup when authorized and current center records are available; however, the location request does not continuously track the device, and a displayed center or distance does not verify road passability, accessibility, current opening, available capacity, or safety. The custom Web Administration System provides authenticated dashboards and governed workflows for map and rainfall references, parameter review, DSS content and structured flows, evacuation-center records and staged center import, data sources, and audit history. Ordinary administrators cannot modify raw Expert System rules, rule conditions, priorities, conflict-resolution logic, or algorithm code, and the separate technical Django Admin is reserved for developer or research maintenance; however, the confirmed fixed-schema CSV export and import workflow for permitted Expert System parameter values remains unimplemented and is distinct from the available evacuation-center import workflow. The 47-barangay layer provides administrative reference and has undergone technical processing for software use; however, it is not automatically City-issued or City-verified, contains no flood-susceptibility facts by itself, and no approved operational Bacoor City susceptibility dataset is established. FloodSense does not provide real-time rainfall or water-level monitoring, automatic rainfall updates, background polling, background GPS tracking, autonomous alerts, official forecasts or warnings, automatic evacuation orders, safe-road routing, live center occupancy, machine-learning prediction, self-modifying rules, complete offline operation, resident assessment history, completed evaluation, or verified production deployment. Provisional geographic materials, processing scripts, mockups, screenshots, and technically valid geometries do not establish institutional approval or operational integration. Automated tests provide evidence of software behavior within their test cases, but they do not establish scientific validity, institutional approval, usability, or production readiness. Personal information and one-time foreground location must be handled under Republic Act No. 10173 and the approved consent and privacy process, and direct identifiers from the ongoing survey must not be included in the analysis dataset or public results."
)
replace_text(scope_paragraph, scope_text)
node = scope_paragraph._p.getnext()
while node is not None and node is not definitions_heading._p:
    nxt = node.getnext()
    if node.tag.endswith("}p"):
        node.getparent().remove(node)
    node = nxt

# Chapter 2: move official services and platforms to literature/existing-system discussions.
foreign_lit_anchor = find_paragraph(doc, "Related Local Literature")
body_template = find_paragraph(doc, "Flood susceptibility expresses the relative tendency of a location to experience flooding when relevant conditions are present. It differs from flood hazard, which concerns the probability and physical characteristics of an event; exposure, which concerns people or assets in affected places; vulnerability, which concerns the capacity to be harmed; and risk, which combines hazard and potential consequences. Membele, Naidu, and Mutanga (2022) found that flood-vulnerability mapping in developing countries varies substantially in terminology, indicators, spatial scale, data availability, and validation practice. Their review shows why a public system must name the type of assessment it performs and avoid presenting a general vulnerability or hazard map as a locally validated susceptibility conclusion.")
subheading_template = find_paragraph(doc, "Flood Susceptibility and Scenario-Based Assessment")

foreign_platforms = [
    ("Existing Foreign Flood-Information Platforms", True),
    ("Foreign platforms demonstrate several forms of flood information without constituting transferable evidence for Bacoor City. FloodAdapt uses SFINCS-based scenario simulation and Delft-FIAT impact analysis to compare flood events and adaptation options (Deltares, 2025). The Global Flood Awareness System applies large-scale hydrological modeling to operational river-flood forecasting (Copernicus Emergency Management Service, 2026), while the Iowa Flood Information System combines Iowa-specific observations, radar rainfall, forecasting, and mapped information (Iowa Flood Center, 2024). These services illustrate scenario analysis, forecasting, and public communication, but their models, data, scales, and authority differ from FloodSense.", False),
    ("The Resilience Analysis and Planning Tool combines United States hazard, infrastructure, demographic, and community-resilience layers for emergency-management planning (Federal Emergency Management Agency, 2025). England's National Assessment of Flood and Coastal Erosion Risk similarly communicates modeled long-term risk from several flood sources at national scale (Environment Agency, 2024). These official resources show the value of visible limitations, data context, and geographic lookup; however, they are platforms and institutional assessments rather than FloodSense research studies, and their datasets and classifications cannot validate a Bacoor City susceptibility result.", False),
]
for text, is_heading in foreign_platforms:
    clone_paragraph_before(foreign_lit_anchor, subheading_template if is_heading else body_template, text)

local_platform_anchor = find_paragraph(doc, "Related Foreign Studies")
local_platforms = [
    ("Existing Philippine Hazard-Information and Planning Platforms", True),
    ("UP NOAH remains an important Philippine reference for accessible geographic hazard information. Lagmay et al. (2017) documented the dissemination of near-real-time hazard information and flood maps through Web-GIS using open-source code and application programming interfaces. The dated publication supports discussion of public map-based risk communication, while the general platform should not be treated as a research study with an invented SDLC. UP NOAH differs from FloodSense because its documented hazard-information resources were institutionally developed, whereas FloodSense is a scenario-based research prototype and is not an official forecasting or warning platform.", False),
    ("HazardHunterPH is a GeoRisk Philippines application that produces site-specific indicative multi-hazard assessments from official agency information (GeoRisk Philippines, 2023). Its rapid geographic lookup and source-based reporting are relevant to FloodSense's emphasis on accessible information, provenance, and visible limitations. It remains an official platform rather than a substitute for locally validated FloodSense rules or data, and its reports do not confer forecasting or warning authority on FloodSense.", False),
    ("GeoMapperPH supports governed collection of hazard, exposure, vulnerability, and coping-capacity information by trained government and agency users, while PlanSmartPH Ready to Rebuild organizes information for local-government rehabilitation and recovery planning (DOST-PHIVOLCS, 2022a, 2022b). These platforms illustrate structured data governance and institutional workflows. They are discussed as existing Philippine systems, not retained as Related Studies, because the inspected official pages do not provide a complete research methodology and results section suitable for the TA-required study comparison.", False),
]
for text, is_heading in local_platforms:
    clone_paragraph_before(local_platform_anchor, subheading_template if is_heading else body_template, text)

# Resolve the remaining undated HazardHunterPH citation in the existing local-literature discussion.
hazardhunter_lit = find_paragraph(doc, "GeoRisk Philippines presents HazardHunterPH as a national platform for indicative hazard assessment while advising users that results depend on available government data and do not replace official hazard reports from mandated agencies (GeoRisk Philippines, n.d.). Its source notices and limitations demonstrate responsible hazard communication. The same principle applies to FloodSense: sources, dates, limitations, and validation status should remain visible, and an indicative or provisional record should not be represented as an absolute prediction.")
replace_text(hazardhunter_lit, "GeoRisk Philippines presents HazardHunterPH as a national platform for site-specific indicative multi-hazard assessment using official information from participating government agencies (GeoRisk Philippines, 2023). Its source notices and limitations demonstrate responsible hazard communication. The same principle applies to FloodSense: sources, dates, limitations, and validation status should remain visible, and an indicative or provisional record should not be represented as an absolute prediction.")

# Replace Related Foreign Studies with sources whose methods and results can be inspected.
foreign_heading, local_heading = remove_between(doc, "Related Foreign Studies", "Related Local Studies")
foreign_entries = [
    (
        "Examining Flood Vulnerability Mapping Approaches in Developing Countries: A Scoping Review",
        "Membele, Naidu, and Mutanga (2022) examined how flood-vulnerability mapping was conducted in developing countries. Following PRISMA-ScR procedures, they searched Web of Science and Scopus for 2010-2020 publications, screened 2,549 records, and reviewed 130 eligible studies. They found frequent use of GIS with indicator-based and multicriteria methods, but limited stakeholder participation, sensitivity analysis, and validation, together with recurring problems involving coarse or outdated data.",
        "The review is relevant because it shows that mapping terminology, indicators, spatial scale, participation, and validation must be reported before a flood-related output can be interpreted responsibly. Like FloodSense, the reviewed work concerns geographically organized flood information; however, it synthesizes vulnerability-mapping studies rather than building a resident scenario-assessment application. Its findings identify methodological weaknesses but do not supply Bacoor-specific factors, weights, rules, or validation, and the review was limited to studies that produced a flood-vulnerability map.",
    ),
    (
        "Actionable Information in Flood Risk Communications and the Potential for New Web-Based Tools for Long-Term Planning for Individuals and Community",
        "Mostafiz, Rohli, Friedland, and Lee (2022) investigated how flood-risk information could better support long-term decisions by individuals and communities. They reviewed flood communication through mass media, citizen science, and social media and examined 50 freely available web-based flood-risk portals. Their review found that many tools emphasize short-term information, while micro-scale, actionable, long-term planning information and explicit cost-benefit support remain limited.",
        "The study supports FloodSense's use of understandable result explanations and separate preparedness guidance. Both emphasize that technical flood information must be connected to actions a user can understand; however, the study evaluated a cross-section of portals rather than a deterministic Philippine susceptibility system. Its portal sample was a time-bounded snapshot, and its findings cannot establish what information Bacoor residents need without local survey and stakeholder evidence.",
    ),
    (
        "Accelerated and Interpretable Flood Susceptibility Mapping Through Explainable Deep Learning with Hydrological Prior Knowledge",
        "Wang, Sanderson, Iqbal, and Woo (2025) developed and evaluated an explainable deep-learning approach for flood-susceptibility mapping in a 1,516-square-kilometer area of Northumberland, United Kingdom. They combined official flood records and terrain and environmental variables, compared support-vector-machine, decision-tree, random-forest, standard U-Net, and hydrology-aware U-Net models, used spatial five-fold cross-validation, and applied Grad-CAM interpretation. The proposed model achieved a reported AUC of 0.966 and faster convergence than the standard U-Net.",
        "The study demonstrates how susceptibility models can be evaluated and interpreted, but it also reports static inputs, an inventory that excluded surface-water flooding, and the need for retraining or fine-tuning in other regions. It is relevant to FloodSense's concerns with susceptibility and explanation, yet it differs fundamentally because FloodSense uses deterministic rules rather than machine-learning prediction. The model, features, and reported performance cannot be transferred to Bacoor City without local data, validation, and a different approved method.",
    ),
]
for title, p1, p2 in foreign_entries:
    clone_paragraph_before(local_heading, subheading_template, title)
    clone_paragraph_before(local_heading, body_template, p1)
    clone_paragraph_before(local_heading, body_template, p2)

# Replace Related Local Studies with Philippine studies that report methods and outputs.
local_heading, synthesis_heading = remove_between(doc, "Related Local Studies", "Synthesis of the Reviewed Literature")
local_entries = [
    (
        "GIS-Based Mapping of Flashflood Susceptible Areas in Bukang Liwayway, Kibawe, Bukidnon",
        "Puno et al. (2022) generated a flashflood-susceptibility map for Barangay Bukang Liwayway, Kibawe, Bukidnon using GIS and frequency-ratio analysis. The researchers collected 160 flashflood inventory points and evaluated nine conditioning factors, including terrain, drainage, land cover, vegetation, and rainfall, using satellite, elevation, rainfall, and field information. They reported that 89.8% of the area fell within low, 4.4% within moderate, and 5.7% within high susceptibility, with 25 households exposed to the high category.",
        "The study is similar to FloodSense in producing spatially organized susceptibility information and emphasizing local planning, but its statistical frequency-ratio method and rural geographic context differ from FloodSense's deterministic, user-requested scenario flow. Its sample points, factor relationships, weights, land-cover conditions, and proposed evacuation-site analysis are specific to Bukang Liwayway. They cannot be generalized to Bacoor City without locally appropriate data, expert review, field validation, and institutional authorization.",
    ),
    (
        "Flood Risk Assessment Using GIS-Based Analytical Hierarchy Process in the Municipality of Odiongan, Romblon, Philippines",
        "Gacu, Monjardin, Senoro, and Tan (2022) assessed flood risk in Odiongan, Romblon using a GIS-based analytical hierarchy process under the Sendai Framework. The study combined five hazard, nine vulnerability, and three exposure parameters, obtained comparative judgments from ten experts and end-users, created weighted-overlay maps, and checked results against ground assessment and historical events. It identified nine barangays with high flood risk and reported that 3.26% of the municipality was in the high-risk class.",
        "The work demonstrates a documented approach for combining spatial criteria, expert judgment, and validation. It is relevant to FloodSense's emphasis on geographic evidence and expert review, but it assesses risk by combining hazard, vulnerability, and exposure rather than producing FloodSense's bounded susceptibility output. Its judgments and weights are subjective and locally derived, so they cannot be reused as Bacoor rules or parameters without renewed local evidence and authorization.",
    ),
    (
        "Web-Based Solution for Flood Warning Decision Support in the Province of Leyte, Philippines",
        "Bentoso et al. (2021) developed a sensor-based engineering prototype intended to support flood-warning decisions in the Province of Leyte. The system collected rainfall and water-level readings, transmitted the data to a web interface, presented graphical information, and used thresholds to send warning messages; the authors also reported functional testing of the prototype. The documented output was a working monitoring and alert interface for responsible users, not evidence of province-wide operational deployment.",
        "The Leyte study is relevant because it separates sensed conditions, decision support, and warning communication within a digital flood application. It differs from FloodSense because it relies on continuous sensors and alert thresholds, whereas FloodSense has no continuous monitoring, automatic polling, autonomous alerts, or official warning authority. Its sensor placement, communications, thresholds, and local results cannot be transferred to Bacoor City without new scientific, technical, and institutional validation.",
    ),
]
for title, p1, p2 in local_entries:
    clone_paragraph_before(synthesis_heading, subheading_template, title)
    clone_paragraph_before(synthesis_heading, body_template, p1)
    clone_paragraph_before(synthesis_heading, body_template, p2)

# Synthesis and research gap updated to match the verified study set.
synthesis_heading, gap_heading = remove_between(doc, "Synthesis of the Reviewed Literature", "Research Gap")
synthesis_paragraphs = [
    "The reviewed literature and studies show that flood susceptibility, vulnerability, hazard, exposure, and risk are related but not interchangeable. Membele et al. (2022) identified wide variation in mapping terminology, indicators, scale, participation, and validation, while Puno et al. (2022) and Gacu et al. (2022) demonstrated how local studies combine geographic factors, documented analytical procedures, and validation. Their methods support disciplined geographic assessment, but their factor weights, classifications, and findings cannot be imported into Bacoor City because the study areas, data, scales, and intended outputs differ.",
    "Mostafiz et al. (2022) showed that flood information is more useful when it supports understandable action, and official platforms such as UP NOAH and HazardHunterPH illustrate public geographic communication with visible institutional sources. FloodSense therefore separates its deterministic susceptibility result from preparedness guidance and identifies the scenario, location, source, and limitations presented to the user. Bentoso et al. (2021) demonstrates a different approach based on sensors and warning messages; its contribution is relevant to flood decision support but does not justify adding real-time monitoring or autonomous alerts to FloodSense.",
    "Explainability and validation remain distinct responsibilities. Wang et al. (2025) showed that an interpretable machine-learning model can be evaluated quantitatively, but its static inputs, inventory limitations, and retraining requirement also show why performance in one area does not establish validity elsewhere. FloodSense instead uses deterministic rule matching so that the same approved facts and rule version produce the same result, yet repeatability alone does not validate the rules. Geographic provenance, local expert review, institutional authorization, and appropriate Bacoor data remain necessary.",
    "Across the sources, geographic technologies support storage, analysis, visualization, and communication, but technical capability does not establish scientific or institutional authority. Administrative boundaries can support location resolution without containing flood-susceptibility facts. Official platforms, reviewed studies, and FloodSense also serve different users and disaster-management phases; therefore, similarities in mapping or decision support must not be interpreted as equivalence in data, functions, or authority.",
]
for text in synthesis_paragraphs:
    clone_paragraph_before(gap_heading, body_template, text)

gap_heading, comparison_heading = remove_between(doc, "Research Gap", "Comparative Synthesis of Related Studies")
gap_paragraphs = [
    "Existing research already contributes flood-susceptibility mapping, GIS-based multicriteria assessment, explainable modeling, risk communication, sensor-based warning, and public hazard platforms. The inspected evidence nevertheless leaves a context-specific need for public-facing, location-specific, scenario-based susceptibility information that clearly explains a deterministic result and then presents preparedness guidance as a separate form of support. This is not a claim that FloodSense is the first flood application or that no comparable system exists.",
    "The remaining gap also concerns provenance and transferability. The reviewed studies do not provide an approved set of Bacoor City susceptibility factors, rule thresholds, validated geographic inputs, or preparedness content that can be adopted automatically. External formulas, weights, model performance, sensor thresholds, and map classifications were developed for other data and places. FloodSense therefore requires local expert and institutional validation and remains a pre-event research prototype rather than an official forecasting, warning, evacuation, or hazard-certification system.",
]
for text in gap_paragraphs:
    clone_paragraph_before(comparison_heading, body_template, text)

# Comparative synthesis prose and table captions.
replace_text(find_paragraph(doc, "The following comparisons summarize the documented approaches and contributions of the foreign and Philippine systems discussed above. Official platforms are identified by their responsible institutions; the methodology column describes the documented technical approach and does not imply that every platform is a peer-reviewed research study."),
             "The following tables compare the verified methodologies, objectives, and contributions of the foreign and Philippine research studies retained in this chapter. Official platforms are discussed separately in the literature and existing-systems discussion because their documentation does not automatically constitute a research methodology and results section.")
replace_text(find_paragraph(doc, "Table 1 compares the related foreign systems by responsible institution, documented approach, objective, and contribution."),
             "Table 1 compares the retained foreign studies by author, verified research or technical procedure, objective, and documented contribution.")
replace_text(find_paragraph(doc, "Table 2 compares the related Philippine systems using the same categories."),
             "Table 2 compares the retained Philippine studies using the same evidence-based categories.")

# Table 1: verified foreign studies.
foreign_rows = [
    ["SYSTEM", "AUTHOR(S)/YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"],
    [
        "Examining Flood Vulnerability Mapping Approaches in Developing Countries: A Scoping Review",
        "Membele, Naidu, and Mutanga (2022)",
        "PRISMA-ScR scoping review of 130 studies selected from Web of Science and Scopus.",
        "The review aimed to identify and compare flood-vulnerability mapping approaches used in developing countries.",
        "It documented major variation in concepts, indicators, scales, participation, sensitivity analysis, and validation practice.",
    ],
    [
        "Actionable Information in Flood Risk Communications and the Potential for New Web-Based Tools for Long-Term Planning for Individuals and Community",
        "Mostafiz, Rohli, Friedland, and Lee (2022)",
        "Literature review of flood-risk communication and comparative review of 50 freely available web portals.",
        "The study aimed to assess how flood information could better support long-term individual and community decisions.",
        "It found limited micro-scale, actionable long-term planning information and identified needs for clearer decision support.",
    ],
]
foreign_continued_rows = [
    ["SYSTEM", "AUTHOR(S)/YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"],
    [
        "Accelerated and Interpretable Flood Susceptibility Mapping Through Explainable Deep Learning with Hydrological Prior Knowledge",
        "Wang, Sanderson, Iqbal, and Woo (2025)",
        "Comparative machine-learning and U-Net modeling with spatial five-fold cross-validation and Grad-CAM interpretation.",
        "The study aimed to improve the speed, predictive performance, and interpretability of flood-susceptibility mapping.",
        "The proposed model reported an AUC of 0.966 and faster convergence, while identifying data and transferability limits.",
    ],
]
resize_rows(doc.tables[0], len(foreign_rows))
for i, row in enumerate(foreign_rows):
    set_row(doc.tables[0], i, row)
resize_rows(doc.tables[1], len(foreign_continued_rows))
for i, row in enumerate(foreign_continued_rows):
    set_row(doc.tables[1], i, row)

# Table 2: verified Philippine studies.
local_rows = [
    ["SYSTEM", "AUTHOR(S)/YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"],
    [
        "GIS-Based Mapping of Flashflood Susceptible Areas in Bukang Liwayway, Kibawe, Bukidnon",
        "Puno et al. (2022)",
        "GIS frequency-ratio analysis using 160 inventory points, nine conditioning factors, and field validation.",
        "The study aimed to generate a localized flashflood-susceptibility map and identify exposed households and a possible evacuation site.",
        "It mapped low, moderate, and high susceptibility and reported 25 households within the high-susceptibility area.",
    ],
    [
        "Flood Risk Assessment Using GIS-Based Analytical Hierarchy Process in the Municipality of Odiongan, Romblon, Philippines",
        "Gacu, Monjardin, Senoro, and Tan (2022)",
        "GIS-based analytical hierarchy process and weighted overlay using expert judgments, ground assessment, and historical-event checks.",
        "The study aimed to map municipal flood risk by combining hazard, vulnerability, and exposure parameters.",
        "It identified nine high-risk barangays and reported the mapped proportions of the municipality's risk classes.",
    ],
]
local_continued_rows = [
    ["SYSTEM", "AUTHOR(S)/YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"],
    [
        "Web-Based Solution for Flood Warning Decision Support in the Province of Leyte, Philippines",
        "Bentoso et al. (2021)",
        "Sensor-based engineering prototype using rainfall and water-level acquisition, web visualization, threshold alerts, and functional testing.",
        "The study aimed to develop a web-based aid for monitoring flood conditions and supporting warning decisions in Leyte.",
        "It reported graphical presentation of sensor readings and warning messages received by responsible authorities.",
    ],
]
resize_rows(doc.tables[2], len(local_rows))
for i, row in enumerate(local_rows):
    set_row(doc.tables[2], i, row)
resize_rows(doc.tables[3], len(local_continued_rows))
for i, row in enumerate(local_continued_rows):
    set_row(doc.tables[3], i, row)

for table in doc.tables[:4]:
    mark_table_rows(table)

# Table 3 remains a system-feature comparison; correct unsupported/misleading legend wording.
table3_intro = find_paragraph(doc, "Table 3 shows that FloodAdapt (S1) and UP NOAH (S6) support exploration of flood scenarios, while GloFAS (S2), IFIS (S3), and the Leyte prototype (S10) provide forecasts or sensor-based information. The England service (S5) and HazardHunterPH (S7) offer public location-based lookup, whereas GeoMapperPH (S8) and PlanSmartPH (S9) focus on institutional data collection and planning. Among the systems compared, FloodSense combines a user-confirmed hypothetical rainfall scenario, a selected or foreground-confirmed location, deterministic rule-based susceptibility assessment, and separate preparedness guidance in a resident-initiated flow. This comparison does not imply that FloodSense provides live monitoring or forecasts, or that its local susceptibility inputs have already been validated.")
replace_text(table3_intro,
             "Table 3 compares documented functions of selected existing platforms and the Leyte research prototype. FloodAdapt (S1) and UP NOAH (S6) support flood-scenario or hazard-map exploration, while GloFAS (S2), IFIS (S3), and the Leyte prototype (S10) provide forecasts or sensor-based information. England's national flood-risk resource (S5) and HazardHunterPH (S7) support public geographic lookup, whereas GeoMapperPH (S8) and PlanSmartPH (S9) support governed institutional data or planning workflows. FloodSense combines a user-confirmed hypothetical rainfall scenario, a selected or foreground-confirmed location, deterministic rule-based susceptibility assessment, and separate preparedness guidance in a resident-initiated flow. The comparison does not imply that FloodSense provides live monitoring or forecasts or that its local susceptibility inputs are already validated.")
replace_text(find_paragraph(doc, "S5 - Check the long term flood"),
             "S5 - National Assessment of Flood and Coastal Erosion Risk in England 2024: An Official Long-Term Risk Information Resource")
set_cell(doc.tables[4].rows[3].cells[9], "")  # PlanSmartPH: interactive geographic display not established by the inspected source.
mark_table_rows(doc.tables[4])

# Chapter 2 references: remove undated entries and add only sources required by the revised chapter.
references_heading = find_paragraph(doc, "REFERENCES")
appendices_heading = find_paragraph(doc, "PROVISIONAL APPENDICES (ORDER AND LETTERS TO BE CONFIRMED)")
reference_template = doc.paragraphs[paragraph_index(doc, "REFERENCES") + 1]
node = references_heading._p.getnext()
while node is not None and node is not appendices_heading._p:
    nxt = node.getnext()
    if node.tag.endswith("}p"):
        node.getparent().remove(node)
    node = nxt

references = [
    "Android Developers. (2026). Request location permissions. https://developer.android.com/develop/sensors-and-location/location/permissions",
    "Bentoso, L. D., Juan, E. O., Brosas, D. G., Paragas, J. R., Nuevas, L. K., & Velarde, M. W. C. (2021). Web-based solution for flood warning decision support in the Province of Leyte, Philippines. In 2021 3rd International Conference on Research and Academic Community Services (ICRACOS) (pp. 185-190). IEEE. https://doi.org/10.1109/ICRACOS53680.2021.9701990",
    "Buchanan, B. G., & Shortliffe, E. H. (Eds.). (1984). Rule-based expert systems: The MYCIN experiments of the Stanford Heuristic Programming Project. Addison-Wesley.",
    "Copernicus Emergency Management Service. (2026). GloFAS homepage. European Centre for Medium-Range Weather Forecasts. https://confluence.ecmwf.int/spaces/CEMS/pages/315561239/3.1+GloFAS+homepage",
    "Deltares. (2025, December 2). FloodAdapt. https://www.deltares.nl/en/software-and-data/products/floodadapt",
    "Department of Science and Technology-Philippine Institute of Volcanology and Seismology. (2022a, December 6). DOST-PHIVOLCS conducts training on the use of GeoRiskPH platforms for the Municipality of General Luna, Quezon. GeoRisk Philippines Initiative. https://www.georisk.gov.ph/articles/2022/12/dost-phivolcs-conducts-training-on-the-use-of-georiskph-platforms-for-the-municipality-of-general-luna-quezon",
    "Department of Science and Technology-Philippine Institute of Volcanology and Seismology. (2022b, September 28). DOST-PHIVOLCS conducts PlanSmart Ready to Rebuild training of trainers. GeoRisk Philippines Initiative. https://www.georisk.gov.ph/articles/2022/9/dost-phivolcs-conducts-plansmart-ready-to-rebuild-training-of-trainers",
    "Environment Agency. (2024). National assessment of flood and coastal erosion risk in England 2024. GOV.UK. https://www.gov.uk/government/publications/national-assessment-of-flood-and-coastal-erosion-risk-in-england-2024/national-assessment-of-flood-and-coastal-erosion-risk-in-england-2024",
    "Fakhruddin, B., Kirsch-Wood, J., Niyogi, D., Guoqing, L., Murray, V., & Frolova, N. (2022). Harnessing risk-informed data for disaster and climate resilience. Progress in Disaster Science, 16, 100254. https://doi.org/10.1016/j.pdisas.2022.100254",
    "Federal Emergency Management Agency. (2025). Resilience Analysis and Planning Tool user guide. https://www.fema.gov/sites/default/files/documents/fema_rapt-user-guide_2025.pdf",
    "Gacu, J. G., Monjardin, C. E. F., Senoro, D. B., & Tan, F. J. (2022). Flood risk assessment using GIS-based analytical hierarchy process in the Municipality of Odiongan, Romblon, Philippines. Applied Sciences, 12(19), 9456. https://doi.org/10.3390/app12199456",
    "GeoRisk Philippines. (2023, August 3). DOST-PHIVOLCS' HazardHunterPH won a 7-Star Rating and the Category Award for Risk Management in the 9th International Best Practice Competition! https://www.georisk.gov.ph/articles/2023/8/dost-phivolcs-hazar-hunter-ph-won-a-7-star-rating-and-the-category-award-for-risk-management-in-the-9th-international-best-practice-competition",
    "International Organization for Standardization. (2023). ISO/IEC 25010:2023 systems and software engineering - Systems and software Quality Requirements and Evaluation (SQuaRE) - Product quality model. https://www.iso.org/standard/78176.html",
    "Iowa Flood Center. (2024). A strong purpose. IIHR Currents 2023-24. University of Iowa. https://iihr.uiowa.edu/news/iihr-currents/iihr-currents-2023-24/strong-purpose",
    "Japan International Cooperation Agency. (2021, October 1). Flood risk protection project in Cavite inaugurated. https://www.jica.go.jp/english/overseas/philippine/information/press/2021/211001.html",
    "Kostopoulos, G., Davrazos, G., & Kotsiantis, S. (2024). Explainable artificial intelligence-based decision support systems: A recent review. Electronics, 13(14), 2842. https://doi.org/10.3390/electronics13142842",
    "Lagmay, A. M. F. A., Racoma, B. A., Aracan, K. A., Alconis-Ayco, J., & Saddi, I. L. (2017). Disseminating near-real-time hazards information and flood maps in the Philippines through Web-GIS. Journal of Environmental Sciences, 59, 13-23. https://doi.org/10.1016/j.jes.2017.03.014",
    "Membele, G. M., Naidu, M., & Mutanga, O. (2022). Examining flood vulnerability mapping approaches in developing countries: A scoping review. International Journal of Disaster Risk Reduction, 69, 102766. https://doi.org/10.1016/j.ijdrr.2021.102766",
    "Mostafiz, R. B., Rohli, R. V., Friedland, C. J., & Lee, Y.-C. (2022). Actionable information in flood risk communications and the potential for new web-based tools for long-term planning for individuals and community. Frontiers in Earth Science, 10, 840250. https://doi.org/10.3389/feart.2022.840250",
    "National Disaster Risk Reduction and Management Council. (2024). National disaster response plan 2024: Hydro-meteorological hazards. https://ndrrmc.gov.ph/attachments/article/4125/National_Disaster_Response_Plan_NDRP_2024.pdf",
    "Papadopoulos, P., Soflano, M., Chaudy, Y., Adejo, W., & Connolly, T. M. (2022). A systematic review of technologies and standards used in the development of rule-based clinical decision support systems. Health and Technology, 12, 713-727. https://doi.org/10.1007/s12553-022-00672-9",
    "Philippine Atmospheric, Geophysical and Astronomical Services Administration. (2024). 2023 annual climate bulletin. https://pubfiles.pagasa.dost.gov.ph/pagasaweb/files/cad/Climate%20Bulletin%202023_v5.pdf",
    "Philippine Statistics Authority. (2023). Third quarter 2023 PSGC updates. https://psa.gov.ph/classification/psgc/node/1684061390",
    "Philippine Statistics Authority. (2025). City of Bacoor: Barangays in the City of Bacoor. https://psa.gov.ph/classification/psgc/barangays/0402103000",
    "Philippine Statistics Authority. (2026). Highlights of the Compendium of Philippine Environment Statistics 2016-2025, Component 4: Extreme events and disasters. https://psa.gov.ph/statistics/environment-statistics/highlights/component-4-extreme-events-and-disaster",
    "PostGIS Development Group. (2026). PostGIS manual: Introduction. https://postgis.net/docs/manual-3.7/en/postgis_introduction.html",
    "Power, D. J. (2002). Decision support systems: Concepts and resources for managers. Quorum Books.",
    "Puno, G. R., Olpenda, A. S., Puno, R. C. C., Aguinsatan, R. G., Lacandula, L. F., Maghuyop, I. V., & Villar, R. G. (2022). GIS-based mapping of flashflood susceptible areas in Bukang Liwayway, Kibawe, Bukidnon. Central Mindanao University Journal of Science, 26(1), 19-29. https://doi.org/10.52751/bwyx7405",
    "Republic Act No. 10173. (2012). Data Privacy Act of 2012. Republic of the Philippines. https://privacy.gov.ph/data-privacy-act/",
    "Tullis, J. A., & Kar, B. (2021). Where is the provenance? Ethical replicability and reproducibility in GIScience and its critical applications. Annals of the American Association of Geographers, 111(5), 1318-1328. https://doi.org/10.1080/24694452.2020.1806029",
    "Wang, H., Sanderson, D., Iqbal, J., & Woo, W. L. (2025). Accelerated and interpretable flood susceptibility mapping through explainable deep learning with hydrological prior knowledge. Remote Sensing, 17(9), 1540. https://doi.org/10.3390/rs17091540",
]
for reference in references:
    clone_paragraph_before(appendices_heading, reference_template, reference)

# Verify protected manuscript areas did not change.
paragraphs = doc.paragraphs
if text_hash([p.text for p in paragraphs[:111]]) != before_intro:
    raise RuntimeError("Protected content before Scope and Limitation changed")
definition_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Definition of Terms")
chapter2_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "REVIEW OF RELATED LITERATURE")
if text_hash([p.text for p in paragraphs[definition_idx:chapter2_idx]]) != before_definitions:
    raise RuntimeError("Protected definitions changed")
method_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "METHODOLOGY")
refs_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "REFERENCES")
if text_hash([p.text for p in paragraphs[method_idx:refs_idx]]) != before_chapter3:
    raise RuntimeError("Protected Chapter 3 changed")
appendix_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "PROVISIONAL APPENDICES (ORDER AND LETTERS TO BE CONFIRMED)")
if text_hash([p.text for p in paragraphs[appendix_idx:]]) != before_appendices:
    raise RuntimeError("Protected appendices changed")

# Final content guards.
chapter2_text = "\n".join(p.text for p in paragraphs[chapter2_idx:refs_idx])
if "n.d." in chapter2_text:
    raise RuntimeError("An n.d. entry remains in Chapter 2")
for forbidden in ("according to our TA consultation", "the TA instructed", "after the developers' handoff", "this revised manuscript", "the previous version"):
    if forbidden.lower() in "\n".join(p.text for p in paragraphs).lower():
        raise RuntimeError(f"Revision-history phrase remains: {forbidden}")

doc.save(DOCX)
print(f"Saved in place: {DOCX}")
print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}")
print("Protected section hashes verified before save.")
