from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Revised_Manuscript_Chapters_1_to_3.docx")
FONT = "Arial"
BODY_SIZE = 11


def set_run_font(run, size=BODY_SIZE, bold=False, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), FONT)
    rpr.rFonts.set(qn("w:hAnsi"), FONT)
    rpr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_page(section):
    section.page_width = Cm(21.03)
    section.page_height = Cm(29.70)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run)


def set_page_start(section, number):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        section._sectPr.append(node)
    node.set(qn("w:start"), str(number))


def set_paragraph_border(paragraph, top=False, bottom=False):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    for edge, enabled in (("top", top), ("bottom", bottom)):
        if enabled:
            item = OxmlElement(f"w:{edge}")
            item.set(qn("w:val"), "single")
            item.set(qn("w:sz"), "12")
            item.set(qn("w:space"), "3")
            item.set(qn("w:color"), "000000")
            borders.append(item)


def body(doc, text="", first_indent=True, italic=False, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_indent:
        pf.first_line_indent = Inches(0.5)
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)
    return p


def body_parts(doc, parts, first_indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if first_indent:
        pf.first_line_indent = Inches(0.5)
    for text, bold, italic in parts:
        run = p.add_run(text)
        set_run_font(run, bold=bold, italic=italic)
    return p


def chapter_heading(doc, text, page_break=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(25.3)
    pf.keep_with_next = True
    pf.page_break_before = page_break
    run = p.add_run(text.upper())
    set_run_font(run, bold=True)
    return p


def section_heading(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(25.3)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, bold=True)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, bold=True)
    return p


def numbered_item(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent = Inches(0.75)
    pf.first_line_indent = Inches(-0.25)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(f"{number}.  {text}")
    set_run_font(run)
    return p


def lead_paragraph(doc, lead, text):
    return body_parts(doc, [(lead, True, False), (text, False, False)])


def placeholder(doc, text, space_before=18, space_after=18):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.keep_together = True
    run = p.add_run(f"[{text}]")
    set_run_font(run, italic=True)
    return p


def reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(12.65)
    run = p.add_run(text)
    set_run_font(run)
    return p


doc = Document()
configure_page(doc.sections[0])

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(BODY_SIZE)
normal.font.color.rgb = RGBColor(0, 0, 0)

# SURECUT-style title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_after = Pt(135)
r = p.add_run("FLOODSENSE: AN ANDROID-BASED EXPERT SYSTEM FOR FLOOD\nSUSCEPTIBILITY ASSESSMENT AND PRE-EVENT PREPAREDNESS IN\nBACOOR CITY")
set_run_font(r, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_after = Pt(132)
r = p.add_run("Undergraduate Thesis\nSubmitted to the Faculty of the\nDepartment of Computer Studies\nCavite State University - Bacoor City Campus\nCity of Bacoor, Cavite")
set_run_font(r)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_after = Pt(125)
r = p.add_run("In partial fulfillment\nof the requirements for the degree\nBachelor of Science in Computer Science")
set_run_font(r)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
r = p.add_run("REYMART V. GOC-ONG\nMARTIN LORENZ D. JUANITES\nROMAR T. PULAO")
set_run_font(r, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
r = p.add_run("May 2027")
set_run_font(r)

# Main manuscript section; page 1 is counted but its number is hidden, as in SURECUT.
main = doc.add_section(WD_SECTION.NEW_PAGE)
configure_page(main)
main.header.is_linked_to_previous = False
main.first_page_header.is_linked_to_previous = False
main.different_first_page_header_footer = True
add_page_number(main.header.paragraphs[0])
main.first_page_header.paragraphs[0].clear()
set_page_start(main, 1)

# Repeated manuscript-identification block on the first numbered page.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_after = Pt(28)
r = p.add_run("FloodSense: An Android-Based Expert System for Flood Susceptibility\nAssessment and Pre-Event Preparedness in Bacoor City")
set_run_font(r, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_after = Pt(28)
r = p.add_run("Reymart V. Goc-ong\nMartin Lorenz D. Juanites\nRomar T. Pulao")
set_run_font(r, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(24)
set_paragraph_border(p, top=True, bottom=True)
r = p.add_run(
    "An undergraduate thesis manuscript submitted to the faculty of the Department of Computer Studies, "
    "Cavite State University - Bacoor City Campus, City of Bacoor, Cavite, in partial fulfillment of the "
    "requirements for the degree of Bachelor of Science in Computer Science with Contribution No. "
    "__________________. Prepared under the supervision of ________________________________."
)
set_run_font(r)

# CHAPTER 1
chapter_heading(doc, "INTRODUCTION")
intro_paragraphs = [
    "Flooding is one of the most frequent and destructive natural hazards affecting Philippine communities. Intense and prolonged rainfall can interrupt transportation, damage homes and livelihoods, displace families, and create uncertainty about when protective action is necessary. Bacoor City, Cavite, includes low-lying and rapidly urbanizing areas connected to the Imus and Bacoor river systems. Flood-mitigation facilities in the Imus River Basin were established to reduce peak floodwater volume in Cavite, including low portions of Bacoor and nearby areas, confirming that flood exposure remains a continuing local concern (Japan International Cooperation Agency [JICA], 2021). These conditions establish the need for accessible, understandable, and location-specific information that residents can consult before a flood event develops.",
    "To examine the information problem, the researchers conducted an initial requirements survey among 50 Bacoor City residents from 12 barangays. Ninety-two percent had experienced flooding, 58% described previous flooding as severe or waist-deep and higher, only 20% considered themselves well informed about flood-prone areas, and 64% reported making an unsafe or incorrect decision because sufficient area-specific information was unavailable (FloodSense Research Team, 2026a). The findings indicate that receiving a general warning does not always help a resident understand the susceptibility of a particular location or identify the preparations appropriate to that situation.",
    "A supplementary requirements survey involving 40 respondents from 14 barangays produced a comparable pattern. Only 22.5% considered themselves well informed, 65% were dissatisfied or very dissatisfied with the flood information available to them, 72.5% reported an unsafe or incorrect decision, and 57.5% experienced difficulty determining the risk of a particular barangay. Respondents also expressed strong interest in plain-language results, severity-based action checklists, and distance-ranked evacuation-center information (FloodSense Research Team, 2026b). The two surveys were administered to different respondent groups and are therefore reported as separate requirements datasets rather than combined into one sample.",
    "Interviews with the Bacoor Disaster Risk Reduction and Management Office further showed that historical records, hazard information, preparedness guidance, and evacuation-center information require institutional custody and validation. The office explained that susceptibility can vary within a barangay and that an application should not assign one conclusion to every place inside the boundary when more specific evidence is available. Coordinates, sitios, zones, affected areas, verified boundaries, source dates, and operational limitations must therefore be preserved when supported data are used (FloodSense Research Team, 2026c; 2026d).",
    "The City Planning and Development Coordinator Office identified geographic materials that may support the study, including elevation, land-use, flood-hazard, and barangay-boundary information (FloodSense Research Team, 2026e). These materials cannot automatically be treated as operational FloodSense inputs. Each dataset must be reviewed according to its source, date, geographic resolution, processing history, limitations, permission, and suitability for the research method. HazardHunterPH likewise explains that scenario-based hazard maps are references that should be considered together with current observations, field assessment, and other hazard information (GeoRisk Philippines, 2026).",
    "To address the identified gap, the researchers developed FloodSense, an Android-based Rule-Based Expert System for flood susceptibility assessment and pre-event preparedness in Bacoor City. A resident selects or confirms a hypothetical rainfall intensity, rainfall duration, and supported location. The backend applies a fixed and documented inference procedure to eligible, versioned knowledge and returns a susceptibility state with an explanation. When the available knowledge is incomplete, conflicting, provisional, or outside the supported area, the system returns a limitation state instead of fabricating a classification.",
    "FloodSense combines a guided mobile flow, an interactive geographic map, explainable susceptibility results, separately sourced preparedness guidance, and verified resource information. The application explains its scenario-based purpose, limitations, privacy practices, and location use before an assessment is performed. The Expert System determines the classification, while the Decision Support System presents approved pre-event guidance without changing that classification or issuing an evacuation order. A role-controlled Web Administration System maintains authorized geographic records, approved scenario parameters, source metadata, guidance, evacuation-center records, and publication status. Ordinary administrators cannot create or modify the raw inference algorithm or expert rules; such changes require research ownership, expert validation, versioning, testing, and explicit authorization (FloodSense Research Team, 2026f).",
    "The current study therefore positions FloodSense as a complementary preparedness application rather than a live monitoring, forecasting, or emergency-dispatch platform. It does not continuously poll rainfall stations, run a background timer, issue autonomous alerts, determine road passability, report current evacuation-center occupancy, or replace PAGASA and BDRRMO advisories. This boundary prevents unvalidated water-level, rainfall-station, and geographic relationships from being presented as established scientific conclusions.",
]
for paragraph in intro_paragraphs:
    body(doc, paragraph)

section_heading(doc, "Statement of the Problem")
body(doc, "The study aimed to develop an Android-based Expert System that provides explainable and location-specific flood susceptibility assessment and pre-event preparedness support for residents of Bacoor City. The research problems were derived from two existing requirements surveys, interviews with the BDRRMO and CPDCO, the thesis adviser consultation, current system records, and recent related literature.")
body_parts(doc, [
    ("Residents lack a consistent way to translate rainfall intensity, duration, geographic attributes, and historical information into an understandable susceptibility state for a supported location. The first survey found that only 20% of respondents considered themselves well informed and that 64% had made an unsafe or incorrect decision because sufficient information was unavailable. The supplementary survey likewise found that only 22.5% felt well informed, while 57.5% had difficulty determining the risk of a particular barangay. ", False, False),
    ("How can an Android-based Expert System use validated rainfall-scenario and geographic facts to provide residents with an understandable and explainable flood susceptibility assessment for a supported location?", False, True),
])
body_parts(doc, [
    ("The BDRRMO and CPDCO hold or identify institutional knowledge, historical records, geographic materials, preparedness guidance, and evacuation-center information, but these sources require authorization, validation, updating, and careful interpretation before public use. The BDRRMO also explained that susceptibility can vary within a barangay and that evacuation-center records may become outdated. ", False, False),
    ("How can FloodSense organize source-backed institutional knowledge, geographic information, preparedness content, and verified resource records while preserving provenance, review status, geographic limitations, and the authority of the responsible offices?", False, True),
])
body_parts(doc, [
    ("Residents currently depend on fragmented information channels and may need to perform several actions before reaching guidance relevant to their situation. The thesis adviser requested a more automatic and guided flow, location support, map response to the selected scenario, preparedness guidance, and first-launch privacy information. At the same time, the adviser restricted ordinary administrators from changing expert rules, and the research team retained a scenario-based boundary instead of background monitoring. ", False, False),
    ("How can the mobile and web components provide a guided, low-friction, and privacy-aware assessment process while keeping the inference method fixed, protecting expert knowledge from unauthorized editing, and avoiding unsupported real-time claims?", False, True),
])
body_parts(doc, [
    ("Given these concerns, the study seeks to answer the overarching question: ", False, False),
    ("How can FloodSense integrate validated rule-based inference, hypothetical rainfall scenarios, supported geographic selection, explainable map visualization, protected data governance, preparedness guidance, and verified resource information to strengthen flood susceptibility awareness and pre-event preparedness in Bacoor City without replacing official forecasts, warnings, or emergency decisions?", False, True),
])

section_heading(doc, "Objectives of the Study")
body(doc, "Generally, the study aimed to develop FloodSense, an Android-based Rule-Based Expert System for explainable flood susceptibility assessment and pre-event preparedness in Bacoor City.")
body(doc, "Specifically, it aimed to:")
objectives = [
    "identify the flood-information, location-awareness, preparedness, governance, and usability needs of Bacoor City residents and relevant local government offices using the available requirements surveys, interviews, consultation records, and the ongoing validated data-gathering process;",
    "analyze authorized flood-related knowledge, geographic records, rainfall-scenario references, historical information, preparedness guidance, and evacuation-center data to define supported facts, limitations, provenance requirements, and the fixed inference procedure;",
    "develop an Android application and supporting web services that provide first-launch privacy and limitation information, a guided rainfall-scenario and location flow, explainable susceptibility results, an interactive map, sourced preparedness guidance, and verified resource details;",
    "develop a role-controlled Web Administration System for maintaining authorized reference data, approved parameters, sources, Decision Support System guidance, geographic records, and evacuation-center information without allowing ordinary administrators to edit the raw expert rules or inference algorithm; and",
    "test and evaluate the functional suitability, performance efficiency, compatibility, interaction capability, reliability, security, maintainability, flexibility, and safety of FloodSense using appropriate software tests and a validated user-evaluation instrument based on ISO/IEC 25010:2023.",
]
for index, item in enumerate(objectives, 1):
    numbered_item(doc, index, item)

section_heading(doc, "Theoretical Framework")
body(doc, "The theoretical framework of FloodSense explains how approved research and institutional inputs are transformed into an explainable scenario-based result. The framework separates the source and governance layer, fixed Expert System inference, spatial processing, resident interaction, Decision Support System guidance, and verified resource presentation. This separation prevents an ordinary content update from silently changing the susceptibility conclusion and makes each public output traceable to its controlling source and version.")
placeholder(doc, "Insert the adviser-approved FloodSense theoretical framework as Figure 1")
lead_paragraph(doc, "Authentication, Onboarding, and Privacy Module. ", "This module presents the application's purpose, scenario-based nature, limitations, privacy information, and location-use explanation before precise location is requested. Location access, if included in the final approved design, is user-triggered and foreground-only. Manual map or area selection remains the fallback.")
lead_paragraph(doc, "Scenario and Location Input Module. ", "This module allows the resident to choose or confirm hypothetical rainfall intensity and duration and to select a supported location. The values are planning inputs and do not describe current weather. A temporary pin may be resolved against approved geographic polygons, while unsupported points and incomplete scenarios produce limitation responses.")
lead_paragraph(doc, "Protected Knowledge and Inference Module. ", "This module contains the fixed forward-chaining procedure and eligible, versioned research knowledge. It evaluates facts for the selected scenario and location and returns Low, Moderate, High, Very High, Uncertain, Insufficient Data, or Outside Supported Area as appropriate. Raw expert rules, priority handling, and algorithm behavior are not operational-administrator controls.")
lead_paragraph(doc, "Spatial Processing and Dynamic Map Module. ", "GeoDjango and PostGIS manage supported polygons, points, containment queries, and geographic relationships. The Android map displays the result using both color and text. The basemap provides geographic context only; it is not the source of susceptibility knowledge. A color change is produced by the evaluated scenario and published data rather than by an administrator manually assigning a conclusion.")
lead_paragraph(doc, "Explainable Result and Decision Support Module. ", "The result screen presents the supported area, scenario, susceptibility state, matched basis, source version, limitations, and warnings. A separate Decision Support System presents validated pre-event guidance. Guidance can be maintained by authorized personnel when it is sourced, reviewed, and published, but guidance edits cannot alter the susceptibility class.")
lead_paragraph(doc, "Administration, Provenance, and Resource Module. ", "The management portal maintains approved parameters, geographic records, sources, guidance, and verified evacuation-center records through draft, review, approval, activation, and rollback states. Technical rule changes follow a separate research and expert-validation process. Every public record must identify its source, effective date, status, and limitations.")
body(doc, "Overall, the framework shows how FloodSense combines protected rule-based reasoning, geographic processing, explainable results, sourced preparedness guidance, and controlled publishing in one mobile and web environment. The system supports resident understanding before a flood event while preserving the authority of PAGASA, the BDRRMO, barangay officials, and emergency responders.")

section_heading(doc, "Significance of the Study")
lead_paragraph(doc, "For Bacoor City Residents. ", "FloodSense provides an additional way to explore the susceptibility of a supported location under a hypothetical rainfall scenario, understand why a result was produced, review pre-event preparedness guidance, and identify verified resource information without presenting itself as an official warning or guarantee of safety.")
lead_paragraph(doc, "For the Bacoor Disaster Risk Reduction and Management Office. ", "The study demonstrates a governed channel for publishing approved preparedness content, source information, and verified resource records while preserving the office's authority over official advisories and emergency directions. It also documents why arbitrary administrator editing of expert rules is inappropriate.")
lead_paragraph(doc, "For the City Planning and Development Coordinator Office. ", "The study illustrates how verified geographic and planning information may support community-oriented spatial assessment when source, scale, processing, and limitations are preserved.")
lead_paragraph(doc, "For the Department of Computer Studies. ", "The project provides a documented case of combining rule-based reasoning, spatial databases, mobile development, administrative governance, explainability, and user-centered requirements within an undergraduate computer science study.")
lead_paragraph(doc, "For Future Researchers. ", "The study provides a reproducible boundary for scenario-based flood susceptibility tools and identifies unresolved areas that require further investigation, including water-level integration, automated-rainfall-gauge data, expert validation, spatial granularity, and larger-scale evaluation.")

section_heading(doc, "Time and Place of the Study")
body(doc, "The formal research and development period extends from September 2026 to May 2027. Requirements gathering began earlier in May 2026 through resident surveys and key informant interviews in Bacoor City. The BDRRMO and CPDCO interviews were conducted at the Bacoor Government Center, while software development, academic consultation, integration, testing, and manuscript preparation were conducted primarily through the Department of Computer Studies at Cavite State University - Bacoor City Campus.")
body(doc, "Development proceeded iteratively as requirements, interface designs, data structures, rule behavior, spatial processing, mobile screens, and administrative functions were reviewed and refined. The current evaluation survey is still being conducted. Its final dates, participant counts, response rates, and findings will be inserted only after the responses have been completed, validated, and analyzed.")

section_heading(doc, "Scope and Limitation of the Study")
scope_paragraphs = [
    "The study covers an Android application for scenario-based flood susceptibility assessment and pre-event preparedness in Bacoor City. A resident selects or confirms a hypothetical rainfall intensity, rainfall duration, and supported location. The fixed Expert System evaluates eligible facts and returns an explainable susceptibility or limitation state. A separate Decision Support System then displays sourced preparedness guidance.",
    "The geographic scope is limited to Bacoor City and to locations for which the researchers possess usable, authorized, sufficiently detailed, and validated records. A barangay boundary identifies an administrative area but does not, by itself, establish one susceptibility conclusion for every place inside it. Where available data do not support finer spatial differences, the application must disclose that limitation rather than imply precision.",
    "The mobile flow includes first-launch limitations and privacy information, rainfall-scenario selection, supported location selection or temporary pin placement, assessment confirmation, map and result presentation, explanation, preparedness guidance, and verified resource details. Device GPS remains optional pending final adviser approval and cannot be implemented as continuous background tracking. Manual selection remains available.",
    "The map may display supported polygons using green, yellow, orange, and red with accompanying Low, Moderate, High, and Very High text. It does not display live flood extent, current water depth, road closures, traffic, river level, satellite rainfall, or evacuation-center occupancy. OpenStreetMap, when used, supplies geographic context rather than susceptibility knowledge.",
    "Evacuation-center information is limited to records obtained from an authorized custodian and reviewed for currency, coordinates, capacity or facilities when available, verification date, and operational limitations. Distance from a selected point does not prove route safety, accessibility, available capacity, or current opening status.",
    "The custom Web Administration System permits role-controlled management of approved datasets, scenario references, geographic records, sources, Decision Support System guidance, evacuation-center records, publication states, and audits. It does not permit ordinary administrators to create or edit raw expert rules, conflict-resolution logic, or the inference algorithm.",
    "Water-level and automated-rainfall-gauge datasets are outside the active inference method until their fields, units, temporal resolution, spatial coverage, missing values, datum, license, and methodological contribution have been examined and validated. CSV or Excel import and export are likewise outside the committed implementation until the required direction, schema, validation, approval, and rollback process are confirmed.",
    "FloodSense requires an internet connection for current server-managed assessment, map, guidance, and resource records. The present scope excludes a fully offline version, iOS release, live meteorological integration, machine-learning retraining, autonomous alerts, continuous timers, background location services, emergency dispatch, and official forecasting.",
]
for paragraph in scope_paragraphs:
    body(doc, paragraph)

section_heading(doc, "Definition of Terms")
terms = [
    ("Android Application. ", "The resident-facing FloodSense software that provides the guided scenario, supported location selection, assessment result, interactive map, preparedness guidance, and verified resource information."),
    ("Decision Support System. ", "The component that presents sourced pre-event preparedness guidance using the latest valid susceptibility result and approved situational inputs without changing the classification or making the resident's final decision."),
    ("Expert System. ", "A knowledge-based software component that applies a fixed inference procedure to eligible facts and validated rules to derive an explainable result."),
    ("Flood Susceptibility. ", "The relative potential of a supported location to experience rainfall-related flooding under a defined scenario. In FloodSense, it is not a prediction of the exact time, depth, or extent of an actual flood."),
    ("Forward Chaining. ", "The inference approach that begins with available scenario and geographic facts and evaluates applicable rules until a conclusion or limitation state is reached."),
    ("Geographic Information System. ", "A system for storing, processing, querying, and presenting information connected to geographic locations."),
    ("Knowledge Base. ", "The controlled collection of validated facts, rules, conditions, priorities, explanations, versions, and sources used by the Expert System."),
    ("Location Resolution. ", "The process of determining which approved geographic feature contains or corresponds to a temporary point selected by the resident."),
    ("Pre-event Preparedness. ", "Actions completed before a flood emergency, such as preparing essential supplies, planning family communication, protecting important documents, and monitoring official advisories."),
    ("Provenance. ", "Information that records where a dataset, parameter, rule, guidance item, geographic feature, or resource record came from and how it was reviewed, changed, and approved."),
    ("Scenario-based Assessment. ", "An assessment initiated from hypothetical rainfall and location inputs chosen or confirmed by the resident while the application is open."),
    ("Susceptibility State. ", "The system output of Low, Moderate, High, Very High, Uncertain, Insufficient Data, or Outside Supported Area."),
    ("Verified Resource. ", "An evacuation-center or preparedness record obtained from an authorized custodian and reviewed for source, date, status, coordinates, and stated limitations."),
    ("Web Administration System. ", "The role-controlled browser interface used to maintain approved data, sources, guidance, geographic records, resource records, and publication states without exposing raw rule editing to ordinary administrators."),
]
for lead, text in terms:
    lead_paragraph(doc, lead, text)

# CHAPTER 2
chapter_heading(doc, "REVIEW OF RELATED LITERATURE", page_break=True)
body(doc, "This chapter presents a review of literature, articles, and studies that provided the guidelines for defining the requirements and boundaries of FloodSense. The discussion examines flood susceptibility, geographic information, public risk communication, explainable decision support, mobile interaction, institutional data governance, and iterative software development. It also distinguishes scenario-based susceptibility assessment from operational forecasting and identifies the research gap addressed by an explainable Android application for Bacoor City.")

section_heading(doc, "Related Foreign Literature")
subheading(doc, "Flood Susceptibility and Spatial Factors")
body(doc, "Flood susceptibility describes the relative tendency of an area to experience flooding based on relevant physical, hydrological, environmental, and human-related factors. It differs from a real-time forecast because it characterizes potential under specified conditions rather than predicting the exact occurrence of an event. Debnath et al. (2024) demonstrated that geospatial and expert-informed models can integrate multiple flood-conditioning factors and classify areas into ordered susceptibility zones. Their work also identified data scarcity, spatial resolution, validation, and the availability of long-term water-level and discharge records as continuing limitations.")
body(doc, "Nguyen et al. (2024) combined a Geographic Information System and an analytical hierarchy process to evaluate urban flood susceptibility using nine factors, expert consultation, and historical flood points. The resulting map classified areas into five levels and was validated using receiver operating characteristic analysis. The study supports explicit factor definitions, expert review, spatial normalization, and independent validation. FloodSense adopts the need for traceability and validation but does not copy the method or weights because Bacoor-specific parameters have not yet been approved.")

subheading(doc, "Interactive Mapping and Risk Communication")
body(doc, "Spatial information is most useful to the public when map symbols, legends, labels, and explanations are understandable. Oubennaceur et al. (2021) found that interactive story maps can combine hazard maps, explanatory text, and user-oriented interaction to communicate flood risk to non-specialists. The authors emphasized clear legends, simple categories, careful color selection, and limited technical terminology. These principles support FloodSense's use of text-supported color classes, result explanations, and a guided interface instead of requiring residents to interpret raw technical layers.")
body(doc, "Risk communication should also lead users toward appropriate protective understanding. Lin (2023) reported relationships among flood-risk information seeking, mobile-application use, coping beliefs, collective efficacy, and community action. This suggests that a digital platform should not stop at displaying a classification. It should connect the result to comprehensible and source-backed preparedness information while avoiding statements that exceed the authority or evidence of the system.")
body(doc, "Inclusive communication requires more than transmitting a technical message. The United Nations Office for Disaster Risk Reduction (UNDRR, 2025) described how mobile technology can expand access to early-warning communication for groups that may otherwise be missed. Although FloodSense is not an official warning service, the evidence supports readable mobile presentation, clear limitations, and alternative input methods that do not make one device permission the only route to preparedness information.")

subheading(doc, "Decision Support for Flood Preparedness")
body(doc, "Alabbad et al. (2022) developed a web-based flood-mitigation decision-support framework that combined maps, property characteristics, scenarios, guidance, and decision-tree analysis. The study illustrates how a Decision Support System can organize alternatives and present context-specific information instead of making an unexplained decision for the user. FloodSense applies this principle to pre-event household preparedness and keeps the Decision Support System logically separate from susceptibility classification.")
body(doc, "Mattos et al. (2022) connected flood modeling, forecast inputs, map visualization, and a web application in an operational flood-alert context. Their architecture demonstrates the value of integrating data, models, and public communication, but it also shows the extensive infrastructure required for actual forecasting. Because FloodSense does not yet have validated continuous hydrometeorological feeds or a calibrated forecasting model, the study supports retaining scenario-based assessment rather than imitating a real-time warning service.")

subheading(doc, "Explainable Rule-Based Decision Support")
body(doc, "Kostopoulos et al. (2024) reviewed explainable decision-support approaches and identified visual, rule-based, case-based, natural-language, and knowledge-based explanations. Production rules and IF-THEN explanations are particularly relevant when a system must show a traceable decision path. The review also notes that domain knowledge and human-understandable explanations are central to user trust. FloodSense consequently exposes the facts, matched basis, version, and limitations of a result without exposing rule editing to ordinary administrators.")
body(doc, "Galanti et al. (2023) likewise showed that explanations must be intelligible to the people who act on a decision-support output. Although their predictive-process context differs from flood susceptibility, their user evaluation reinforces an important design requirement: accuracy measures alone do not establish that an explanation is understandable. FloodSense must therefore test the wording, sequence, and visual presentation of explanations with its intended users.")

subheading(doc, "Human-Centered Mobile Flood Applications")
body(doc, "Alsabhan and Dudin (2023) examined human-computer-interaction concerns in mobile flood forecasting and warning applications. Their work emphasized the difficulty of presenting complex hydrological information on mobile devices and the need to consider accessibility, interaction, and the conditions under which users make decisions. FloodSense addresses the interface problem through short guided steps, visible limitations, plain-language scenario labels, and consistent result presentation while avoiding the real-time claims of the systems reviewed in that study.")

subheading(doc, "Iterative Requirements and Software Quality")
body(doc, "Hoy and Xu (2023) found that agile requirements engineering can respond to changing client needs through short cycles, user-oriented requirements, use cases, test cases, and concise acceptance criteria, while still requiring disciplined communication and documentation. This supports the iterative process reflected in FloodSense's consultation records, interface designs, implementation guides, and tests. The project therefore uses an Agile iterative software-development lifecycle, while Design Thinking activities are limited to early problem discovery and interface exploration.")
body(doc, "ISO/IEC 25010:2023 defines a product-quality model for specifying, measuring, and evaluating information and communication technology products. Its characteristics provide a current basis for evaluating FloodSense beyond whether individual features run. Functional suitability, performance efficiency, compatibility, interaction capability, reliability, security, maintainability, flexibility, and safety will be considered when the final evaluation instrument and participant plan are approved (International Organization for Standardization [ISO], 2023).")

section_heading(doc, "Related Local Literature")
subheading(doc, "Flood Conditions and Institutional Responsibility in Cavite")
body(doc, "JICA (2021) reported the inauguration of flood-mitigation facilities in the Imus River Basin, including structures intended to reduce flooding in low-lying portions of Bacoor and Imus. The project establishes the continuing relevance of basin-level flood management to Bacoor. It also shows that a resident-facing application must coexist with engineering interventions and government operations rather than imply that information technology alone controls the hazard.")

subheading(doc, "Official Flood Information and Warning Boundaries")
body(doc, "The Philippine Atmospheric, Geophysical and Astronomical Services Administration described flood bulletins as products of near-real-time monitoring, forecast rainfall, water-level trends, possible impacts, and warning protocols coordinated with local disaster offices (PAGASA, 2024). FloodSense does not possess those operational inputs or that institutional authority. Its interface must therefore label rainfall values as hypothetical planning inputs and direct residents to official channels during actual events.")

subheading(doc, "Scenario-Based Hazard Information")
body(doc, "GeoRisk Philippines (2026) presents HazardHunterPH as a platform for examining hazard information while cautioning that map outputs are references rather than absolute predictions. The platform's disclaimer stresses the need to combine maps with real-time data and field assessment. This boundary supports FloodSense's limitation states, source labels, and refusal to treat a barangay boundary or basemap as sufficient proof of susceptibility.")

subheading(doc, "Local Knowledge and Preparedness Requirements")
body(doc, "The BDRRMO interviews documented the need to preserve authoritative data, avoid barangay-wide generalization when conditions vary inside a boundary, validate evacuation-center records, and maintain preparedness guidance from authorized sources (FloodSense Research Team, 2026c; 2026d). The CPDCO interview identified available geographic materials but did not authorize the researchers to treat every identified layer as a validated system input (FloodSense Research Team, 2026e). These records establish the governance and data-validation requirements that distinguish FloodSense from a generic map application.")

section_heading(doc, "Related Foreign Studies")
body(doc, "Recent foreign studies demonstrate several possible flood-information architectures. Oubennaceur et al. (2021) focused on interactive communication of mapped flood risk; Alabbad et al. (2022) developed a decision-support framework for mitigation alternatives; Mattos et al. (2022) connected models and a web application for flood alerts; Noori and Bonakdari (2023) used expert-informed Geographic Information System modeling; Alsabhan and Dudin (2023) examined a human-centered mobile warning application; Debnath et al. (2024) compared geospatial and expert-informed susceptibility models; and Nguyen et al. (2024) developed and validated an urban susceptibility map. Together, these studies show that public-facing flood systems require defensible data, a clearly defined method, validation, understandable maps, and communication appropriate to the authority of the system.")
body(doc, "The studies also clarify what FloodSense does not yet possess. Operational alert applications depend on current sensor or forecast data and calibrated hydrological or hydraulic models. Susceptibility studies depend on selected factors, weights, spatial layers, and validation observations. Because those components cannot be transferred to Bacoor without expert and methodological approval, FloodSense uses only approved local knowledge and returns Insufficient Data when a reliable conclusion cannot be supported.")

section_heading(doc, "Related Local Studies")
body(doc, "Johnson et al. (2021) modeled urban expansion and flood exposure across the Philippines using open geospatial data. Their findings indicate that future urban growth may increase the number of people and developed land exposed to flooding, supporting the need for spatially explicit preparedness information. However, their national-scale exposure modeling cannot be treated as a location-level Bacoor assessment without local validation.")
body(doc, "Nagumo et al. (2022) created a three-dimensional flood-hazard map for a flood-prone Philippine area and emphasized that visualization can help residents understand local flood characteristics where geographic information and disaster records are limited. The study supports accessible map presentation but also demonstrates that visualization must be linked to a documented model and appropriate local data.")
body(doc, "These Philippine studies demonstrate the value of structured spatial data and accessible visualization, but neither provides a validated set of Bacoor-specific thresholds for FloodSense. Their methods and findings are therefore treated as related evidence rather than as formulas or weights that an administrator may import or invent.")

section_heading(doc, "Synthesis")
body(doc, "The reviewed literature establishes five principles relevant to FloodSense. First, susceptibility is spatial and depends on clearly defined factors and geographic resolution. Second, interactive maps require simple legends, explanations, and contextual limitations. Third, a Decision Support System should organize evidence and guidance while leaving the final decision to people and authorized institutions. Fourth, rule-based explanations can improve transparency, but rules and methods must remain governed and validated. Fifth, operational forecasting requires data and models beyond those presently approved for FloodSense.")
body(doc, "Existing systems commonly emphasize either technical mapping, real-time monitoring, public warning, or general risk communication. FloodSense addresses a narrower local gap: it connects a user-confirmed rainfall scenario and supported Bacoor location to a deterministic and explainable susceptibility result and then to separately sourced pre-event preparedness guidance. Its contribution is not the invention of a new forecast model. It is the governed integration of knowledge, scenario interaction, spatial resolution, explanation, guidance, verified resources, and explicit limitation states in an Android and web architecture.")
body(doc, "The local requirements evidence reinforces this gap. Residents reported limited area-specific awareness, dissatisfaction with available information, unsafe decisions, difficulty determining location-specific risk, and strong interest in actionable guidance and evacuation-center information. The BDRRMO and CPDCO records identify both potential sources and necessary restrictions. These findings justify a system that prioritizes clarity and traceability while refusing unsupported claims of precision or real-time authority.")

p = doc.add_paragraph()
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_before = Pt(12.65)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Table 1. Comparative Synthesis of Selected Related Studies and Systems")
set_run_font(r, italic=True)

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
table.autofit = False
widths = [Inches(1.25), Inches(1.15), Inches(1.55), Inches(1.82)]
headers = ["SYSTEM", "AUTHOR(S)", "METHODOLOGY", "CONTRIBUTION AND REMAINING GAP"]
for index, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
    cell.width = widths[index]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=True)

comparison_rows = [
    ("Flood Risk StoryMaps", "Oubennaceur et al. (2021)", "Interactive flood-risk communication", "Supports plain-language spatial communication; it does not provide Bacoor-specific inference."),
    ("Flood Mitigation DSS", "Alabbad et al. (2022)", "Web DSS and decision tree", "Supports structured mitigation guidance; it differs from household pre-event susceptibility assessment."),
    ("Flood Alert Web Application", "Mattos et al. (2022)", "Model-based flood alert architecture", "Shows real-time infrastructure requirements that FloodSense currently excludes."),
    ("Mobile Flood Application", "Alsabhan and Dudin (2023)", "Human-centered mobile application", "Supports guided mobile interaction but operates in a forecasting context."),
    ("Flood Susceptibility Model", "Debnath et al. (2024)", "GIS and expert-informed models", "Supports multiple spatial factors and validation; parameters cannot be transferred without local review."),
    ("Urban Susceptibility Map", "Nguyen et al. (2024)", "GIS and analytical hierarchy process", "Supports expert weighting and validation; the method is not approved for FloodSense."),
    ("HazardHunterPH", "GeoRisk Philippines (2026)", "National multi-hazard reference platform", "Provides hazard-reference access and expressly warns that maps are not absolute predictions."),
    ("FloodSense", "Present Study", "Scenario-based rule inference, map, DSS, and governance", "Targets explainable Bacoor assessment; it still requires validated local data and final user evaluation."),
]
for row in comparison_rows:
    cells = table.add_row().cells
    for index, (cell, text) in enumerate(zip(cells, row)):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_run_font(run, size=9)

tr_pr = table.rows[0]._tr.get_or_add_trPr()
tbl_header = OxmlElement("w:tblHeader")
tbl_header.set(qn("w:val"), "true")
tr_pr.append(tbl_header)

# CHAPTER 3
chapter_heading(doc, "METHODOLOGY", page_break=True)
body(doc, "This chapter presents the requirements for developing FloodSense and discusses the materials used by the researchers, the selected development method, the participants of the study, and the statistical techniques planned for analyzing the gathered data. Completed requirements-gathering activities are distinguished from the new survey and final system evaluation that are still in progress.")

section_heading(doc, "Materials")
body(doc, "FloodSense is implemented through a Flutter and Dart Android application, a Python Django and Django REST Framework backend, GeoDjango spatial services, and a PostgreSQL database with PostGIS. Flutter's documented application-architecture guidance supports the separation of interface and data responsibilities in maintainable applications (Flutter, 2026). The mobile interface uses flutter_map for interactive map presentation, while the custom Web Administration System uses Django templates and controlled administrative workflows. Visual Studio Code, Git, Figma or an equivalent interface-design tool, internet access, and document-processing software support development and documentation.")
body(doc, "The software repository contains unit, widget, application programming interface, spatial-query, permission, publication-workflow, and integration tests. Fictional or provisional development records are separated from approved research data so that demonstration content cannot be presented as official Bacoor information. Source files, geographic records, and preparedness information are accepted for public use only after their ownership, date, status, limitations, and permitted use have been reviewed.")
placeholder(doc, "PENDING RESEARCH INFORMATION: Insert the verified development-computer and Android test-device specifications after the researchers confirm the actual equipment used", 12, 12)

section_heading(doc, "Method")
placeholder(doc, "Insert the revised Agile iterative software-development lifecycle diagram here; no diagram has been embedded because the researchers will revise and place it", 12, 18)
body(doc, "The development of FloodSense is guided by an Agile iterative software-development lifecycle. The method reflects the team's actual process: requirements are gathered and bounded, converted into manageable features, implemented in small increments, tested, reviewed during consultations, and revised when evidence or adviser feedback changes the design. Agile requirements engineering supports short feedback cycles, testable requirements, and adaptation when stakeholder needs evolve (Hoy & Xu, 2023). Design Thinking activities are used during early discovery and interface exploration, but they are not presented as the complete software-development lifecycle.")
body(doc, "Each development cycle preserves the same research boundary. FloodSense remains a scenario-based assessment and preparedness application unless a later adviser-approved method establishes an operational relationship among current rainfall, water level, station coverage, spatial resolution, and susceptibility. Features that depend on unexamined data or unresolved governance are documented as pending rather than implemented as unsupported conclusions.")

subheading(doc, "Requirements and Evidence Review")
body(doc, "During the first activity, the researchers review survey summaries, interview transcripts, adviser-consultation records, concept documents, feature specifications, geographic-data notes, existing interface designs, and implementation records. Each requested capability is classified as supported, unresolved, deferred, or outside the current scope. Requirements that depend on unexamined datasets or unspecified scientific relationships are not converted into algorithm behavior.")

subheading(doc, "Planning and Analysis")
body(doc, "The researchers define the scenario-based boundary, user roles, assessment inputs, limitation states, output explanations, data provenance, and authorization rules. The five-M categories of Measurement, Manpower, Machinery, Material, and Method connect observed causes to the research problems. User stories and acceptance criteria are prepared for the resident flow, map interaction, location handling, assessment, guidance, verified resources, and administration.")

subheading(doc, "Design")
body(doc, "The design activity organizes the Android application, web services, spatial database, Expert System, Decision Support System, and Web Administration System into components with distinct responsibilities. Interface designs are revised toward a guided sequence instead of one long form. First-launch privacy and limitation information, readable susceptibility states, map legends, explanation content, and manual fallbacks are included in the design.")
placeholder(doc, "The revised system, data, process, and interface diagrams are intentionally omitted from Chapter 3 and will be inserted by the researchers after completing the adviser-required corrections", 12, 12)

subheading(doc, "Implementation")
body(doc, "Implementation proceeds through small and testable increments. The Android application sends validated requests to versioned backend endpoints. The backend resolves supported locations, assembles scenario facts, evaluates the fixed rule set, retrieves source-backed guidance and resources, and returns structured responses. Ordinary administrative routes exclude raw rule editing, and provisional data remain visibly separated from approved records.")

subheading(doc, "Testing")
body(doc, "Each increment is tested at the unit and integration levels before it becomes part of the end-to-end resident or administrator flow. Tests cover request validation, deterministic rule execution, conflict and insufficient-data states, geographic containment, stale response handling, permissions, publication states, map consistency, guidance separation, and safe error messages. The final product evaluation will use the approved ISO/IEC 25010:2023 instrument after the current survey and evaluator plan are completed.")

subheading(doc, "Review and Iteration")
body(doc, "Consultation findings and test failures are converted into documented revisions. A change that affects the research method, scientific parameters, privacy behavior, or system boundary requires corresponding updates to the manuscript, specifications, diagrams, implementation, and tests. The cycle continues until the approved requirements are represented consistently in the research document and the working software.")

section_heading(doc, "Participants of the Study")
body(doc, "Completed requirements gathering included an initial survey of 50 residents from 12 barangays and a supplementary survey of 40 respondents from 14 barangays. These samples are reported separately because they were not administered as one combined survey. Key informants included representatives of the BDRRMO and CPDCO whose institutional responsibilities were relevant to flood information, preparedness, geographic data, and the validation of public records.")
body(doc, "The final participants for the new survey and system evaluation will be reported according to the approved sampling plan, inclusion criteria, respondent classifications, geographic coverage, and completed valid responses. Existing requirements respondents will not be represented as the final evaluation population unless the approved evaluation procedure expressly includes them.")
placeholder(doc, "PENDING RESEARCH DATA: The new survey and final system evaluation are still being conducted. Insert the approved sampling method, target population, final respondent classifications, frequency, percentage, inclusion criteria, and response rate only after validation", 12, 12)

section_heading(doc, "Statistical Treatment of Data")
body(doc, "After the researchers complete and validate the survey records, the following tools may be used to analyze and interpret the gathered data. A statistic will be retained in the final manuscript only when it matches the approved questionnaire, sampling procedure, measurement scale, and analysis plan.")
numbered_item(doc, 1, "Percentage. Percentage will be used to determine the proportion of valid responses for each categorical option and may be computed as P = (f / n) x 100, where P is the percentage, f is the frequency of a response, and n is the number of valid responses for the item.")
numbered_item(doc, 2, "Weighted Mean. For an approved scaled item, the weighted mean may be used to summarize the responses and may be computed as WM = sum(fx) / N, where f is the frequency, x is the assigned scale value, and N is the total number of valid responses.")
numbered_item(doc, 3, "Sample Size Determination. The final sample-size method and computation will be inserted only after the target population, sampling frame, margin of error, confidence assumptions, feasibility limits, and adviser-approved procedure have been established.")
numbered_item(doc, 4, "Likert Scale. The final scale points, verbal interpretations, cutoffs, direction of scoring, and treatment of incomplete responses will be inserted after the evaluation questionnaire and scoring guide have been validated.")
placeholder(doc, "PENDING RESEARCH DATA: Insert final statistical tables and evaluation results only after the ongoing survey and adviser-approved analysis are complete", 12, 12)

# REFERENCES
chapter_heading(doc, "REFERENCES", page_break=True)
section_heading(doc, "Articles")
article_refs = [
    "Flutter. (2026). Architecting Flutter apps. Retrieved from https://docs.flutter.dev/app-architecture",
    "GeoRisk Philippines. (2026). HazardHunterPH: Hazard assessment at your fingertips. Retrieved from https://hazardhunter.georisk.gov.ph/map",
    "International Organization for Standardization. (2023). ISO/IEC 25010:2023 systems and software engineering - Systems and software Quality Requirements and Evaluation (SQuaRE) - Product quality model. Retrieved from https://www.iso.org/standard/78176.html",
    "Japan International Cooperation Agency. (2021, October 1). Flood risk protection project in Cavite inaugurated. Retrieved from https://www.jica.go.jp/english/overseas/philippine/information/press/2021/211001.html",
    "Philippine Atmospheric, Geophysical and Astronomical Services Administration. (2024). Annual report 2024. Retrieved from https://pubfiles.pagasa.dost.gov.ph/pagasaweb/files/transparency/Fiscal%20Year%202024%20Annual%20Report.pdf",
    "United Nations Office for Disaster Risk Reduction. (2025). Mobile technology expanding inclusive early warning communication. Retrieved from https://www.undrr.org/resource/case-study/mobile-technology-expanding-inclusive-early-warning-communication",
]
for item in article_refs:
    reference(doc, item)

section_heading(doc, "Journals")
journal_refs = [
    "Alabbad, Y., Yildirim, E., & Demir, I. (2022). Flood mitigation data analytics and decision support framework: Iowa Middle Cedar Watershed case study. Science of the Total Environment, 814, 152768. https://doi.org/10.1016/j.scitotenv.2021.152768",
    "Alsabhan, W., & Dudin, B. (2023). Real-time flood forecasting and warning: A comprehensive approach toward HCI-centric mobile app development. Multimodal Technologies and Interaction, 7(5), 44. https://doi.org/10.3390/mti7050044",
    "Debnath, J., Sahariah, D., Nath, N., Saikia, A., Lahon, D., Islam, M. N., Hashimoto, S., Meraj, G., Kumar, P., Singh, S. K., Kanga, S., & Chand, K. (2024). Modelling on assessment of flood risk susceptibility at the Jia Bharali River basin in Eastern Himalayas by integrating multicollinearity tests and geospatial techniques. Modeling Earth Systems and Environment, 10, 2393-2419. https://doi.org/10.1007/s40808-023-01912-1",
    "Galanti, R., Coma-Puig, B., de Leoni, M., Carmona, J., & Navarin, N. (2023). An explainable decision support system for predictive process analytics. Engineering Applications of Artificial Intelligence, 120, 105904. https://doi.org/10.1016/j.engappai.2023.105904",
    "Hoy, Z., & Xu, M. (2023). Agile software requirements engineering challenges-solutions: A conceptual framework from systematic literature review. Information, 14(6), 322. https://doi.org/10.3390/info14060322",
    "Johnson, B. A., Estoque, R. C., Li, X., Kumar, P., Dasgupta, R., Avtar, R., & Magcale-Macandog, D. B. (2021). High-resolution urban change modeling and flood exposure estimation at a national scale using open geospatial data: A case study of the Philippines. Computers, Environment and Urban Systems, 90, 101704. https://doi.org/10.1016/j.compenvurbsys.2021.101704",
    "Kostopoulos, G., Davrazos, G., & Kotsiantis, S. (2024). Explainable artificial intelligence-based decision support systems: A recent review. Electronics, 13(14), 2842. https://doi.org/10.3390/electronics13142842",
    "Lin, C. A. (2023). Flood risk management via risk communication, cognitive appraisal, collective efficacy, and community action. Sustainability, 15(19), 14191. https://doi.org/10.3390/su151914191",
    "Mattos, T. S., Oliveira, P. T. S., Bruno, L. S., Carvalho, G. A., Pereira, R. B., Crivellaro, L. L., Lucas, M. C., & Roy, T. (2022). Towards reducing flood risk disasters in a tropical urban basin by the development of flood alert web application. Environmental Modelling & Software, 151, 105367. https://doi.org/10.1016/j.envsoft.2022.105367",
    "Nagumo, N., Ohara, M., Fujikane, M., Inoue, T., Hiramatsu, Y., & Jaranilla-Sanchez, P. A. (2022). Creation of a 3D flood hazard map for a flood-prone area in the Republic of the Philippines and dissemination of the mapping technology. E-journal GEO, 17(1), 123-136. https://doi.org/10.4157/ejgeo.17.123",
    "Nguyen, H. N., Fukuda, H., & Nguyen, M. N. (2024). Assessment of the susceptibility of urban flooding using GIS with an analytical hierarchy process in Hanoi, Vietnam. Sustainability, 16(10), 3934. https://doi.org/10.3390/su16103934",
    "Noori, A., & Bonakdari, H. (2023). A GIS-based fuzzy hierarchical modeling for flood susceptibility mapping: A case study in Ontario, Eastern Canada. Environmental Sciences Proceedings, 25(1), 62. https://doi.org/10.3390/ECWS-7-14242",
    "Oubennaceur, K., Chokmani, K., El Alem, A., & Gauthier, Y. (2021). Flood risk communication using ArcGIS StoryMaps. Hydrology, 8(4), 152. https://doi.org/10.3390/hydrology8040152",
]
for item in journal_refs:
    reference(doc, item)

section_heading(doc, "Research Records")
research_refs = [
    "FloodSense Research Team. (2026a). FloodSense first resident requirements survey summary [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026b). FloodSense supplementary resident requirements survey analysis [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026c). First key informant interview transcript with the Bacoor Disaster Risk Reduction and Management Office [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026d). Second key informant interview transcript with the Bacoor Disaster Risk Reduction and Management Office [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026e). Key informant interview transcript with the Bacoor City Planning and Development Coordinator Office [Unpublished research record]. Cavite State University - Bacoor City Campus.",
    "FloodSense Research Team. (2026f). Thesis adviser consultation transcript and current system decision record [Unpublished research record]. Cavite State University - Bacoor City Campus.",
]
for item in research_refs:
    reference(doc, item)

# Set compact, even cell margins in the comparison table.
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            tc = cell._tc
            tc_pr = tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for edge in ("top", "left", "bottom", "right"):
                node = tc_mar.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    tc_mar.append(node)
                node.set(qn("w:w"), "80")
                node.set(qn("w:type"), "dxa")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
