"""Expand only the Introduction in the current FloodSense working copy."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from lxml import etree


TARGET = Path(
    r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense"
    r"\1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11"
    r"\FloodSense_Chapters_1_to_3_Introduction_Chapter2_Working_Revision_2026-10-12.docx"
)

INTRODUCTION = [
    "The first decisions about a flood are often made before floodwater is visible. A household may protect belongings, change a planned trip, or make arrangements for someone who may need help leaving home. These choices depend on how people understand weather information and the places where they live and move. Preparedness thus begins before an emergency response, when there is still time to consider protective action.",
    "Flood conditions are not uniform across a city. Rainfall interacts with elevation, drainage, land cover, and river systems, so a broad weather update cannot describe the same conditions at every location. A Philippine GIS study used several geographic factors to map flash-flood susceptibility within one barangay, illustrating the importance of local context (Puno et al., 2022). A susceptibility map, a rainfall bulletin, and an official warning also communicate different kinds of information. Their distinctions matter when people interpret what they see before a possible flood.",
    "In the Philippines, flood-related information carries different levels of public authority. The Philippine Atmospheric, Geophysical and Astronomical Services Administration issues official hydro-meteorological information, while national and local disaster-management bodies coordinate action under established response arrangements (National Disaster Risk Reduction and Management Council [NDRRMC], 2024). Information used to consider possible conditions before an event is different from an official forecast, warning, or evacuation instruction. Keeping these roles distinct helps readers understand the purpose and authority of each message.",
    "Bacoor City offers a concrete local setting for examining preparedness. The Japan International Cooperation Agency (JICA, 2021) reported completion of the Imus River Basin subproject, including retarding basins in Bacoor and Imus intended to reduce peak floodwater volume in low-lying Cavite areas. These structures show that flood management in the locality involves river behavior, built infrastructure, and cooperation among public institutions. They also place residents' preparations within a wider setting of public works and local disaster management.",
    "Against this background, the meaning and limits of flood information deserve careful attention. Residents may encounter weather bulletins, barangay announcements, maps, and preparedness guidance as they make decisions before heavy rainfall. Each resource has a different source and purpose, and its relevance depends partly on the location and decision at hand. Examining residents' experiences together with available institutional records provides a grounded basis for the study's research questions.",
]

doc = Document(TARGET)
paragraphs = list(doc.paragraphs)
intro_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "INTRODUCTION")
sop_heading = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Statement of the Problem")
existing = paragraphs[intro_heading + 1 : sop_heading]
if len(existing) != 2 or not existing[0].text.startswith("Flood preparedness begins before water"):
    raise RuntimeError("The working copy's Introduction has changed; refusing to overwrite it.")

anchor = paragraphs[sop_heading]
template = existing[0]
for text in INTRODUCTION:
    element = deepcopy(template._p)
    for child in list(element):
        if etree.QName(child).localname != "pPr":
            element.remove(child)
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))

for paragraph in existing:
    paragraph._p.getparent().remove(paragraph._p)

doc.save(TARGET)
print(TARGET)
