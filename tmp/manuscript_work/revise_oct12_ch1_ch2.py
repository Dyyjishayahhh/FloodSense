"""Create the authorized FloodSense Introduction–Chapter 2 working revision."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
SOURCE = ROOT / "1. FLOODSENSE-MAIN" / "WORKING_REVISION_2026-10-11" / "FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"
OUTPUT = ROOT / "1. FLOODSENSE-MAIN" / "WORKING_REVISION_2026-10-11" / "FloodSense_Chapters_1_to_3_Introduction_Chapter2_Working_Revision_2026-10-12.docx"

doc = Document(SOURCE)
original = list(doc.paragraphs)


def new_paragraph_before(anchor, template, parts):
    """Copy paragraph formatting, not its old text, and insert before anchor."""
    element = deepcopy(template._p)
    for child in list(element):
        if etree.QName(child).localname != "pPr":
            element.remove(child)
    anchor._p.addprevious(element)
    result = Paragraph(element, anchor._parent)
    old_run = template.runs[0] if template.runs else None
    for part in parts:
        value, italic = part[:2]
        run = result.add_run(value)
        if old_run is not None and old_run._r.rPr is not None:
            run._r.insert(0, deepcopy(old_run._r.rPr))
        if italic:
            run.italic = True
        if len(part) > 2:
            run.bold = part[2]
    return result


def replace(start, stop, contents, template_index=None):
    anchor = original[start]
    template = original[start if template_index is None else template_index]
    for item in contents:
        if isinstance(item, str):
            parts = [(item, False)]
            use_template = template
        else:
            parts, use_template = item
        new_paragraph_before(anchor, use_template, parts)
    for idx in range(start, stop + 1):
        old = original[idx]._p
        old.getparent().remove(old)


replace(66, 71, [
    "Flood preparedness begins before water reaches a household or street. People must interpret information about rainfall, familiar places, and public advisories while making decisions for themselves and those in their care. In the Philippines, disaster risk reduction depends in part on accessible information and coordinated local action (NDRRMC, 2024).",
    "Bacoor City lies within the Imus River Basin, where flood-mitigation works have been undertaken for low-lying areas of Cavite (JICA, 2021). The local setting makes it important to examine how residents obtain and understand flood-related information before heavy rainfall. It also calls attention to the geographic and institutional evidence on which responsible preparedness decisions depend.",
])

replace(73, 80, [
    ([("Bacoor residents obtain flood-related information through public advisories, barangay announcements, online sources, and personal contacts. These sources do not necessarily answer the same location-specific or preparedness questions. In the initial purposive survey of 40 eligible adults from 14 Bacoor barangays, 29 respondents (72.5%) described obtaining needed flood information as difficult, very difficult, or usually impossible. This finding describes the participating group, not the prevalence of the problem across the city. ", False), ("What flood-information and preparedness needs do the surveyed residents report, and what evidence is required to address those needs responsibly?", True)], original[73]),
    ([("Flood experiences can disrupt household activities and affect decisions about protection and movement. Among the 40 respondents, 30 (75.0%) said that inadequate information often or always affected their decisions during flooding. The response indicates that information access matters to the surveyed residents, although it does not establish the effectiveness of any proposed application. ", False), ("What flood experiences, information-access difficulties, and preparedness needs are reported by the participating Bacoor residents?", True)], original[73]),
    ([("Understanding conditions at a particular place and finding dependable evacuation information are distinct tasks. Thirty respondents (75.0%) found it difficult, very difficult, or did not know how to determine their location's flood susceptibility, while 33 (82.5%) found verified evacuation-center information difficult to locate or did not know where to find it. Institutional records also require attention to source ownership, authorization, and geographic validation. ", False), ("What geographic, knowledge, preparedness, privacy, provenance, and administrative requirements can be established from the resident, interview, and documentary evidence?", True)], original[73]),
    ([("Most survey participants could access an Android device: 34 owned one and five shared one. However, reliable internet access was not universal, and the participants' technology conditions alone do not determine a suitable assessment design. The study must organize the documented needs without treating hypothetical scenarios as live observations or unverified geographic records as official hazard data. ", False), ("How can the established requirements be developed into a resident-requested, pre-event assessment and a distinct preparedness-support process with appropriate administrative controls?", True)], original[73]),
    ([("A developed application cannot be regarded as technically dependable merely because its screens and records exist. Its resident, geographic, assessment, guidance, and administrative functions must be checked against defined behaviors and data boundaries. ", False), ("To what extent do these functions operate correctly and consistently under specified technical test cases?", True)], original[73]),
    ([("Technical verification does not by itself establish how intended users and evaluators perceive the system's quality. A separate evaluation must use an approved instrument and defined evaluator groups; initial needs-survey responses cannot substitute for post-use findings. ", False), ("How will the relevant ISO/IEC 25010 product-quality characteristics of the completed system be evaluated?", True)], original[73]),
])

replace(82, 93, [
    ([("Generally, the study seeks to develop and evaluate FloodSense as a scenario-based Android system for explainable flood-susceptibility assessment and separate pre-event preparedness support in Bacoor City.", False)], original[82]),
    ([("Specifically, it seeks to:", False)], original[83]),
    ([("1. identify the flood experiences, information-access difficulties, preparedness needs, and technology conditions reported by eligible Bacoor City respondents;", False)], original[84]),
    ([("2. determine the geographic, knowledge, rainfall-scenario, privacy, provenance, preparedness, and administrative requirements supported by resident, interview, and documentary evidence;", False)], original[84]),
    ([("3. develop FloodSense as a Flutter-based Android resident application and Django-based Web Administration System through an iterative Agile development process, with these major functions:", False)], original[84]),
    ([("• authenticated resident and authorized-administrator access with role-based permissions and privacy controls;", False)], original[87]),
    ([("• user-confirmed rainfall scenario and location inputs, including barangay selection, manual map pin, and optional foreground location;", False)], original[87]),
    ([("• deterministic susceptibility assessment, explanation, and geographic presentation;", False)], original[87]),
    ([("• separate preparedness guidance and evacuation-center reference information; and", False)], original[87]),
    ([("• controlled administration of geographic references, parameters, guidance, sources, and review records;", False)], original[87]),
    ([("4. test the resident, administrative, geographic, Expert System, and Decision Support System functions through applicable unit, integration, API, spatial, and end-to-end tests; and", False)], original[84]),
    ([("5. evaluate selected product-quality characteristics of FloodSense using an approved ISO/IEC 25010-based instrument.", False)], original[84]),
])

# Theoretical Framework paragraphs 94–102, including its figure, are untouched.
replace(103, 103, [
    ([("System Architecture", False)], original[104]),
    ([("[Insert System Architecture]", False)], original[103]),
])

replace(111, 111, [
    "The study is centered in Bacoor City, Cavite, with development and manuscript preparation at Cavite State University–Bacoor City Campus. Interviews with the Bacoor City Disaster Risk Reduction and Management Office and City Planning and Development Office were conducted in May 2026, and the initial resident survey collected responses in October 2026. The principal development and manuscript period extends from September 2026 through the planned completion in May 2027.",
])

replace(113, 118, [
    "FloodSense covers a connected Android application for Bacoor City residents and a browser-based administration system for authorized personnel. Residents can authenticate, select or confirm a hypothetical rainfall intensity and duration, identify a location by barangay or map pin or use optional one-time foreground GPS, confirm the barangay, and explicitly request a pre-event assessment. A deterministic Rule-Based Expert System provides an explained susceptibility classification and map context; a separate Decision Support System provides source-backed preparedness guidance and evacuation-center reference information. The administration system governs authorized access and review of rainfall and map references, assessment parameters, preparedness content, evacuation-center records, sources, and audit records. Ordinary administrators cannot rewrite raw rules, priorities, conflict-resolution logic, or algorithm code.",
    "However, FloodSense is limited to scenario-based assessment and preparedness support within its documented Bacoor administrative coverage. It does not continuously monitor rainfall or water level, provide official forecasts or warnings, issue autonomous alerts or evacuation orders, track GPS in the background, determine safe routes, report live center occupancy, or use machine-learning prediction. Its 47-barangay layer is an administrative reference rather than a City-verified flood-susceptibility dataset; the authority and local validity of assessment inputs and guidance remain dependent on source documentation, expert review, and institutional validation. Evacuation-center distance is a straight-line reference and does not establish road safety, accessibility, opening status, or capacity. Connected functions require a compatible device and network access, while personal and temporary location information remain subject to consent and purpose limitation. The purposive survey of 40 respondents cannot represent every Bacoor resident.",
])

# Correct operational definitions that implied expert or institutional validation.
replace(133, 133, [
    ([("Knowledge Base. ", False, True), ("The structured collection of documented facts, conditional rules, geographic attributes, rainfall parameters, susceptibility categories, and provenance records used by the Expert System. These materials require appropriate expert and institutional validation before they can be treated as authoritative local evidence.", False, False)], original[133]),
])
replace(127, 127, [
    ([("Expert Knowledge and Validation. ", False, True), ("Flood-related facts, parameters, rules, preparedness guidance, and supporting information that require review by qualified experts and adequate provenance before being treated as locally valid. Repository presence or technical validation alone does not establish institutional approval.", False, False)], original[127]),
])
replace(134, 134, [
    ([("Rainfall Scenario Input. ", False, True), ("A user-selected hypothetical rainfall intensity and duration used to request a pre-event assessment. The choices are scenario labels, not automatically updated observations, official forecasts, or official warning thresholds.", False, False)], original[134]),
])

# Each retained related study now follows purpose, methodology, results, and
# discussion in four continuous paragraphs under its existing study heading.
replace(212, 213, [
    "Wisudawan (2021) investigated a mobile-connected Flood Early Warning System for a reservoir at Universitas Islam Indonesia in Sleman, Yogyakarta. The work addressed the need to observe local environmental conditions and communicate warning levels to a mobile user. Its purpose was continuous site monitoring, not the assessment of a hypothetical rainfall scenario.",
    "The researcher assembled an Arduino Nano with rainfall, water-level, temperature, and humidity sensors, GSM communication, and the Blynk Android application. The prototype was checked in sensor tests and then observed at the reservoir for 30 days. Readings were categorized into Normal, Advisory, Watch, and Warning states; the source did not name a formal software-development life-cycle model.",
    "During the observation period, the reported mean temperature was 28.6 degrees Celsius, mean humidity was 63.7%, and mean water level was 54.5 centimeters. The paper documented transmission and display of readings and warning states. No flood occurred during field observation, so the study did not demonstrate warning performance during an actual flood event.",
    "The prototype illustrates how environmental measurements can be presented understandably on a mobile interface. Its one-reservoir setting, short observation period, and sensor thresholds limit transferability to other places. FloodSense instead uses a resident-confirmed hypothetical scenario and deterministic susceptibility rules; the Indonesian sensor readings cannot validate Bacoor City conditions.",
])
replace(215, 216, [
    "Meléndez-Landaverde and Sempere-Torres (2023) examined A4alerts, a mobile application for disseminating site-specific, impact-based flood warnings in Catalonia, Spain. Their research asked whether users considered the application and its warning-related functions useful in an emergency context. The app drew on an operational warning environment rather than a stand-alone hypothetical susceptibility exercise.",
    "Seventy-six civil-protection, emergency-response, and local stakeholders completed an online interactive scenario experiment between January and March 2022. The researchers assessed perceived usefulness, functional appropriateness, and willingness to adopt the app and compared feature rankings using Friedman's test and Bonferroni-adjusted Wilcoxon comparisons. The participant experience was a demonstration, not observation of behavior during a real flood.",
    "Ninety-seven percent of participants considered A4alerts useful during flood emergencies, while 93% said they would use or recommend it. Site-specific impact warnings, official regional warnings, and a self-protection checklist were among the most valued functions. These are reported perceptions of the recruited participants, not evidence of an observed reduction in flood losses.",
    "The study supports careful presentation of place-specific information and actionable guidance. Its stakeholder sample, Catalan warning infrastructure, and scenario experiment constrain generalization to Bacoor City residents. FloodSense can draw on its communication lessons but does not disseminate official operational warnings or inherit the evaluated app's performance.",
])
replace(218, 219, [
    "Alsabhan and Dudin (2023) developed and evaluated a human-computer-interaction framework and prototype for mobile real-time flood forecasting and warning. The study focused on the tasks and information needs of front-end users, weather-station users, and administrators. Its research problem concerned a multi-role forecasting environment rather than pre-event assessment initiated by a resident.",
    "Activity theory and scenario-based design guided nominal-group sessions with five prospective users and five experts. Participants generated 63 situations, from which 12 priority scenarios informed requirements and prototype tasks. The researchers evaluated task performance and satisfaction across usefulness, ease of use, learnability, satisfaction, and overall experience; no Agile, Waterfall, or other SDLC was stated as the study's development model.",
    "The process produced 87 initial requirements and six additional requirements after testing, for a total of 93. Mean satisfaction-category scores ranged from 3.93 to 4.08, although some analytical tasks were harder for participants. The reported prototype and evaluation therefore give evidence about requirements and interaction quality, not about forecast accuracy in Bacoor City.",
    "This study shows why role-specific access and understandable interactions matter in a flood-information application. Its small participant groups and forecasting architecture limit wider inference. FloodSense shares the need for comprehensible outputs and permissions, but its deterministic, resident-requested susceptibility assessment is separate from real-time forecasting and warning.",
])
replace(221, 222, [
    "Msusa et al. (2025) investigated a climate-smart Web GIS application for multi-hazard early warning in Tanzania. Flood analysis focused on Dar es Salaam, while drought analysis focused on Dodoma. The authors sought to bring modeled hazard information and local knowledge into an accessible geographic application for planning and warning purposes.",
    "For flooding, the researchers used Height Above Nearest Drainage analysis, GloFAS forecast information, and community-mapped flood points for spatial comparison. For drought, they used geospatial climate inputs and a ConvLSTM model. The paper expressly identified Rapid Application Development as its Web GIS development process; these hydrological, machine-learning, and development procedures served different methodological roles.",
    "The mapped flood points showed qualitative spatial agreement with the flood-model output. The reported median ROC-AUC of 0.87 and median F1 score of 0.81 refer to the drought model, not to the flood analysis. The application brought these multi-hazard outputs together in a Web GIS interface.",
    "The study demonstrates the value of combining spatial analysis, local validation, and accessible map presentation. Its forecast data, modeling choices, and Tanzanian geographic conditions cannot be transferred directly to Bacoor City. FloodSense uses deterministic scenario-based assessment and cannot claim the study's early-warning or machine-learning results as its own.",
])
replace(225, 226, [
    "Bentoso et al. (2021) presented a web-based flood-warning decision-support prototype for the Province of Leyte, Philippines. The project addressed the need to display rainfall and water-level conditions in a form usable by authorities making warning and flood-control decisions. Its subject was measured environmental conditions, not a resident-selected hypothetical scenario.",
    "The available institutional publication describes rainfall and water-level sensing, transmission of readings to a web interface, graphical reporting, and alert messages for responsible personnel. This is a documented technical procedure, but the accessible institutional record does not establish a named research design or software-development model. Details of sample size, field duration, and threshold calibration were not available for independent confirmation from that record.",
    "The authors reported a functional prototype that graphically displayed the flooding situation and sent alert-level messages to support monitoring and planning. The accessible record does not establish province-wide deployment or measured preparedness outcomes. The reported output is therefore treated as a prototype contribution rather than proof of operational effectiveness throughout Leyte.",
    "The work is relevant to digital presentation of flood information for decision support. Its sensor-and-alert design differs from FloodSense's resident-requested susceptibility assessment and separate preparedness guidance. Leyte-specific readings and thresholds cannot establish the validity of Bacoor City classifications, and the limited accessible methodological detail constrains comparison.",
])
replace(228, 229, [
    "Carpio, Asnaul, Said, and Amerin (2022) developed an Arduino and GSM flood-monitoring system for Riverside Barangay in Isabela City, Basilan. The system was intended to observe water conditions and communicate warning signals to local disaster personnel. It addressed a local monitoring problem rather than the interpretation of a hypothetical resident-chosen rainfall scenario.",
    "The authors explicitly used an Agile Development Model for incremental system construction and a quantitative survey design for evaluation. Twenty City Disaster Risk Reduction and Management Office personnel participated in a demonstration and initial device test. They assessed functionality, design, security, efficiency, and reliability, making the survey an evaluation of the prototype under the reported conditions.",
    "The paper reported an overall mean of 3.70, interpreted by its authors as Acceptable, and described warning-signal operation during testing. It also reported that battery operation was limited to approximately ten hours. These results describe the tested device and evaluator group; they do not show sustained warning performance across Isabela City.",
    "The Philippine setting provides a useful comparison for local disaster-information systems and for distinguishing development method from evaluation design. However, its hardware, threshold conditions, respondent group, and measured outcomes cannot be transferred to Bacoor City. FloodSense does not continuously monitor water level or send autonomous SMS warnings.",
])
replace(231, 232, [
    "Dela Cruz et al. (2024) developed an IoT flood-alert prototype for Barangay Narvacan in Sto. Tomas, La Union. The project responded to a community need for water-level information and alert delivery through a connected system. Its purpose was to develop the local prototype and examine its acceptability among intended community participants.",
    "The authors identified descriptive and developmental research designs and used Agile for application development. Interviews and observations informed the prototype, while purposive sampling selected nine barangay council officials and 21 residents from seven sectors, for 30 evaluators in all. The User Technology Acceptability Questionnaire measured perceived usefulness, information quality, and interface quality after testing.",
    "On the study's seven-point scale, the reported criterion means were 6.87 for usefulness, 6.83 for information quality, and 6.82 for interface quality. The source also prints a grand mean of 6.87, which does not match the simple average of those three displayed criterion means; this review therefore reports the individual criterion results without adopting the inconsistent aggregate. The ratings document acceptability among the selected evaluators, not field effectiveness during a flood.",
    "The work offers a local example of combining community consultation, iterative development, and user evaluation. Its single-barangay sample, continuous water-level sensing, and SMS-alert function differ from FloodSense's pre-event scenario assessment. Neither its acceptability ratings nor its sensor thresholds can validate a Bacoor City susceptibility result.",
])
replace(234, 235, [
    "Delos Reyes, Cabreros, Paulino, and Garcia (2020) developed BantayBaha for Barangay Virgen Delas Flores in Baliwag, Bulacan. The project addressed the barangay's manual water-level measurement and megaphone-based warning practices. It proposed a connected monitoring, website-logging, and SMS-notification prototype for that setting.",
    "The authors interviewed local officials and rescue personnel to understand practices and define water-level states. They built an IoT prototype, carried out verification, validation, and repeated tests, and then asked 70 Baliuag University students to evaluate it. The source documents these procedures but does not expressly name a formal SDLC model.",
    "The student evaluators rated functionality at 4.70, user interface at 4.50, and reliability at 4.57 on the paper's five-point scale; the authors interpreted each as Very Satisfactory. The prototype also demonstrated water-level recording and SMS notification under its test conditions. These ratings do not constitute a field trial of warning impact among all barangay residents.",
    "BantayBaha illustrates a Philippine effort to improve local communication of monitored water levels. Its single-barangay location, student evaluator group, and hardware requirements limit generalization. FloodSense does not adopt continuous sensing or automated SMS alerts, and the Bulacan results do not establish the validity of Bacoor City scenario assessments.",
])

replace(245, 245, [
    "The following tables compare the documented methodologies or technical procedures, objectives, and contributions of the foreign and Philippine system studies retained in this chapter. Each entry draws on an original publication or institutional record that describes a research or development procedure and an output, although the accessible methodological detail is not equal for every source.",
])

for index in [210, 211, 214, 217, 220, 223, 224, 227, 230, 233, 236, 241, 244, 246, 247]:
    original[index].paragraph_format.keep_with_next = True

# Keep the manually formatted table; correct one internally inconsistent statistic.
cell = doc.tables[1].cell(3, 4)
cell.paragraphs[0].runs[0].text = (
    "It reported means of 6.87 for usefulness, 6.83 for information quality, "
    "and 6.82 for interface quality among 30 evaluators."
)
for paragraph in cell.paragraphs:
    for run in paragraph.runs[1:]:
        run.text = ""

# Match SURECUT's readable, left-aligned body cells without changing the
# manually chosen column widths or table indents.
for table in doc.tables[:2]:
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        if row_index == 0 and row_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is None:
            row_properties.append(OxmlElement("w:tblHeader"))
        if row_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if row_index == 0
                                       else WD_ALIGN_PARAGRAPH.LEFT)
                paragraph.paragraph_format.left_indent = Pt(0)
                paragraph.paragraph_format.right_indent = Pt(0)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8 if row_index == 0 else 9)

doc.save(OUTPUT)
print(OUTPUT)
