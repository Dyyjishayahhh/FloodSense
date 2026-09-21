from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


SOURCE = Path(r"1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Chapter_1_MAIN.docx")
OUT = Path(r"1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Chapters_1_to_3_Development_Aligned_2026-09-20.docx")
FONT = "Arial"
SIZE = 11


def set_run(run, size=SIZE, bold=False, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph(paragraph, text, lead=None):
    clear_paragraph(paragraph)
    if lead and text.startswith(lead):
        set_run(paragraph.add_run(lead), bold=True)
        set_run(paragraph.add_run(text[len(lead):]))
    else:
        set_run(paragraph.add_run(text))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Inches(0.5)
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def body(doc, text, first_indent=True, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if first_indent:
        fmt.first_line_indent = Inches(0.5)
    set_run(p.add_run(text), italic=italic)
    return p


def lead(doc, heading_text, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Inches(0.5)
    set_run(p.add_run(heading_text), bold=True)
    set_run(p.add_run(text))
    return p


def chapter(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.page_break_before = True
    fmt.keep_with_next = True
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(24)
    set_run(p.add_run(text.upper()), bold=True)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.keep_with_next = True
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(24)
    fmt.space_after = Pt(0)
    set_run(p.add_run(text), bold=True)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.keep_with_next = True
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    set_run(p.add_run(text), bold=True)
    return p


def note(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = p.paragraph_format
    fmt.keep_together = True
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(12)
    set_run(p.add_run(text), italic=True)
    return p


def reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.5)
    fmt.first_line_indent = Inches(-0.5)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(12)
    set_run(p.add_run(text))
    return p


def configure_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, result, end])
    set_run(run)


def set_page_start(section, number):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        section._sectPr.append(node)
    node.set(qn("w:start"), str(number))


doc = Document(SOURCE)

# Preserve the title/preliminary material and the original theoretical framework.
# Remove the original reference list because Chapters 2 and 3 require one combined list.
reference_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "REFERENCES")
for p in list(doc.paragraphs[reference_index:]):
    p._element.getparent().remove(p._element)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(SIZE)

# These changes are outside the Theoretical Framework. They apply the confirmed
# administrator boundary and distinguish current functions from future requirements.
replacements = {
    "The Android application presents the assessment through a dynamic GIS-style map": (
        "The current FloodSense prototype presents scenario-based susceptibility assessment through an interactive map. The user selects or confirms a hypothetical rainfall intensity and duration, selects a supported location manually or through an optional one-time foreground location request, and explicitly requests an assessment. The map can display demonstration susceptibility zones and a neutral reference layer representing the 47 current barangay identities; the reference layer is technically validated for geometry but is not City-issued, is not City-verified, and contains no flood-susceptibility facts. The custom Web Administration System currently provides authenticated dashboard access, read-only review of map and rainfall-reference data, parameter-governance views, and controlled workflows for Decision Support System content, evacuation-center records, source records, and audit history. Ordinary administrators do not edit raw Expert System rules, priorities, conflict-resolution logic, or algorithm code. A fixed-schema CSV export-and-import workflow for permitted values is a confirmed requirement, but it is not yet implemented. FloodSense does not provide real-time monitoring, official forecasting, autonomous alerts, or automatic evacuation orders.",
        None,
    ),
    "develop an Android application using Flutter and Dart": (
        "develop an Android application using Flutter and Dart, supported by Django, Django REST Framework, GeoDjango, and PostgreSQL/PostGIS, that integrates a deterministic Rule-Based Expert System, scenario-based rainfall input, an interactive GIS-style map, manual location selection, optional one-time foreground location, explainable susceptibility output, and separate Decision Support System guidance; provide a custom Web Administration System for authorized review and maintenance of permitted content and records; protect raw Expert System rules and algorithm logic from ordinary administrator editing; and establish a fixed-schema CSV exchange requirement for permitted values without representing that requirement as operational until its routes, validation, interface, integration, and tests are completed;",
        None,
    ),
    "For the Bacoor Disaster Risk Reduction and Management Office (BDRRMO).": (
        "For the Bacoor Disaster Risk Reduction and Management Office (BDRRMO). The study provides a prototype channel through which validated susceptibility information and pre-event preparedness guidance may eventually be communicated to residents. The custom Web Administration System currently supports controlled management of DSS guidance, evacuation-center records, source information, and audit history, together with read-only review of map, rainfall-reference, and parameter records. Ordinary administrators are not permitted to rewrite raw Expert System rules or inference logic. The proposed fixed-schema CSV workflow will permit only validated values within system-required fields after implementation and authorization. These functions complement rather than replace official advisories, warnings, and emergency directives.",
        "For the Bacoor Disaster Risk Reduction and Management Office (BDRRMO). ",
    ),
    "Following data gathering, the study progressed through the Define, Ideate, Prototype, and Test phases": (
        "Following data gathering, the study proceeds through the Define, Ideate, Prototype, and Test phases within the research timeline. Requirements, interface wireframes, geographic data structures, susceptibility rules, DSS content, mobile screens, and Web Administration System functions are refined through consultations and development evidence. Final resident evaluation and system-quality analysis remain future activities. Results will be reported only after the approved instrument, sampling procedure, completed responses, and analysis are available.",
        None,
    ),
    "The study developed FloodSense, an Android application using a Rule-Based Expert System": (
        "The study develops FloodSense as a scenario-based, pre-event flood-susceptibility assessment and preparedness-support system for Bacoor City. Its deterministic Rule-Based Expert System evaluates a user-confirmed rainfall scenario and supported location, while the separate Decision Support System presents published preparedness guidance without changing the susceptibility result. The current prototype produces Low, Moderate, High, or Very High demonstration outputs and explanations. Operational use remains blocked until approved local susceptibility data, validated parameters, and institutional authorization are available. FloodSense does not provide real-time flood monitoring, official forecasts, warnings, evacuation orders, or safety guarantees.",
        None,
    ),
    "The primary users of the Android application are Bacoor City residents": (
        "The primary intended users of the Android application are Bacoor City residents. In the current mobile flow, a resident selects Light, Moderate, Heavy, Intense, or Torrential rainfall and a duration of 1, 3, 6, 12, or 24 hours, confirms the hypothetical scenario, selects or confirms a location, and explicitly requests an assessment. Resident registration, login, persistent profiles, and assessment history are not exposed in the current interface and remain unimplemented for the intended user flow.",
        None,
    ),
    "The dynamic GIS-style map is limited to the territorial jurisdiction of Bacoor City": (
        "The interactive map is limited to supported Bacoor City reference and demonstration data. It can pan and zoom, display scenario-dependent demonstration colors, accept manual area selection, and show a neutral 47-barangay administrative reference boundary. The boundary layer represents current barangay identities through a documented derivation and has passed technical geometry checks, but it is not automatically City-issued or City-verified and contains no flood-susceptibility facts. The map does not display live flood extent, water depth, traffic conditions, road closures, river level, center occupancy, or satellite rainfall imagery.",
        None,
    ),
    "Check My Area allows the user to position a pin on the map": (
        "Location functions allow the user to select an area manually, place or adjust a map pin, or request the device's location once while the application is in the foreground. The application can propose a barangay from the coordinate and requires the user to confirm or correct it before assessment. The coordinate is used for the active request and is not continuously collected in the background. The current function does not find safe routes, verify road passability, or continuously track the resident.",
        None,
    ),
    "Evacuation-center locations maintained from authorized records are displayed as map points": (
        "The backend and custom Admin portal include an evacuation-center record workflow. This capability does not yet expose resident-facing nearby-center discovery, distance ranking, live capacity, operating status, or safe-route guidance. Any future resident display must rely on authorized and current records and must direct users to confirm emergency instructions through official channels.",
        None,
    ),
    "The Decision Support System is confined to pre-event preparedness": (
        "The Decision Support System is confined to pre-event preparedness and remains separate from the susceptibility assessment. The current implementation can return published guidance associated with the assessment context. A resident-facing structured multi-question decision tree has not been verified as complete. The DSS does not alter the susceptibility class, issue autonomous evacuation orders, diagnose conditions, guarantee safety, or make the resident's final decision.",
        None,
    ),
    "A separate Web Administration System is included for authorized administrators": (
        "A separate custom Web Administration System provides authenticated access to a dashboard, read-only map-data and rainfall-reference review, parameter-governance views, and workflows for DSS content, evacuation-center records, data sources, and audit history. Ordinary administrators cannot directly edit raw Expert System rules, rule conditions, priorities, conflict-resolution logic, or the algorithm. Technical Django Admin remains a separate developer or research-maintenance interface. The confirmed fixed-schema CSV export-and-import workflow is not yet operational and will require routes, services, schema and value validation, interface controls, permissions, integration, and tests before it can be described as implemented.",
        None,
    ),
    "The feasible implementation stack consists of Flutter and Dart": (
        "The implementation stack uses Flutter and Dart for the mobile application; flutter_map for resident-facing mapping; Python Django and Django REST Framework for backend services and REST interfaces; GeoDjango for geographic operations; PostgreSQL with PostGIS for relational and spatial records; and Django templates with Leaflet for the custom administrative interface. The current prototype requires network access for integrated server functions. iOS support, fully offline operation, live sensors, automated meteorological feeds, background polling, background GPS tracking, emergency dispatch, and machine-learning prediction are outside the current scope.",
        None,
    ),
    "The reliability of the susceptibility assessment depends on the completeness": (
        "The reliability of any susceptibility assessment depends on the completeness, geographic resolution, provenance, institutional approval, and expert validation of the data and knowledge used. No approved operational susceptibility dataset is currently present in the repository. The available 47-barangay layer is an administrative reference only, while provisional research materials remain unvalidated or restricted. Demonstration results must therefore not be presented as official Bacoor City flood-hazard information.",
        None,
    ),
    "Data privacy and security are addressed in accordance with Republic Act No. 10173": (
        "Data privacy and security are addressed in accordance with Republic Act No. 10173, or the Data Privacy Act of 2012. The current location flow requests foreground permission only when the user chooses the location function, uses the coordinate for the active request, permits confirmation or correction of the detected barangay, and does not implement continuous background location tracking. Administrative functions require authenticated and authorized access. Evaluation findings will not be generalized until valid survey and system-evaluation data are available.",
        None,
    ),
    "Android Application. The resident-facing FloodSense software": (
        "Android Application. The resident-facing FloodSense software developed with Flutter. Its current interface supports scenario selection and confirmation, manual or one-time foreground location input, barangay confirmation or correction, assessment requests, map presentation, susceptibility results, explanations, and preparedness guidance. Resident accounts, assessment history, notifications, and offline operation are not currently exposed.",
        "Android Application. ",
    ),
    "Check My Area. A feature that allows a user to position a pin": (
        "Check My Area. The location-selection flow that allows manual area selection, placement or adjustment of a map pin, and an optional one-time foreground location request. A detected barangay is proposed for user confirmation or correction before assessment; the function does not continuously track the device.",
        "Check My Area. ",
    ),
    "Decision Support System (DSS).": (
        "Decision Support System (DSS). The component separate from the Rule-Based Expert System that presents published pre-event preparedness guidance using the assessment context. It does not change the susceptibility result, issue an official warning or evacuation order, or make the user's final decision.",
        "Decision Support System (DSS). ",
    ),
    "Dynamic Map. The GIS-style map": (
        "Dynamic Map. The interactive map that supports panning, zooming, scenario-dependent demonstration coloring, manual area selection, map-pin placement, and display of the neutral 47-barangay reference layer. The reference boundary is not an official flood-hazard dataset.",
        "Dynamic Map. ",
    ),
    "Expert-Provided and Expert-Validated Knowledge.": (
        "Expert-Provided and Expert-Validated Knowledge. Flood-related facts, parameters, rules, preparedness guidance, and supporting information that have been supplied or reviewed by an authorized source and recorded with adequate provenance. Repository presence or technical validation alone does not establish institutional approval.",
        "Expert-Provided and Expert-Validated Knowledge. ",
    ),
    "Flood Susceptibility. The relative potential": (
        "Flood Susceptibility. The relative potential of a supported geographic area to experience rainfall-related flooding under a user-confirmed scenario. The current prototype demonstrates Low, Moderate, High, or Very High outputs but does not yet contain an approved operational Bacoor City susceptibility dataset and is not a real-time prediction.",
        "Flood Susceptibility. ",
    ),
    "PAGASA-Aligned Rainfall Input.": (
        "PAGASA-Aligned Rainfall Input. The current interface labels for Light, Moderate, Heavy, Intense, and Torrential rainfall and durations of 1, 3, 6, 12, or 24 hours. These selections are hypothetical scenario inputs. Their operational thresholds and use require authoritative source confirmation and validation before official deployment.",
        "PAGASA-Aligned Rainfall Input. ",
    ),
    "Preparedness Provenance.": (
        "Preparedness Provenance. Information identifying the source, validation status, responsible actor, relevant date, and version of a preparedness item or related record. Provenance permits users and administrators to distinguish approved, provisional, and unsupported information.",
        "Preparedness Provenance. ",
    ),
    "Web Administration System. The browser-based interface": (
        "Web Administration System. The custom browser-based interface used by authorized administrators for the dashboard, read-only map and rainfall-reference review, parameter governance, DSS-content management, evacuation-center management, data-source management, and audit history. It does not expose raw rule editing. The separate technical Django Admin is reserved for developer or research maintenance. Fixed-schema CSV export and import is a confirmed but currently unimplemented requirement.",
        "Web Administration System. ",
    ),
}

theory_started = False
theory_ended = False
for p in doc.paragraphs:
    stripped = p.text.strip()
    if stripped == "Theoretical Framework":
        theory_started = True
    if stripped == "Significance of the Study":
        theory_ended = True
    # Never alter the Theoretical Framework or its System Architecture content.
    if theory_started and not theory_ended:
        continue
    for start, (new_text, lead_text) in replacements.items():
        if p.text.startswith(start):
            replace_paragraph(p, new_text, lead_text)
            break

# Create a new section before the Introduction so the main text has Arabic page numbering.
intro_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "INTRODUCTION")
prelim_break_p = doc.paragraphs[intro_idx - 1]
ppr = prelim_break_p._p.get_or_add_pPr()
if ppr.find(qn("w:sectPr")) is None:
    sect = deepcopy(doc.sections[-1]._sectPr)
    for tag in (qn("w:headerReference"), qn("w:footerReference"), qn("w:pgNumType")):
        for node in list(sect.findall(tag)):
            sect.remove(node)
    type_node = sect.find(qn("w:type"))
    if type_node is None:
        type_node = OxmlElement("w:type")
        sect.insert(0, type_node)
    type_node.set(qn("w:val"), "nextPage")
    ppr.append(sect)

doc.save(OUT)
doc = Document(OUT)

# Configure regular pages but retain the original full-page framework section.
for idx, section in enumerate(doc.sections):
    if idx != 2:
        configure_page(section)

if len(doc.sections) >= 2:
    main = doc.sections[1]
    main.header.is_linked_to_previous = False
    clear_paragraph(main.header.paragraphs[0])
    add_page_number(main.header.paragraphs[0])
    set_page_start(main, 1)
    for section in doc.sections[2:]:
        section.header.is_linked_to_previous = True


# CHAPTER 2
chapter(doc, "REVIEW OF RELATED LITERATURE")
body(doc, "This chapter reviews current literature and studies relevant to the FloodSense concept. The discussion covers flood susceptibility, scenario-based assessment, deterministic and explainable reasoning, decision support, geographic information, location privacy, data provenance, disaster preparedness, and software quality. Findings, factors, weights, and formulas from other settings are not treated as Bacoor City parameters unless supported by local evidence, institutional approval, and expert validation.")

heading(doc, "Related Foreign Literature")
subheading(doc, "Flood Susceptibility and Scenario-Based Assessment")
body(doc, "Flood susceptibility describes the relative tendency of a location to experience flooding under stated conditions, rather than the timing or certainty of a particular flood. Membele, Naidu, and Mutanga (2022) observed that susceptibility and vulnerability mapping in developing countries is shaped by the availability, scale, and quality of spatial and socioeconomic data. Debnath et al. (2024) likewise showed that geospatial and expert-informed approaches require explicit factors, defensible spatial preparation, and validation. These findings support FloodSense's use of explicit scenario inputs and limitation states, while also showing why demonstration data cannot be represented as official local susceptibility information.")
body(doc, "Nguyen, Fukuda, and Nguyen (2024) applied a Geographic Information System and analytical hierarchy process to urban flood-susceptibility assessment using multiple conditioning factors, expert consultation, and historical validation points. The method demonstrates the need to document how factors are selected and validated. Its weights and conclusions are not transferred to FloodSense because they were developed for a different location and evidence base.")

subheading(doc, "Explainable Rule-Based Reasoning")
body(doc, "Explainability is important when a system presents a classification that may influence preparedness decisions. Kostopoulos, Davrazos, and Kotsiantis (2024) reviewed explainable decision-support systems and emphasized transparency, interpretability, and understandable justification. FloodSense uses deterministic rule matching so that a result can be associated with the applicable scenario facts, matched rule, and explanation. This design does not involve machine-learning prediction or self-modifying rules.")
body(doc, "The need for explainability does not require exposing rule editing to ordinary administrators. FloodSense separates the presentation of understandable result explanations from the technical maintenance of rules and algorithm logic. The custom Admin portal is therefore bounded to approved administrative workflows, while developer or research maintenance remains separate.")

subheading(doc, "Decision Support for Preparedness")
body(doc, "A Decision Support System organizes information and alternatives to help people make decisions without removing human judgment. Alabbad, Yildirim, and Demir (2022) presented a flood-mitigation decision-support framework that combined geographic information, scenarios, analysis, and decision support. For FloodSense, the susceptibility result is produced by the Rule-Based Expert System, while preparedness guidance is delivered by a separate DSS component. The guidance does not alter the assessment and does not become an official warning or evacuation order.")

subheading(doc, "GIS and Public Risk Communication")
body(doc, "Oubennaceur et al. (2021) demonstrated that interactive maps can combine spatial information, text, and user interaction to communicate flood risk to non-specialist audiences. Clear legends, labels, explanatory text, and careful limits are therefore necessary in FloodSense. A boundary or basemap identifies location but does not, by itself, justify a flood-susceptibility classification.")
body(doc, "Fang et al. (2023) described a service-oriented approach that integrates geospatial resources and task chains for disaster decision support. The study supports separating data, processing, and user-facing services while maintaining traceable connections between them. FloodSense follows a comparable separation through a mobile client, REST services, deterministic assessment logic, a spatial database, and administrative workflows, while remaining narrower in scope and pre-event purpose.")

subheading(doc, "Provenance and Responsible Geographic Data Use")
body(doc, "Tullis and Kar (2021) explained that provenance is essential to ethical replicability and reproducibility in GIScience. Fakhruddin et al. (2022) likewise emphasized risk-informed data and the importance of understanding how disaster and climate data are produced and used. These principles support source records, validation status, version information, and explicit distinction among approved, provisional, technically checked, and unsupported FloodSense data.")

subheading(doc, "Foreground Location and Privacy")
body(doc, "Mobile location is sensitive personal information because it can reveal a person's movements and context. Android's permission model distinguishes foreground access from background access and encourages requesting only the access needed for the active function (Android Developers, 2026). FloodSense therefore limits its current location behavior to an optional user-initiated foreground request, permits confirmation or correction of the proposed barangay, and excludes continuous background tracking.")

subheading(doc, "Software Quality Evaluation")
body(doc, "ISO/IEC 25010:2023 defines a product-quality model for specifying and evaluating software quality. It provides nine characteristics that can guide the eventual FloodSense evaluation. The standard supplies an evaluation framework; it does not provide FloodSense scores. Quality findings will be reported only after an approved instrument has been administered and valid responses have been analyzed.")

heading(doc, "Related Local Literature")
subheading(doc, "Flood Context and Preparedness in the Philippines")
body(doc, "The Philippines experiences repeated losses from extreme events and disasters, making accessible preparedness information a continuing public need (Philippine Statistics Authority, 2026). In Cavite, the Japan International Cooperation Agency (2021) reported the inauguration of flood-protection facilities in the Imus River Basin intended to reduce floodwater affecting low-lying portions of Bacoor and Imus. These sources establish the local relevance of flood preparedness, but they do not supply operational FloodSense susceptibility parameters.")

subheading(doc, "Responsible Use of Official and Local Data")
body(doc, "National and local information must be used within the authority and limits of its source. The FloodSense research records identify BDRRMO and CPDCO as relevant institutional sources for flood, preparedness, and geographic information. However, a record's existence, a proposal, a technical geometry check, or inclusion in the repository does not establish institutional approval for operational assessment. The system must retain provenance and avoid presenting provisional material as official Bacoor City information.")

subheading(doc, "Data Privacy")
body(doc, "Republic Act No. 10173, or the Data Privacy Act of 2012, establishes principles for the lawful and proportionate processing of personal information. For FloodSense, this supports purpose-limited location handling, authenticated administrative access, restricted permissions, and avoidance of unnecessary background collection. Privacy protections must be verified in the implemented flow and not assumed solely from a design document.")

heading(doc, "Related Studies")
body(doc, "Recent studies collectively support several aspects of FloodSense. Susceptibility studies demonstrate the need for defensible factors, adequate geographic resolution, and validation; interactive mapping studies show the value of understandable spatial communication; decision-support studies show how structured information can assist without replacing human judgment; and explainability studies support traceable output. None of these studies provides a ready-made rule set, parameter table, or official Bacoor City dataset for FloodSense.")
body(doc, "Local research and institutional records identify a need for location-specific information and preparedness support. The completed initial resident survey described in Chapter 1 and the available key-informant records are requirements evidence. They are not substitutes for validation of the assessment knowledge, the geographic data, or the completed software. The newer survey package contains an instrument but no validated responses, so no result is inferred from it.")

heading(doc, "Synthesis of the Reviewed Literature")
body(doc, "The literature establishes six design implications. FloodSense must: use an explicitly confirmed scenario; keep deterministic susceptibility assessment separate from DSS guidance; explain results in understandable terms; use GIS layers according to their actual provenance and meaning; collect foreground location only when requested; and evaluate software quality with a valid instrument after implementation is sufficiently complete. The current prototype demonstrates the technical flow, but operational use depends on approved susceptibility data, verified parameters, complete integration, and evaluation evidence.")

p = doc.add_paragraph()
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.page_break_before = True
set_run(p.add_run("Table 1. Comparative Synthesis of Selected Literature"), italic=True)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(1.35), Inches(1.2), Inches(1.45), Inches(1.75)]
headers = ["SOURCE", "FOCUS", "CONTRIBUTION", "FLOODSENSE BOUNDARY"]
for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
    cell.width = widths[i]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run(p.add_run(text), size=9, bold=True)

rows = [
    ("Membele et al. (2022)", "Mapping approaches", "Highlights data and method limitations", "Does not supply Bacoor parameters"),
    ("Oubennaceur et al. (2021)", "Interactive flood communication", "Supports understandable spatial presentation", "Does not validate FloodSense outputs"),
    ("Alabbad et al. (2022)", "Flood decision support", "Supports scenario and decision-support separation", "Does not authorize warnings or orders"),
    ("Kostopoulos et al. (2024)", "Explainable DSS", "Supports transparency and understandable reasons", "Does not require ordinary-admin rule editing"),
    ("Tullis & Kar (2021)", "GIS provenance", "Supports traceability and responsible reuse", "Technical checks do not equal institutional approval"),
    ("ISO/IEC 25010:2023", "Software product quality", "Provides an evaluation framework", "Provides no FloodSense evaluation result"),
]
for row in rows:
    cells = table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, row)):
        cell.width = widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), size=9)


# CHAPTER 3
chapter(doc, "METHODOLOGY")
body(doc, "This chapter presents the research and development method for FloodSense, the current system design supported by the repository, the completed requirements-gathering activities, and the procedures planned for testing and evaluation. Completed functions are described in the present or past tense. Ongoing, planned, confirmed but unimplemented, and data-dependent work is identified accordingly. No survey result, software-quality score, acceptance result, or statistical finding is reported without completed and validated evidence.")

heading(doc, "Research Design")
body(doc, "The study uses a developmental research design to produce and examine a mobile and web-based information system. Design Thinking organizes the research activities into Empathize, Define, Ideate, Prototype, and Test. The approach links resident information needs and institutional evidence to an iterative prototype while allowing requirements, interfaces, and validation boundaries to be refined as evidence becomes available.")
body(doc, "The initial Empathize activity included the survey of 50 residents from 12 Bacoor City barangays and key-informant interviews described in Chapter 1. These completed activities are used as requirements evidence. The newer survey package is treated only as a research instrument because it does not contain validated responses or analysis.")

heading(doc, "Research Setting and Participants")
body(doc, "The research is conducted at Cavite State University - Bacoor City Campus, with Bacoor City as the geographic and institutional context of the application. The completed requirements evidence involves the 50 initial resident respondents and the available BDRRMO and CPDCO key-informant records identified in Chapter 1. Their contributions concern information needs, institutional context, and potential data sources; they do not by themselves validate every current feature or dataset.")
body(doc, "Participants for the final resident survey and system evaluation will be reported only after the sampling method, inclusion criteria, valid response count, and completed responses have been approved and verified.")
note(doc, "[AUTHOR NOTE: Survey data collection is ongoing. Complete this section after the validated survey responses and approved analysis are available.]")

heading(doc, "Research Instruments")
body(doc, "The completed requirements-gathering instruments consist of the initial resident survey and the available key-informant interview guides and records. Repository documentation, approved school documents, current source code, database migrations, interfaces, and automated tests are used as development evidence. Plans, mockups, diagrams, screenshots, and provisional datasets are used only according to their stated status and are not treated as proof of an implemented or approved feature.")
body(doc, "The newer survey questionnaire and evaluation materials remain instruments rather than results. Before use in final analysis, their validity, administration procedure, respondent eligibility, response completeness, and statistical treatment must be documented.")

heading(doc, "Development Procedure")
subheading(doc, "Empathize")
body(doc, "The researchers identify resident information problems and institutional practices through the completed initial survey, interviews, and document review. The activity focuses on access to location-specific susceptibility information, clarity and timing of preparedness information, and the institutional sources needed for responsible geographic and preparedness content.")
subheading(doc, "Define")
body(doc, "The requirements are translated into a bounded problem definition: a resident selects or confirms a hypothetical rainfall scenario and a supported location, requests an assessment, receives a deterministic susceptibility result with an explanation, and views separate pre-event preparedness guidance. Real-time monitoring, official forecasting, autonomous alerts, background tracking, automatic evacuation orders, live center occupancy, and safe-road routing are excluded.")
subheading(doc, "Ideate")
body(doc, "Alternative mobile interactions, map behaviors, result presentations, administrative workflows, and data-governance controls are represented in requirements documents and interface designs. Each proposed element is checked against the research objective, the privacy boundary, data availability, and the separation between susceptibility assessment and DSS guidance.")
subheading(doc, "Prototype")
body(doc, "The prototype is implemented through a Flutter mobile application, a Django and Django REST Framework backend, PostgreSQL/PostGIS spatial storage, and a custom web Admin portal. The prototype is refined by comparing mobile screens, routes, services, models, migrations, permissions, templates, and tests with the documented requirements.")
subheading(doc, "Test and Iterate")
body(doc, "Testing and iteration cover input validation, deterministic inference, rule matching and conflict handling, geographic lookup, API behavior, permissions, mobile interaction, custom Admin workflows, and appropriate limitation responses. Final acceptance and quality conclusions will be added only after a complete, dated test run and the approved user evaluation are available.")

note(doc, "[DIAGRAM PLACEHOLDER: Insert the approved Design Thinking or system-development process diagram here.]")

heading(doc, "System Design and Current Implementation")
subheading(doc, "Mobile Application")
body(doc, "The Flutter mobile application implements rainfall-intensity choices of Light, Moderate, Heavy, Intense, and Torrential and duration choices of 1, 3, 6, 12, and 24 hours. The user confirms the scenario before requesting an assessment. The location flow supports manual selection, a movable map pin, and an optional one-time foreground GPS request. A detected barangay is proposed for user confirmation or correction. The application does not implement continuous background tracking.")
body(doc, "The mobile map supports panning, zooming, demonstration scenario coloring, manual selection, and display of the neutral 47-barangay reference layer. The current layer is an administrative identity reference produced through a documented derivation. Its geometry has been technically checked, but it is not established as City-issued or City-verified and contains no flood-susceptibility facts.")

subheading(doc, "Backend, REST Services, and Spatial Database")
body(doc, "Django and Django REST Framework provide validation, permissions, APIs, business logic, and administrative services. GeoDjango and PostgreSQL/PostGIS support spatial records and point-to-area processing. Models and migrations provide database structure, but a model is not treated as a complete user function unless the relevant route, service, interface, integration, permissions, tests, and required data are also present.")

subheading(doc, "Deterministic Rule-Based Expert System")
body(doc, "The Expert System uses deterministic rule matching for a confirmed rainfall scenario and supported geographic context. It produces a susceptibility class and explanation. The backend includes rule matching, priorities, conflict handling, validation, and related tests. The system does not use machine-learning prediction, automatic retraining, or self-modifying rules.")

subheading(doc, "Decision Support System")
body(doc, "The DSS is maintained separately from assessment logic. Its current backend and Admin workflow manage preparedness content that can be published and returned as guidance. The DSS consumes the assessment context but does not change the susceptibility class. A complete resident-facing structured multi-question flow has not been verified and remains developing or planned.")

subheading(doc, "Custom Admin Portal and Technical Django Admin")
body(doc, "The custom Admin portal implements authentication, a dashboard, read-only map-data review, rainfall-reference review, parameter-governance views, DSS-content management, evacuation-center management, data-source management, and audit-history review. Ordinary administrators do not directly edit raw rules, rule conditions, priorities, conflict resolution, or the inference algorithm. Technical Django Admin is a separate maintenance surface for developer or research use and is not the ordinary administrator workflow.")

subheading(doc, "Fixed-Schema CSV Requirement")
body(doc, "The confirmed design requires the administrator to export a canonical CSV table with system-required column names, parameter identifiers, and data types. Authorized editors may change only permitted values within the fixed fields. Import must validate the structure, required fields, parameter identifiers, values, and data types and reject incompatible files. This workflow will not expose raw rule editing or permit algorithm changes.")
body(doc, "CSV export and import is a confirmed requirement but is not implemented in the current repository. It will remain described as future implementation until routes, services, validation, interface controls, permissions, integration, and tests are present and verified.")

subheading(doc, "Data Sources, Provenance, and Validation")
body(doc, "The backend and custom Admin portal support source records, validation status, and audit history. The approved-data directory currently contains guidance but no approved operational susceptibility dataset. Provisional data remains fictional, incomplete, unvalidated, restricted, or awaiting approval. Processing scripts demonstrate that a preparation or verification method exists, but they do not prove institutional approval or operational database use.")

note(doc, "[DIAGRAM PLACEHOLDER: Insert the approved context diagram here.]")
note(doc, "[DIAGRAM PLACEHOLDER: Insert the approved data-flow diagram here.]")
note(doc, "[DIAGRAM PLACEHOLDER: Insert the approved use-case diagram here.]")
note(doc, "[DIAGRAM PLACEHOLDER: Insert the approved entity-relationship diagram here.]")

status_heading = heading(doc, "Current Development Status")
status_heading.paragraph_format.page_break_before = True
p = doc.add_paragraph()
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(6)
set_run(p.add_run("Table 2. Summary of Current Prototype Status"), italic=True)
status_table = doc.add_table(rows=1, cols=3)
status_table.style = "Table Grid"
status_table.alignment = WD_TABLE_ALIGNMENT.CENTER
status_table.autofit = False
status_widths = [Inches(1.75), Inches(1.75), Inches(2.25)]
for i, (cell, text) in enumerate(zip(status_table.rows[0].cells, ["FUNCTION", "STATUS", "METHODOLOGICAL TREATMENT"])):
    cell.width = status_widths[i]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    set_run(p.add_run(text), size=9, bold=True)

status_rows = [
    ("Scenario selection, confirmation, deterministic assessment, explanation", "Implemented and currently verified in code", "Included in prototype testing"),
    ("Manual selection, map pin, foreground GPS, barangay confirmation", "Implemented and currently verified in code", "Verify permissions, correction, and limitation behavior"),
    ("Custom Admin dashboard and governed content workflows", "Implemented and currently verified in code", "Test role permissions and workflow validation"),
    ("Fixed-schema CSV export and import", "Confirmed requirement - not yet implemented", "Future implementation and testing"),
    ("Resident nearest-center discovery", "Not implemented", "Do not evaluate as an available resident function"),
    ("Resident accounts, assessment history, notifications, offline use", "Not implemented or outside current scope", "Exclude from current effectiveness claims"),
    ("Approved operational susceptibility data", "Blocked or awaiting data", "Demonstration output only; no official claim"),
    ("Final survey and ISO/IEC 25010 evaluation", "Ongoing or not yet conducted", "Report only after validation and analysis"),
]
for row in status_rows:
    cells = status_table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, row)):
        cell.width = status_widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(text), size=9)

heading(doc, "Testing and Verification")
body(doc, "The repository contains automated tests for backend models, services, APIs, permissions, geographic processing, Expert System behavior, custom Admin workflows, and mobile functions. These tests constitute implementation evidence when their scope and execution record are known. Historical test files alone are not reported as a current passing result.")
body(doc, "During the documentation preparation, all 117 Python source files passed a syntax-tree parse. The full Django and Flutter test suites were not rerun because the available documentation environment lacked Django, Django REST Framework, pytest, psycopg, Flutter, and Dart. Final Chapter 3 test results must identify the environment, date, commit, commands, passed and failed tests, and any unavailable dependencies.")

heading(doc, "System Evaluation Plan")
body(doc, "The final system evaluation will use criteria derived from ISO/IEC 25010:2023 and the approved institutional instrument. The selected characteristics, indicators, scale, respondent groups, administration procedure, and interpretation method will be documented before analysis. Accessibility and context-specific research questions may be reported separately when they are not direct ISO product-quality characteristics.")
note(doc, "[AUTHOR NOTE: Survey data collection is ongoing. Complete this section after the validated survey responses and approved analysis are available.]")

heading(doc, "Statistical Treatment of Data")
body(doc, "Frequencies and percentages may be used for completed categorical survey items. When an approved Likert-type evaluation instrument is available, the study may use the institutionally approved scoring and summary method. The manuscript will state the formula, scale anchors, interpretation ranges, number of valid responses, missing-data treatment, and reliability procedure actually used. No weighted mean, ranking, acceptance level, or quality score is reported at this stage.")

heading(doc, "Ethical, Privacy, and Data-Governance Considerations")
body(doc, "Research participation must be voluntary and based on the approved consent and data-handling procedure. Personally identifiable information must be limited to what is necessary, protected from unauthorized access, and retained only according to the approved research process. Application location is requested only when the user initiates the function, used for the active assessment flow, and not continuously collected in the background.")
body(doc, "Geographic, susceptibility, rainfall-reference, preparedness, and evacuation-center records must retain source and validation information. Provisional or technically derived data must be labeled accurately. Operational publication requires the appropriate institutional approval and must not be inferred from repository presence, a diagram, a mockup, or a processing script.")

heading(doc, "Methodological Limitations")
limitations = body(doc, "The current prototype demonstrates the application flow and supporting administration but does not yet have an approved operational susceptibility dataset. The 47-barangay boundary is a neutral administrative reference, not a flood-hazard layer. CSV exchange, resident nearest-center discovery, resident account flows, assessment history, notifications, offline operation, final survey analysis, full test execution in the required environments, and production deployment are incomplete or outside the current scope. These limits prevent claims of operational forecasting, official warning capability, or completed user acceptance.")
limitations.paragraph_format.keep_together = True


# REFERENCES
chapter(doc, "REFERENCES")
for item in [
    "Alabbad, Y., Yildirim, E., & Demir, I. (2022). Flood mitigation data analytics and decision support framework: Iowa Middle Cedar Watershed case study. Science of the Total Environment, 814, 152768. https://doi.org/10.1016/j.scitotenv.2021.152768",
    "Android Developers. (2026). Request location permissions. https://developer.android.com/develop/sensors-and-location/location/permissions",
    "Buchanan, B. G., & Shortliffe, E. H. (Eds.). (1984). Rule-based expert systems: The MYCIN experiments of the Stanford Heuristic Programming Project. Addison-Wesley.",
    "Debnath, J., Sahariah, D., Nath, N., Saikia, A., Lahon, D., Islam, M. N., Hashimoto, S., Meraj, G., Kumar, P., Singh, S. K., Kanga, S., & Chand, K. (2024). Modelling on assessment of flood risk susceptibility at the Jia Bharali River basin in Eastern Himalayas by integrating multicollinearity tests and geospatial techniques. Modeling Earth Systems and Environment, 10, 2393-2419. https://doi.org/10.1007/s40808-023-01912-1",
    "Fakhruddin, B., Kirsch-Wood, J., Niyogi, D., Guoqing, L., Murray, V., & Frolova, N. (2022). Harnessing risk-informed data for disaster and climate resilience. Progress in Disaster Science, 16, 100254. https://doi.org/10.1016/j.pdisas.2022.100254",
    "Fang, Z., Yue, P., Zhang, M., Xie, J., Wu, D., & Jiang, L. (2023). A service-oriented collaborative approach to disaster decision support by integrating geospatial resources and task chain. International Journal of Applied Earth Observation and Geoinformation, 117, 103217. https://doi.org/10.1016/j.jag.2023.103217",
    "International Organization for Standardization. (2023). ISO/IEC 25010:2023 systems and software engineering - Systems and software Quality Requirements and Evaluation (SQuaRE) - Product quality model. https://www.iso.org/standard/78176.html",
    "Japan International Cooperation Agency. (2021, October 1). Flood risk protection project in Cavite inaugurated. https://www.jica.go.jp/english/overseas/philippine/information/press/2021/211001.html",
    "Kostopoulos, G., Davrazos, G., & Kotsiantis, S. (2024). Explainable artificial intelligence-based decision support systems: A recent review. Electronics, 13(14), 2842. https://doi.org/10.3390/electronics13142842",
    "Longley, P. A., Goodchild, M. F., Maguire, D. J., & Rhind, D. W. (2015). Geographic information science and systems (4th ed.). Wiley.",
    "Membele, G. M., Naidu, M., & Mutanga, O. (2022). Examining flood vulnerability mapping approaches in developing countries: A scoping review. International Journal of Disaster Risk Reduction, 69, 102766. https://doi.org/10.1016/j.ijdrr.2021.102766",
    "Nguyen, H. N., Fukuda, H., & Nguyen, M. N. (2024). Assessment of the susceptibility of urban flooding using GIS with an analytical hierarchy process in Hanoi, Vietnam. Sustainability, 16(10), 3934. https://doi.org/10.3390/su16103934",
    "Oubennaceur, K., Chokmani, K., El Alem, A., & Gauthier, Y. (2021). Flood risk communication using ArcGIS StoryMaps. Hydrology, 8(4), 152. https://doi.org/10.3390/hydrology8040152",
    "Philippine Atmospheric, Geophysical and Astronomical Services Administration. (2019, January 16). Official statement of DOST-PAGASA on the heavy rainfall event associated with Tropical Depression Usman. https://www.pagasa.dost.gov.ph/article/38",
    "Philippine Statistics Authority. (2026). Component 4: Extreme events and disasters. https://psa.gov.ph/statistics/environment-statistics/highlights/component-4-extreme-events-and-disaster",
    "Power, D. J. (2002). Decision support systems: Concepts and resources for managers. Quorum Books.",
    "Republic Act No. 10173. (2012). Data Privacy Act of 2012. Republic of the Philippines. https://privacy.gov.ph/data-privacy-act/",
    "Tullis, J. A., & Kar, B. (2021). Where is the provenance? Ethical replicability and reproducibility in GIScience and its critical applications. Annals of the American Association of Geographers, 111(5), 1318-1328. https://doi.org/10.1080/24694452.2020.1806029",
]:
    reference(doc, item)

# Compact table cell margins and keep all generated table text in Arial.
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for edge in ("top", "left", "bottom", "right"):
                node = tc_mar.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    tc_mar.append(node)
                node.set(qn("w:w"), "80")
                node.set(qn("w:type"), "dxa")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
