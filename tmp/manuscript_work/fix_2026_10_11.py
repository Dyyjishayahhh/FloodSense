from copy import deepcopy
from pathlib import Path

from docx import Document


docx_path = Path(
    r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense\1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11\FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"
)
doc = Document(docx_path)

scope_index = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "Scope and Limitation of the Study")
scope = doc.paragraphs[scope_index + 1]
plain = next(p for p in doc.paragraphs if p.text.startswith("The study is scheduled from September 2026"))
if not scope.runs or not plain.runs:
    raise RuntimeError("Expected formatted runs were not found")
scope_run = scope.runs[0]._r
if scope_run.rPr is not None:
    scope_run.remove(scope_run.rPr)
if plain.runs[0]._r.rPr is not None:
    scope_run.insert(0, deepcopy(plain.runs[0]._r.rPr))

old = "GeoMapperPH supports governed collection of hazard, exposure, vulnerability, and coping-capacity information by trained government and agency users, while PlanSmartPH Ready to Rebuild organizes information for local-government rehabilitation and recovery planning (DOST-PHIVOLCS, 2022a, 2022b). These platforms illustrate structured data governance and institutional workflows. They are discussed as existing Philippine systems, not retained as Related Studies, because the inspected official pages do not provide a complete research methodology and results section suitable for the TA-required study comparison."
new = "GeoMapperPH supports governed collection of hazard, exposure, vulnerability, and coping-capacity information by trained government and agency users, while PlanSmartPH Ready to Rebuild organizes information for local-government rehabilitation and recovery planning (DOST-PHIVOLCS, 2022a, 2022b). These platforms illustrate structured data governance and institutional workflows. They are discussed as existing Philippine systems, not retained as Related Studies, because the inspected official pages do not provide a complete research methodology and results section suitable for the comparative study tables."
matches = [p for p in doc.paragraphs if p.text == old]
if len(matches) != 1:
    raise RuntimeError(f"Expected one platform paragraph; found {len(matches)}")
matches[0].runs[0].text = new
for run in matches[0].runs[1:]:
    run._element.getparent().remove(run._element)

doc.save(docx_path)
print(f"Corrected formatting and formal wording in {docx_path}")
