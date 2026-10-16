from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


root = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
path = root / "1. FLOODSENSE-MAIN" / "WORKING_REVISION_2026-10-11" / "FloodSense_Chapters_1_to_3_Introduction_Chapter2_Working_Revision_2026-10-12.docx"
doc = Document(path)

for paragraph in doc.paragraphs:
    if paragraph.text.startswith("The principal study period began in September 2026"):
        paragraph.text = (
            "The principal study period began in September 2026 and is expected to continue through May 2027. "
            "The initial resident needs survey was completed on October 11, 2026, with 40 qualified respondents "
            "from 14 Bacoor City barangays."
        )
    elif paragraph.text.startswith("Authorized administrators use the web system to review and manage rainfall and map references"):
        paragraph.text = (
            "Authorized administrators review rainfall and map references, parameters, preparedness content, "
            "evacuation-center records, data sources, and audits through the web system. They cannot edit raw "
            "rules, priorities, conflict-resolution logic, or algorithm code. Location access is user-initiated; "
            "geographic records require documented sources and validation before authoritative use."
        )
    elif paragraph.text.startswith("Table 3 shows that the reviewed systems primarily emphasize"):
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
        paragraph.paragraph_format.line_spacing = 1.8
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif paragraph.text.startswith("However, FloodSense is limited to scenario-based assessment"):
        paragraph.paragraph_format.keep_together = True

doc.tables[0].rows[2].cells[3].text = (
    "The study evaluated A4alerts' usefulness, suitability, and potential adoption among users."
)

widths = [1.33, 0.90, 1.22, 1.00, 1.25]
for table in doc.tables[:4]:
    assert len(table.columns) == 5
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
    for row_index, row in enumerate(table.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5 if row_index == 0 else 10)
    for cell in table.rows[-1].cells:
        tcpr = cell._tc.get_or_add_tcPr()
        borders = tcpr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcpr.append(borders)
        bottom = borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "14")
        bottom.set(qn("w:color"), "000000")

# The compact matrix font keeps feature names readable without changing its
# SURECUT-like horizontal-rule presentation.
for table in doc.tables[4:5]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
    for cell in table.rows[-1].cells:
        tcpr = cell._tc.get_or_add_tcPr()
        borders = tcpr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tcpr.append(borders)
        bottom = borders.find(qn("w:bottom"))
        if bottom is None:
            bottom = OxmlElement("w:bottom")
            borders.append(bottom)
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "14")
        bottom.set(qn("w:color"), "000000")

doc.save(path)
print(path)
