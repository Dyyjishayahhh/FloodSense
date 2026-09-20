from pathlib import Path

import pdfplumber
from docx import Document


files = [
    Path("1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Chapter_1_REvision_Version_1.docx"),
    Path("1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Revised_Feature_Specifications_SRS.docx"),
    Path("1. FLOODSENSE-MAIN/MAIN FILESS/FloodSense_Revised_Theoretical_Framework_and_System_Architecture_Vertical_Format.docx"),
    Path("2. FLOODSENSE-CODE AND FILES/FILES/data gathered/Flood 1st Survey Summary.pdf"),
    Path("2. FLOODSENSE-CODE AND FILES/FILES/data gathered/FloodSense 2nd Survey AnalysisReport (not reflected in chapter 1).pdf"),
    Path("2. FLOODSENSE-CODE AND FILES/FILES/data gathered/BDRRMO_ 1st INTERVIEW TRANSCRIPTION.pdf"),
    Path("2. FLOODSENSE-CODE AND FILES/FILES/data gathered/BDRRMO_2nd interview-TRANSCRIPTION (not reflected in chapter 1).pdf"),
    Path("2. FLOODSENSE-CODE AND FILES/FILES/data gathered/CPDO_1st Interview Transcript.pdf"),
    Path("2. FLOODSENSE-CODE AND FILES/FILES/data gathered/TA consultation Transciption.pdf"),
]

out_dir = Path("tmp/manuscript_work/extracted")
out_dir.mkdir(parents=True, exist_ok=True)

for source in files:
    output = out_dir / f"{source.stem}.txt"
    if source.suffix.lower() == ".docx":
        doc = Document(source)
        with output.open("w", encoding="utf-8") as stream:
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    stream.write(text + "\n")
            for t_index, table in enumerate(doc.tables, start=1):
                stream.write(f"\n[TABLE {t_index}]\n")
                for row in table.rows:
                    stream.write(" | ".join(cell.text.strip() for cell in row.cells) + "\n")
    else:
        with pdfplumber.open(source) as pdf, output.open("w", encoding="utf-8") as stream:
            for index, page in enumerate(pdf.pages, start=1):
                stream.write(f"\n===== PAGE {index} =====\n")
                stream.write(page.extract_text() or "")
                stream.write("\n")
    print(f"{source} -> {output}")
