from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from hashlib import sha256
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
SOURCE = ROOT / "1. FLOODSENSE-MAIN" / "MAIN FILESS" / "FloodSense_Chapters_1_to_2_DRAFT.docx"
OUTPUT = ROOT / "1. FLOODSENSE-MAIN" / "MAIN FILESS" / "FloodSense_Chapters_1_and_2_SURECUT_Table_Aligned_2026-09-21.docx"

# This script may be rerun only to improve the output created in this task.

source_digest = sha256(SOURCE.read_bytes()).hexdigest()
doc = Document(SOURCE)
original_chapter1 = [p._p.xml for p in doc.paragraphs[:150]]

reference_heading = next(p for p in doc.paragraphs if p.text.strip() == "REFERENCES")
section_model = next(p for p in doc.paragraphs if p.text.strip() == "Related Local Studies")
body_model = doc.paragraphs[211]
reference_model = doc.paragraphs[269]


def add_before_reference(element):
    reference_heading._p.addprevious(element)


def insert_paragraph(text: str, *, role: str = "body", page_break: bool = False):
    p = doc.add_paragraph()
    model = section_model if role == "section" else body_model
    p.style = model.style
    if model._p.pPr is not None:
        if p._p.pPr is not None:
            p._p.remove(p._p.pPr)
        p._p.insert(0, deepcopy(model._p.pPr))
    run = p.add_run(text)
    run.bold = role == "section"
    if role == "section":
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.page_break_before = page_break
    add_before_reference(p._p)
    return p


def insert_caption(text: str, *, page_break: bool = False):
    p = doc.add_paragraph()
    p.style = body_model.style
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    add_before_reference(p._p)
    return p


def set_cell_margins(cell, top=60, start=55, bottom=60, end=55):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_line_border(parent, side: str, size: int):
    pr = parent
    borders = pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.append(borders)
    node = borders.find(qn("w:" + side))
    if node is None:
        node = OxmlElement("w:" + side)
        borders.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:color"), "000000")


def set_cell_bottom_border(cell, size=4):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        pr.append(borders)
    node = OxmlElement("w:bottom")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:color"), "000000")
    borders.append(node)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def style_table(table, widths, *, font_size=9.5):
    table.autofit = False
    table.alignment = 1
    table_pr = table._tbl.tblPr
    set_line_border(table_pr, "top", 16)
    set_line_border(table_pr, "bottom", 16)
    for row in table.rows:
        no_split(row)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.left_indent = Inches(0)
                p.paragraph_format.right_indent = Inches(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    add_before_reference(table._tbl)


def make_study_table(rows):
    headers = ["SYSTEM", "AUTHOR(S) /YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"]
    table = doc.add_table(rows=1, cols=5)
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    repeat_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_bottom_border(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(table, [1.18, 0.93, 1.15, 1.24, 1.25], font_size=9.4)


foreign = [
    (
        "FloodAdapt",
        "Deltares (2025)",
        "Physics-based flood scenarios using SFINCS and Delft-FIAT impact assessment.",
        "Help communities compare possible flood events, future conditions, and adaptation choices.",
        "Produces flood and impact maps, metrics, and scenario comparisons for planning; it is not FloodSense's rule-based assessment.",
    ),
    (
        "Global Flood Awareness System (GloFAS)",
        "Copernicus Emergency Management Service (2026)",
        "Operational hydrological forecasting, observations, and web-map products.",
        "Provide global and regional flood-monitoring and forecast information.",
        "Presents time-sensitive forecast and monitoring layers; FloodSense neither monitors nor forecasts live conditions.",
    ),
    (
        "Iowa Flood Information System (IFIS)",
        "Iowa Flood Center (2024; n.d.)",
        "Stream-stage sensors, inundation maps, and model-based forecasts presented online.",
        "Make Iowa flood information understandable and accessible to the public.",
        "Combines location-oriented maps with real-time data and forecasts; its Iowa data cannot be transferred to Bacoor.",
    ),
    (
        "Resilience Analysis and Planning Tool (RAPT)",
        "Federal Emergency Management Agency (2025)",
        "GIS integration of hazard, infrastructure, and community indicators.",
        "Support emergency managers and partners in resilience analysis and planning.",
        "Enables map-layer review and contextual comparisons; it does not make FloodSense's scenario-based rule conclusion.",
    ),
    (
        "Check the long term flood risk for an area in England",
        "GOV.UK (n.d.); Environment Agency (2024)",
        "Official area lookup based on national long-term flood-risk datasets.",
        "Explain area-level risk from rivers, the sea, surface water, and other available sources.",
        "Provides public place-based risk information and limitations; an area result is not an individual-property prediction.",
    ),
]

local = [
    (
        "Know Your Hazards and NOAH Studio",
        "UP NOAH Center (n.d.)",
        "Point lookup and map layers from simulated, satellite, and historical hazard information.",
        "Allow Philippine users to inspect mapped flood and other hazard conditions.",
        "Offers location-based hazard views and return-period flood layers, with beta and data-limit notices; not a Bacoor rule assessment.",
    ),
    (
        "HazardHunterPH",
        "GeoRisk Philippines (n.d.)",
        "Point assessment using hazard layers supplied by mandated Philippine agencies.",
        "Give users an indicative, location-specific, multi-hazard report.",
        "Combines official-source layers with explicit method and boundary disclaimers; its output is not an official certification.",
    ),
    (
        "GeoMapperPH",
        "DOST-PHIVOLCS (2022a)",
        "Structured web and mobile collection of hazard, exposure, vulnerability, and coping-capacity data.",
        "Help trained agencies and local governments build exposure geodatabases.",
        "Illustrates governed field-data collection and administrative roles, rather than resident-requested susceptibility assessment.",
    ),
    (
        "PlanSmartPH Ready to Rebuild",
        "DOST-PHIVOLCS (2022b)",
        "Baseline-data assessment and web-supported rehabilitation-and-recovery planning.",
        "Assist technical staff and local governments in preparing recovery plans.",
        "Supports institutional post-disaster planning, not household pre-event scenario assessment.",
    ),
    (
        "Web-Based Solution for Flood Warning Decision Support in the Province of Leyte, Philippines",
        "Bentoso et al. (2021)",
        "Rainfall and water-level sensors, web reports, graphical monitoring, and alert messages.",
        "Monitor local flood conditions and support warning decisions in Leyte.",
        "Reports a tested sensor-based warning prototype; its thresholds and deployment evidence do not transfer to Bacoor.",
    ),
]


insert_paragraph("Synthesis of the Reviewed Literature", role="section")
insert_paragraph(
    "The reviewed literature separates susceptibility from hazard, exposure, vulnerability, and risk. "
    "Membele et al. (2022) show why mapping concepts and validation methods must be identified before a result is communicated. "
    "Mostafiz et al. (2022) add that a technically informative map is less useful to residents when it does not connect its result to understandable action. "
    "Together, these works support FloodSense's separation of a user-requested susceptibility assessment from source-backed preparedness guidance."
)
insert_paragraph(
    "The reviewed systems demonstrate different forms of geographic decision support. Deltares (2025) uses model-based what-if scenarios in FloodAdapt, while the Copernicus Emergency Management Service (2026) and Iowa Flood Center (2024) provide operational forecasts or monitoring information. "
    "In the Philippines, the UP NOAH Center (n.d.) and GeoRisk Philippines (n.d.) present location-based hazard information, whereas DOST-PHIVOLCS (2022a, 2022b) documents agency-oriented data collection and recovery planning. "
    "Bentoso et al. (2021) instead investigated sensor-based warning in Leyte. Their data, geographic scale, intended users, and authority differ from those of FloodSense; their thresholds and findings cannot simply be applied to Bacoor City."
)
insert_paragraph(
    "Tullis and Kar (2021) emphasize provenance in geographic work, and the Philippine platforms' own source and boundary notices illustrate why technical processing must be distinguished from institutional validation. "
    "For FloodSense, a 47-barangay administrative reference layer cannot by itself establish flood susceptibility. "
    "Deterministic rule matching may make a result repeatable and explainable, but its public validity still depends on locally appropriate data, documented rules, expert review, and authorized use."
)
insert_paragraph("Research Gap", role="section")
insert_paragraph(
    "Existing platforms already provide scenario simulation, official hazard-layer lookup, real-time monitoring, forecasts, and institutional planning. "
    "The reviewed evidence does not establish that these distinct functions are routinely combined for Bacoor residents in a single flow that starts with a user-confirmed hypothetical rainfall scenario, resolves a supported location, applies deterministic and explainable rules, and then presents separate pre-event preparedness guidance. "
    "This is a context-specific design gap, not a claim that FloodSense is the first flood application."
)
insert_paragraph(
    "The gap also concerns evidence and authority. Bacoor-specific susceptibility inputs, rule thresholds, geographic layers, and guidance require source documentation and local expert or institutional validation before public conclusions can be treated as operational. "
    "FloodSense therefore remains a prototype for scenario-based assessment and preparedness support, not a substitute for official forecasting, flood warnings, evacuation orders, or verified hazard certification."
)
insert_paragraph("Comparative Synthesis of Related Studies", role="section")
insert_paragraph(
    "The following comparisons summarize the documented approaches and contributions of the foreign and Philippine systems discussed above. "
    "Official platforms are identified by their responsible institutions; the methodology column describes the documented technical approach and does not imply that every platform is a peer-reviewed research study."
)
insert_paragraph(
    "Table 1 compares the related foreign systems by responsible institution, documented approach, objective, and contribution."
)
insert_caption("Table 1. Comparison of related foreign studies")
make_study_table(foreign[:2])
insert_caption("Table 1. Continued", page_break=True)
make_study_table(foreign[2:])
insert_paragraph(
    "Table 2 compares the related Philippine systems using the same categories."
)
insert_caption("Table 2. Comparison of related local studies")
make_study_table(local[:2])
insert_caption("Table 2. Continued", page_break=True)
make_study_table(local[2:])
insert_paragraph(
    "Table 3 identifies only features documented for each reviewed system. A blank cell means that the feature was not established by the inspected source; it does not prove that the system lacks the feature. "
    "The Proposed System column describes the verified FloodSense prototype, not unimplemented requirements or official flood-data approval."
)
insert_caption("Table 3. Comparison of developed study to the other existing systems.", page_break=True)

features = [
    ("Selectable flood scenario", {1, 6, 11}),
    ("Interactive geographic display", {1, 2, 3, 4, 6, 7, 8, 9, 11}),
    ("Public location-based lookup", {3, 5, 6, 7, 11}),
    ("Real-time monitoring or forecast", {2, 3, 10}),
    ("Rule-based susceptibility output", {11}),
    ("Pre-event preparedness guidance", {5, 11}),
    ("Governed administrative data workflow", {8, 9, 11}),
    ("User-initiated foreground location", {11}),
]

matrix = doc.add_table(rows=2 + len(features), cols=12)
matrix.cell(0, 0).text = "Features"
matrix.cell(0, 1).merge(matrix.cell(0, 10)).text = "Existing System"
matrix.cell(0, 11).text = "Proposed\nSystem"
for i in range(10):
    matrix.cell(1, i + 1).text = f"S{i+1}"
for ri, (feature, checked) in enumerate(features, start=2):
    matrix.cell(ri, 0).text = feature
    for i in range(1, 12):
        matrix.cell(ri, i).text = "✔" if i in checked else ""
for row in matrix.rows[:2]:
    repeat_header(row)
for cell in matrix.rows[1].cells:
    set_cell_bottom_border(cell)
style_table(matrix, [1.50] + [0.34] * 10 + [0.85], font_size=8.6)
for row in matrix.rows[:2]:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
for row in matrix.rows[2:]:
    for cell in row.cells[1:]:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                if run.text == "✔":
                    run.font.name = "Segoe UI Symbol"

legend = insert_paragraph("Legend:", role="section", page_break=True)
legend.paragraph_format.space_before = Pt(0)
insert_paragraph("✔ - included in the documented system; blank - not established by the inspected source.")
legend_names = [
    "FloodAdapt",
    "Global Flood Awareness System (GloFAS)",
    "Iowa Flood Information System (IFIS)",
    "Resilience Analysis and Planning Tool (RAPT)",
    "Check the long term flood risk for an area in England",
    "Know Your Hazards and NOAH Studio (UP NOAH)",
    "HazardHunterPH",
    "GeoMapperPH",
    "PlanSmartPH Ready to Rebuild",
    "Web-Based Solution for Flood Warning Decision Support in the Province of Leyte, Philippines",
]
for i, name in enumerate(legend_names, 1):
    p = insert_paragraph(f"S{i} - {name}")
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.space_after = Pt(4)

# The public service is distinct from the Environment Agency's national assessment.
service_p = next(p for p in doc.paragraphs if p.text.startswith("The Environment Agency (2024) supports England"))
service_p.text = service_p.text.replace(
    "while the public Check the Long Term Flood Risk for an Area service allows",
    "while GOV.UK (n.d.) offers the public Check the Long Term Flood Risk for an Area service, which allows",
)

# Retain only entries actually cited by Chapters 1 and 2, including the new tables.
unused = {
    "Alabbad", "Cayamanda", "Dasallas", "Debnath", "Dulawan", "Fang", "Gacul",
    "Gumba", "Johnson", "Longley", "Macalalad", "Mercado", "Nguyen", "Oubennaceur", "Williams",
}
removed = []
for p in list(doc.paragraphs):
    if p._p.getparent() is None:
        continue
    first = re.match(r"^([^,.(]+)", p.text)
    if first and first.group(1) in unused and p.text.startswith(first.group(1) + ","):
        removed.append(first.group(1))
        p._p.getparent().remove(p._p)

next_reference = next(p for p in doc.paragraphs if p.text.startswith("International Organization for Standardization."))
gov_ref = doc.add_paragraph()
gov_ref.style = reference_model.style
if reference_model._p.pPr is not None:
    if gov_ref._p.pPr is not None:
        gov_ref._p.remove(gov_ref._p.pPr)
    gov_ref._p.insert(0, deepcopy(reference_model._p.pPr))
gov_ref.add_run(
    "GOV.UK. (n.d.). Check the long term flood risk for an area in England. "
    "https://www.gov.uk/check-long-term-flood-risk"
)
next_reference._p.addprevious(gov_ref._p)

# The source draft ends in an empty additional section. Keep the reference
# section's properties while removing only this output's trailing blank page.
tail = list(doc.paragraphs)[-3:]
assert all(not p.text.strip() for p in tail)
closing_section = tail[1]._p.xpath("./w:pPr/w:sectPr")
assert len(closing_section) == 1
body = doc._element.body
body.replace(body.sectPr, deepcopy(closing_section[0]))
for p in tail:
    p._p.getparent().remove(p._p)

doc.save(OUTPUT)
reloaded = Document(OUTPUT)
chapter1 = [p._p for p in reloaded.paragraphs[:150]]
assert len(chapter1) == len(original_chapter1)
assert all(a == b.xml for a, b in zip(original_chapter1, chapter1)), "Chapter 1 XML changed"
assert sha256(SOURCE.read_bytes()).hexdigest() == source_digest, "Source manuscript changed"
texts = [p.text for p in reloaded.paragraphs]
assert not any(re.search(r"\bCHAPTER\s*(?:3|III)\b|^METHODOLOGY$", t, re.I) for t in texts)
assert texts[-1].startswith("UP NOAH Center."), "Document does not end after references"
print("OUTPUT", OUTPUT)
print("SOURCE_SHA256", source_digest)
print("CHAPTER1_PARAGRAPHS_PRESERVED", len(chapter1))
print("TABLES", len(reloaded.tables))
print("REMOVED_UNUSED_REFERENCE_ENTRIES", sorted(removed))
print("OUTPUT_SHA256", sha256(OUTPUT.read_bytes()).hexdigest())
