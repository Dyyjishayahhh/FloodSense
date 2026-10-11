from pathlib import Path
import sys

from docx import Document
from pypdf import PdfReader


def extract_docx(path: Path) -> str:
    doc = Document(path)
    out = [f"SOURCE: {path}", f"PARAGRAPHS: {len(doc.paragraphs)}", f"TABLES: {len(doc.tables)}", ""]
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.replace("\t", " ").strip()
        if text:
            style_name = paragraph.style.name if paragraph.style is not None else "No style"
            out.append(f"P{index:04d} [{style_name}] {text}")
    for table_index, table in enumerate(doc.tables):
        out.append("")
        out.append(f"TABLE {table_index + 1} ({len(table.rows)}x{len(table.columns)})")
        for row_index, row in enumerate(table.rows):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            out.append(f"R{row_index:03d}\t" + "\t".join(cells))
    return "\n".join(out)


def extract_pdf(path: Path) -> str:
    reader = PdfReader(path)
    out = [f"SOURCE: {path}", f"PAGES: {len(reader.pages)}", ""]
    for index, page in enumerate(reader.pages, start=1):
        out.append(f"===== PAGE {index} =====")
        out.append(page.extract_text() or "")
    return "\n".join(out)


def main() -> None:
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    destination.parent.mkdir(parents=True, exist_ok=True)
    if source.suffix.lower() == ".docx":
        text = extract_docx(source)
    elif source.suffix.lower() == ".pdf":
        text = extract_pdf(source)
    else:
        text = source.read_text(encoding="utf-8", errors="replace")
    destination.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    main()
