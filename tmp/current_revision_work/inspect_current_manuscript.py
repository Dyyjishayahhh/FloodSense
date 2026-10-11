from pathlib import Path
import sys

from docx import Document


PATH = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense\1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11\FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx")

doc = Document(PATH)
sys.stdout.reconfigure(encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} shapes={len(doc.inline_shapes)} sections={len(doc.sections)}")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip().replace("\t", " ")
    if text:
        print(f"P{i:04d}\t{p.style.name}\t{text}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        values = [c.text.strip().replace("\n", " / ") for c in row.cells]
        print(f"T{ti}R{ri}\t" + " || ".join(values))
