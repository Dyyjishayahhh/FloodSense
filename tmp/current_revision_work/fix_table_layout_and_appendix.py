from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
DOCX = ROOT / r"1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11\FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"


def set_widths(table, widths: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
        for cell in table.columns[index].cells:
            cell.width = Inches(width)
            tcw = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tcw is not None:
                tcw.set(qn("w:w"), str(round(width * 1440)))
                tcw.set(qn("w:type"), "dxa")


def set_cell_text(cell, value: str, *, size: float, bold: bool = False, center: bool = False, symbol: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Segoe UI Symbol" if symbol else "Arial"
    run.font.size = Pt(size)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def no_split_and_repeat(table) -> None:
    for index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if not any(child.tag == qn("w:cantSplit") for child in trpr):
            trpr.append(OxmlElement("w:cantSplit"))
        if index == 0 and not any(child.tag == qn("w:tblHeader") for child in trpr):
            trpr.append(OxmlElement("w:tblHeader"))


doc = Document(DOCX)
assert len(doc.tables) == 6, len(doc.tables)

# Combine the two pairs of manually split study tables. Word can then paginate
# them naturally while repeating their header row and keeping each study intact.
original_tables = list(doc.tables)
for first_index, second_index in [(0, 1), (2, 3)]:
    first, second = original_tables[first_index], original_tables[second_index]
    for row in list(second.rows)[1:]:
        first._tbl.append(deepcopy(row._tr))
    second._tbl.getparent().remove(second._tbl)

for label in ("Table 1. Continued", "Table 2. Continued"):
    paragraph = next(p for p in doc.paragraphs if p.text.strip() == label)
    paragraph._p.getparent().remove(paragraph._p)

for label in ("Table 1. Comparison of related foreign studies", "Table 2. Comparison of related local studies"):
    paragraph = next(p for p in doc.paragraphs if p.text.strip() == label)
    paragraph.paragraph_format.page_break_before = False
    paragraph.paragraph_format.keep_with_next = True

assert len(doc.tables) == 4, len(doc.tables)
for table in doc.tables[:2]:
    set_widths(table, [1.40, 0.95, 1.20, 1.10, 1.05])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            value = cell.text.strip()
            set_cell_text(cell, value, size=7.8 if row_index == 0 else 8.2, bold=row_index == 0, center=row_index == 0 or column_index == 1)
    no_split_and_repeat(table)

matrix = doc.tables[2]
assert len(matrix.columns) == 10 and len(matrix.rows) in (11, 12), (len(matrix.columns), len(matrix.rows))
matrix_labels = ["FEATURES", "S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8", "FloodSense"]
for column_index, label in enumerate(matrix_labels):
    set_cell_text(matrix.cell(0, column_index), label, size=7.0, bold=True, center=True)
if len(matrix.rows) == 12:
    matrix._tbl.remove(matrix.rows[1]._tr)
set_widths(matrix, [2.05] + [0.40] * 8 + [0.45])
for row_index, row in enumerate(matrix.rows):
    if row_index == 0:
        continue
    for column_index, cell in enumerate(row.cells):
        value = cell.text.strip()
        if value == "✔":
            value = "✓"
        set_cell_text(cell, value, size=8.0 if column_index == 0 else 8.5, center=column_index != 0, symbol=value == "✓")
no_split_and_repeat(matrix)

appendices_heading = next(p for p in doc.paragraphs if p.text.strip() == "APPENDICES")
appendix_a = next(p for p in doc.paragraphs if p.text.strip() == "Appendix A. FloodSense Fishbone Diagram")
appendices_heading.paragraph_format.page_break_before = True
appendices_heading.paragraph_format.keep_with_next = True
appendix_a.paragraph_format.page_break_before = False
appendix_a.paragraph_format.keep_with_next = True

doc.save(DOCX)
print(DOCX)
print(f"tables={len(doc.tables)}; foreign/local rows={[len(t.rows) for t in doc.tables[:2]]}; matrix rows={len(matrix.rows)}")
