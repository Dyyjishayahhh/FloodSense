from pathlib import Path

from docx import Document

root = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
path = root / r"1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11\FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"
doc = Document(path)
old_sentence = "Bentoso et al. (2021) developed an automated flood-warning prototype for the Province of Leyte using rainfall and water-level sensors, data transmission, a web interface, graphical reporting, threshold-based alerts, and decision-support outputs."
new_sentence = "Bentoso et al. (2021) described an automated flood-warning prototype for the Province of Leyte using rainfall and water-level sensors, data transmission, a web interface, graphical reporting, alert messages, and decision-support information."
matches = [p for p in doc.paragraphs if old_sentence in p.text]
assert len(matches) == 1, len(matches)
paragraph = matches[0]
assert len(paragraph.runs) == 1, len(paragraph.runs)
paragraph.runs[0].text = paragraph.runs[0].text.replace(old_sentence, new_sentence)
method_cell = doc.tables[1].cell(1, 2)
assert "threshold alerts" in method_cell.text
run = method_cell.paragraphs[0].runs[0]
run.text = "Rainfall and water-level readings transmitted to a web system, converted to graphical reports, and used to send alert messages."
doc.save(path)
print(path)
