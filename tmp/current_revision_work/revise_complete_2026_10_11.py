from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
DOCX = ROOT / r"1. FLOODSENSE-MAIN\WORKING_REVISION_2026-10-11\FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"
FRAMEWORK = ROOT / r"tmp\current_revision_work\FloodSense_Theoretical_Framework.png"
DIAGRAM_DIR = ROOT / r"tmp\diagram_review_2026_10_11\pages"


def find_paragraph(doc: Document, exact: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if p.text.strip() == exact]
    if len(matches) != 1:
        raise RuntimeError(f"Expected exactly one paragraph {exact!r}; found {len(matches)}")
    return matches[0]


def replace_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def clone_before(anchor: Paragraph, template: Paragraph, text: str, *, bold: bool | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    if template._p.pPr is not None:
        element.append(deepcopy(template._p.pPr))
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    run = paragraph.add_run(text)
    if template.runs and template.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(template.runs[0]._r.rPr))
    if bold is not None:
        run.bold = bold
    return paragraph


def clear_between(start: Paragraph, end: Paragraph) -> None:
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def clear_after(heading: Paragraph) -> None:
    node = heading._p.getnext()
    while node is not None:
        nxt = node.getnext()
        if node.tag == qn("w:sectPr"):
            break
        node.getparent().remove(node)
        node = nxt


def add_bullet(anchor: Paragraph, template: Paragraph, text: str) -> Paragraph:
    p = clone_before(anchor, template, f"• {text}")
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    return p


def add_numbered(anchor: Paragraph, template: Paragraph, number: int, text: str) -> Paragraph:
    p = clone_before(anchor, template, f"{number}. {text}")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    return p


def set_cell(cell, text: str, font_size: float, *, bold: bool = False, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def repeat_header_and_prevent_split(table) -> None:
    for idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if idx == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def add_table_before(anchor: Paragraph, data: list[list[str]], widths: list[float], *, font_size: float = 8.5, two_header_rows: bool = False):
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = table_style_name
    except Exception:
        table.style = "Table Grid"
    table.autofit = False
    anchor._p.addprevious(table._tbl)
    for r_idx, row_data in enumerate(data):
        for c_idx, value in enumerate(row_data):
            set_cell(
                table.cell(r_idx, c_idx),
                value,
                font_size,
                bold=r_idx == 0 or (two_header_rows and r_idx == 1),
                center=r_idx == 0 or (two_header_rows and r_idx == 1),
            )
            table.cell(r_idx, c_idx).width = Inches(widths[c_idx])
            tc_pr = table.cell(r_idx, c_idx)._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is not None:
                tc_w.set(qn("w:w"), str(int(widths[c_idx] * 1440)))
                tc_w.set(qn("w:type"), "dxa")
    repeat_header_and_prevent_split(table)
    return table


def add_picture_before(anchor: Paragraph, template: Paragraph, path: Path, width: float) -> Paragraph:
    p = clone_before(anchor, template, "")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.left_indent = None
    p.paragraph_format.right_indent = None
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def add_caption_before(anchor: Paragraph, text: str, *, page_break_before: bool = False) -> Paragraph:
    p = clone_before(anchor, caption_template, text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = page_break_before
    return p


assert DOCX.exists(), DOCX
assert FRAMEWORK.exists(), FRAMEWORK
for required in ("Fishbone_p1.png", "Flowchart_p1.png", "Flowchart_p2.png"):
    assert (DIAGRAM_DIR / required).exists(), required

doc = Document(DOCX)
preliminary_hash_before = sha256("\n".join(p.text for p in doc.paragraphs[:65]).encode("utf-8")).hexdigest()

body_template = doc.paragraphs[67]
major_heading_template = find_paragraph(doc, "Statement of the Problem")
subheading_template = find_paragraph(doc, "Flood Susceptibility and Scenario-Based Assessment")
caption_template = find_paragraph(doc, "Table 1. Comparison of related foreign studies")
reference_template = find_paragraph(doc, "Android Developers. (2026). Request location permissions. https://developer.android.com/develop/sensors-and-location/location/permissions")
table_style_name = doc.tables[0].style.name if doc.tables and doc.tables[0].style is not None else "Table Grid"

# CHAPTER 1: evidence-led context, aligned problems/objectives, restored authentication, formal scope/limitations.
intro_heading = find_paragraph(doc, "INTRODUCTION")
sop_heading = find_paragraph(doc, "Statement of the Problem")
clear_between(intro_heading, sop_heading)
for text in [
    "Flooding remains one of the most persistent hazards affecting Philippine communities. Beyond physical inundation, its effects include interrupted travel, disrupted work and schooling, household displacement, property loss, and difficult decisions about preparation and evacuation. These consequences make understandable and appropriately qualified flood information an important part of community preparedness.",
    "Bacoor City includes low-lying and river-connected communities that experience recurring flood conditions. Flood-mitigation works in the Imus River Basin were established to reduce peak floodwater volume in parts of Cavite, including low portions of Bacoor and its tributaries (Japan International Cooperation Agency [JICA], 2021). City offices also maintain historical, geographic, preparedness, and evacuation information, although the availability, authority, and permitted public use of particular records vary.",
    "The initial needs-assessment survey included 40 consenting adult Bacoor City residents from 14 barangays who had experienced flooding or serious heavy-rain effects. Thirty respondents (75.0%) reported more than five such events during the preceding three years. Reported effects included disruption of work or study (97.5%), flooding of streets or community areas (87.5%), difficult travel routes (85.0%), floodwater entering the home (77.5%), and evacuation or temporary relocation (60.0%).",
    "Access to usable information was also a recurring concern. Twenty-nine respondents (72.5%) reported difficulty obtaining needed flood information, and 30 (75.0%) stated that inadequate information often or always affected their decisions. Thirty respondents (75.0%) found it difficult, very difficult, or were unsure how to determine the susceptibility of their location, while 33 (82.5%) reported the same difficulty regarding verified evacuation-center information. The highest-rated information problems concerned information that was too general for the respondent's location, limited access when internet service or electricity was unreliable, overly general preparedness advice, and unclear source or date information.",
    "The survey also showed that 39 respondents (97.5%) had access to an Android device and that 33 (82.5%) were comfortable with one-time GPS access when checking an area. Manual barangay selection and temporary map-pin use were also acceptable to most respondents, while 75.0% preferred information presented in both English and Filipino. These findings support a user-initiated mobile flow while also emphasizing privacy, manual alternatives, and clear communication.",
    "Within this context, FloodSense is a scenario-based, pre-event flood-susceptibility assessment and preparedness-support system. It requires the resident to confirm a hypothetical rainfall scenario and location and to request an assessment explicitly. A deterministic Rule-Based Expert System produces an explained susceptibility result, while a separate Decision Support System presents preparedness guidance. FloodSense does not replace official observations, forecasts, warnings, emergency directives, or local-government authority.",
]:
    clone_before(sop_heading, body_template, text)

objectives_heading = find_paragraph(doc, "Objectives of the Study")
clear_between(sop_heading, objectives_heading)
for text in [
    "Recurring flooding in Bacoor City affects household routines, mobility, work, schooling, property, and evacuation decisions. The initial survey establishes that residents repeatedly experience these effects and that many have difficulty obtaining information that is location-specific, understandable, timely, and traceable to a source. Local institutional records and interviews further show that geographic and preparedness information must be handled according to its ownership, validation status, and permitted public use.",
    "The central research problem is therefore the limited availability of an accessible and appropriately qualified way for residents to examine how a hypothetical rainfall situation may relate to the susceptibility of a selected Bacoor City location and then obtain separate preparedness guidance. Addressing this problem requires attention to resident information needs, deterministic and explainable assessment, geographic presentation, privacy, data provenance, administrative governance, technical verification, and product-quality evaluation.",
    "Specifically, the study seeks to answer the following questions:",
]:
    clone_before(objectives_heading, body_template, text)
questions = [
    "What flood experiences, information-access difficulties, preparedness needs, and technology conditions are reported by the eligible Bacoor City respondents?",
    "What knowledge, rainfall-scenario, geographic, privacy, provenance, preparedness, evacuation-information, and administrative-governance requirements are necessary for FloodSense?",
    "How can an Android resident application and a Web Administration System be developed through an Agile Development Method to provide user access, scenario and location input, deterministic susceptibility assessment and explanation, map presentation, separate preparedness support, and governed information management?",
    "To what extent do the developed resident, administrative, geographic, Expert System, and Decision Support System functions operate correctly and consistently under defined technical test cases?",
    "How will the relevant product-quality characteristics of FloodSense be evaluated using an approved ISO/IEC 25010-based instrument and defined evaluator groups?",
]
for i, text in enumerate(questions, 1):
    add_numbered(objectives_heading, body_template, i, text)

framework_heading = find_paragraph(doc, "Theoretical Framework")
clear_between(objectives_heading, framework_heading)
clone_before(framework_heading, body_template, "Generally, the study aimed to develop and evaluate FloodSense as a scenario-based Android system for explainable flood-susceptibility assessment and separate pre-event preparedness support for Bacoor City.")
clone_before(framework_heading, body_template, "Specifically, it aimed to:")
objective_texts = [
    "identify the flood experiences, information-access difficulties, preparedness needs, and technology conditions reported by eligible Bacoor City respondents;",
    "determine the knowledge, rainfall-scenario, geographic, privacy, provenance, preparedness, evacuation-information, and administrative-governance requirements of the system;",
    "develop FloodSense through an Agile Development Method as a Flutter-based Android resident application and a Django-based Web Administration System with the following major functions:",
]
for i, text in enumerate(objective_texts, 1):
    add_numbered(framework_heading, body_template, i, text)
for text in [
    "user access, authentication, verification, recovery, privacy acceptance, profiles, roles, and permissions;",
    "user-confirmed rainfall intensity, rainfall duration, barangay, manual map pin, and optional one-time foreground location;",
    "deterministic rule matching, conflict handling, susceptibility classification, explanation, and map presentation;",
    "structured preparedness guidance and evacuation-center information kept separate from susceptibility assessment; and",
    "controlled administrative workflows for references, parameters, DSS content, evacuation-center records, sources, review records, permissions, and audit history;",
]:
    add_bullet(framework_heading, body_template, text)
add_numbered(framework_heading, body_template, 4, "test the resident, administrative, geographic, Expert System, and Decision Support System functions through unit, integration, API, spatial, and end-to-end test cases; and")
add_numbered(framework_heading, body_template, 5, "evaluate the relevant product-quality characteristics of FloodSense using an approved ISO/IEC 25010-based instrument and defined evaluator groups.")

significance_heading = find_paragraph(doc, "Significance of the Study")
clear_between(framework_heading, significance_heading)
framework_paragraphs = [
    "The theoretical framework of FloodSense connects resident access, user-confirmed inputs, deterministic inference, geographic presentation, separate preparedness guidance, and governed information management. Consistent with the study's research problems, related functions are consolidated under five function-named modules rather than being separated according to individual technologies.",
    "User Access and Authentication Module. This module controls resident and administrator identity, registration, sign-in, verification, password recovery, privacy and legal acceptance, profiles, roles, and permissions. Authentication is necessary because FloodSense processes account information, optional location data, and governed administrative records that must not be available to unauthorized users.",
    "Scenario and Location Input Module. This module receives the resident's hypothetical rainfall intensity and duration and the selected location. The resident may choose a barangay, adjust a temporary map pin, or request one-time foreground location access, then confirm or correct the resolved barangay before requesting assessment. These inputs describe a user-selected scenario rather than a live observation or forecast.",
    "Susceptibility Assessment and Explanation Module. This module validates the confirmed facts, applies deterministic rule matching and conflict handling, and returns a susceptibility class with a human-readable explanation and geographic presentation. The same approved facts and rule version produce the same outcome, while the scientific and local validity of the underlying knowledge remains dependent on appropriate data, provenance, expert review, and institutional authorization.",
    "Preparedness Decision Support Module. This module presents source-backed preparedness guidance through structured questions and information flows. It may present authorized evacuation-center information and straight-line distance as reference information. The guidance remains separate from the susceptibility assessment and does not modify the assessment class or issue an evacuation order.",
    "Administration and Data Governance Module. This module supports authenticated and permission-controlled management of rainfall and map references, assessment parameters, DSS content, evacuation-center records, data sources, review records, and audit history. Ordinary administrators do not directly modify raw rule conditions, priorities, conflict-resolution logic, or algorithm code. Technical Django Admin access remains separate and is reserved for developer or research maintenance.",
]
for text in framework_paragraphs:
    clone_before(significance_heading, body_template, text)
add_picture_before(significance_heading, body_template, FRAMEWORK, 6.35)
fig = add_caption_before(significance_heading, "Figure 1. FloodSense Theoretical Framework and System Architecture")
fig.paragraph_format.keep_with_next = True
clone_before(significance_heading, body_template, "The architecture places the resident on the left and the authorized administrator on the right. The Android application and custom Web Administration System communicate with the Django REST services that coordinate authentication, scenario and location processing, deterministic inference, geographic services, preparedness guidance, and evacuation-center information. PostgreSQL/PostGIS and controlled knowledge records support these functions as an implementation foundation rather than as a separate conceptual module.")

time_heading = find_paragraph(doc, "Time and Place of the Study")
clear_between(significance_heading, time_heading)
for text in [
    "FloodSense contributes an evidence-based approach to scenario-oriented flood preparedness by combining an explained susceptibility assessment with separate guidance and visible information boundaries.",
    "For the Bacoor Disaster Risk Reduction and Management Office. The study may provide a complementary channel for presenting validated preparedness and evacuation-center information while retaining source, review, and audit records. It does not replace official warning, response, or emergency-command functions.",
    "For the City Planning and Development Coordinator Office. The study demonstrates how appropriately authorized administrative and geographic information may support location-specific public communication while maintaining provenance and validation requirements.",
    "For Bacoor City residents. The resident application provides a user-initiated process for confirming a hypothetical rainfall scenario and location, requesting an explained susceptibility result, viewing geographic context, and obtaining structured preparedness information.",
    "For future researchers. The study documents the separation of deterministic assessment from preparedness guidance, the importance of authentication and role control, and the distinction among technical processing, scientific validation, and institutional authorization. Its procedures and findings require renewed validation before use in another locality.",
]:
    clone_before(time_heading, body_template, text)

scope_heading = find_paragraph(doc, "Scope and Limitation of the Study")
clear_between(time_heading, scope_heading)
clone_before(scope_heading, body_template, "The study is conducted from September 2026 to May 2027 in Bacoor City, Cavite. System development and manuscript preparation take place primarily at Cavite State University-Bacoor City Campus. The initial resident survey covered respondents from 14 Bacoor barangays and was completed on October 11, 2026, while institutional interviews and document requests involved the Bacoor Disaster Risk Reduction and Management Office and the City Planning and Development Coordinator Office.")

definitions_heading = find_paragraph(doc, "Definition of Terms")
clear_between(scope_heading, definitions_heading)
clone_before(definitions_heading, body_template, "FloodSense covers an Android application for Bacoor City residents and a browser-based administration system for authorized personnel. Resident users register and authenticate, accept the applicable privacy and legal notices, select or confirm a hypothetical rainfall intensity and duration, identify a location through barangay selection, a temporary map pin, or optional one-time foreground GPS, confirm or correct the resolved barangay, and explicitly request an assessment. The deterministic Rule-Based Expert System returns an explained susceptibility classification and map context, while a separate Decision Support System presents source-backed preparedness guidance and authorized evacuation-center reference information. The administration system supports authenticated dashboards and permission-controlled management of rainfall and map references, assessment parameters, preparedness content, evacuation-center records, sources, review records, and audit history. Ordinary administrators do not directly edit raw rules, rule priorities, conflict-resolution logic, or algorithm code.")
clone_before(definitions_heading, body_template, "However, the study is limited to scenario-based, pre-event assessment and preparedness support within Bacoor City's documented administrative coverage. It does not observe rainfall or water levels continuously and does not provide automatic rainfall updates, real-time flood monitoring, official forecasts or warnings, autonomous alerts, background GPS tracking, automatic evacuation orders, safe-road routing, live evacuation-center occupancy, or machine-learning prediction. The 47-barangay layer is used only as an administrative reference and does not itself contain flood-susceptibility information or establish City verification. The authority and local validity of susceptibility results, preparedness content, and evacuation-center records depend on appropriate source documentation, expert review, and institutional validation. Center distance is straight-line reference information and does not establish route safety, accessibility, opening status, available capacity, or safety. The Android application and web services require a compatible device and network access for connected functions. The 40-person purposive survey represents the participating respondents rather than the entire Bacoor City population, and personal and one-time location information remains subject to consent, purpose limitation, role-controlled access, and the Data Privacy Act of 2012.")

web_admin_def = find_paragraph(doc, "Web Administration System. The custom browser-based interface used by authorized administrators for dashboards and governed workflows involving map and rainfall references, parameters, DSS content and structured flows, evacuation centers and staged center imports, data sources, and audit history. It does not expose raw rule editing. Technical Django Admin is separate. Fixed-schema Expert System parameter CSV export/import is confirmed but not implemented.")
replace_text(web_admin_def, "Web Administration System. The custom browser-based interface used by authorized administrators for dashboards and permission-controlled workflows involving map and rainfall references, parameters, DSS content and structured flows, evacuation centers, data sources, review records, and audit history. It does not expose raw rule editing. Technical Django Admin is separate and is reserved for developer or research maintenance.")

# CHAPTER 2: formal boundary language and eight verified system studies with methods, data, results, and limitations.
old_rule_lit = find_paragraph(doc, "Deterministic processing means that the same approved facts and the same published rule version should produce the same result. It does not imply that the rules are scientifically valid merely because the software executes them consistently. Knowledge acquisition, conflict handling, parameter definitions, version control, and expert or institutional validation remain necessary. Accordingly, ordinary FloodSense administrators may manage approved records and permitted values but do not rewrite raw rule conditions, priorities, conflict-resolution logic, or algorithm code. The confirmed fixed-schema CSV exchange is a future controlled-value workflow and is not evidence that such exchange is currently operational.")
replace_text(old_rule_lit, "Deterministic processing means that the same approved facts and the same published rule version should produce the same result. It does not imply that the rules are scientifically valid merely because the software executes them consistently. Knowledge acquisition, conflict handling, parameter definitions, version control, and expert or institutional validation remain necessary. Accordingly, ordinary FloodSense administrators may manage authorized records and permitted values but do not rewrite raw rule conditions, priorities, conflict-resolution logic, or algorithm code.")
old_quality = find_paragraph(doc, "A defensible evaluation must identify the tested build, participants or evaluators, tasks, criteria, scale, procedure, valid responses, and interpretation method. Repository tests and interface demonstrations can support implementation verification, but they are not resident-acceptance findings. Because the final FloodSense survey and evaluation are not complete, Chapter 2 uses ISO/IEC 25010 only as literature supporting future assessment and does not report quality scores.")
replace_text(old_quality, "A defensible product-quality evaluation must identify the tested build, participants or evaluators, tasks, criteria, scale, procedure, valid responses, and interpretation method. Automated tests and interface demonstrations can support technical verification, but they are not user-acceptance findings. The completed initial survey is a needs assessment rather than a post-use system evaluation; therefore, Chapter 2 uses ISO/IEC 25010 as literature supporting the planned product-quality assessment and does not report FloodSense evaluation scores.")

# Move the Philippine mapping studies to local literature because they are not developed systems.
existing_platforms = find_paragraph(doc, "Existing Philippine Hazard-Information and Planning Platforms")
clone_before(existing_platforms, subheading_template, "Philippine Flood-Susceptibility and Risk Mapping")
clone_before(existing_platforms, body_template, "Puno et al. (2022) used GIS frequency-ratio analysis, 160 flashflood inventory points, and nine conditioning factors to map susceptibility in Barangay Bukang Liwayway, Bukidnon. Gacu et al. (2022) combined hazard, vulnerability, and exposure parameters through a GIS-based analytical hierarchy process in Odiongan, Romblon, using judgments from ten experts and end-users and checks against ground assessment and historical events. These studies demonstrate documented local spatial methods and validation practices, but they are mapping investigations rather than resident-facing flood-information systems.")
clone_before(existing_platforms, body_template, "Their factors, weights, classifications, geographic scales, and reported results are specific to their study areas. They support the need for local data and validation but do not provide Bacoor City rules or an operational susceptibility dataset. Accordingly, they are retained as Philippine literature relevant to spatial assessment and are not presented in the system-study comparison tables.")

foreign_studies = find_paragraph(doc, "Related Foreign Studies")
local_studies = find_paragraph(doc, "Related Local Studies")
clear_between(foreign_studies, local_studies)
foreign_entries = [
    (
        "Design and Implementation of Real-Time Flood Early Warning System (FEWS) Based on IoT Blynk Application: An Arduino- and GSM-Based Mobile Flood-Warning System",
        "Wisudawan (2021) designed and implemented a real-time Flood Early Warning System at a reservoir of Universitas Islam Indonesia. The system combined Arduino-based rain, temperature, humidity, and water-level sensing with GSM communication and the Blynk Android application. Its field implementation recorded observations for 30 days and classified readings as Normal, Advisory, Watch, or Warning.",
        "The reported mean temperature, humidity, and water level were 28.6 degrees Celsius, 63.7%, and 54.5 centimeters, and no flood occurred during the observation period. The system demonstrates sensor-to-mobile warning delivery, but its single-site, short-duration test did not include an actual flood event. Unlike FloodSense, it performs continuous environmental monitoring and warning rather than a resident-requested hypothetical susceptibility assessment, and its thresholds cannot be transferred to Bacoor City.",
    ),
    (
        "A User Experience Evaluation of a Mobile Application for Disseminating Site-Specific Impact-Based Flood Warnings: The A4alerts App: A Mobile Application for Site-Specific Impact-Based Flood Warnings",
        "Meléndez-Landaverde and Sempere-Torres (2023) evaluated A4alerts, a mobile crisis application linked to an operational site-specific flood early-warning system in Catalonia, Spain. Seventy-six participants from civil-protection, emergency-response, and local stakeholder groups completed an online interactive experiment from January to March 2022 using realistic flood scenarios. The study measured usefulness, functional appropriateness, and adoption and compared feature rankings through Friedman's test and Bonferroni-adjusted Wilcoxon tests.",
        "Ninety-seven percent considered the application useful during flood emergencies, and 93% stated that they would use or recommend it. Site-specific impact warnings, official regional warnings, and a self-protection checklist were ranked most useful. The evaluation was confined to recruited Catalan stakeholder networks, an online demonstration, and perceived usefulness rather than actual behavior during an event. A4alerts supports FloodSense's emphasis on location context and understandable guidance, but it disseminates operational warnings and does not validate FloodSense for Bacoor City.",
    ),
    (
        "Real-Time Flood Forecasting and Warning: A Comprehensive Approach toward HCI-Centric Mobile App Development: A Human-Centered Mobile Flood-Forecasting and Warning Prototype",
        "Alsabhan and Dudin (2023) developed an HCI-centered framework and functional prototype for mobile real-time flood forecasting and warning. Activity theory and scenario-based design guided a nominal-group process involving five prospective users and five experts, who generated 63 situations; 12 priority scenarios were retained for requirements analysis. Prototype evaluation examined task completion and satisfaction involving usefulness, ease of use, learnability, satisfaction, and overall experience.",
        "The process produced 93 design requirements and identified front-end users, weather-station users, and administrators. Reported mean satisfaction-category scores ranged from 3.93 to 4.08, while some analytical tasks remained more difficult. The study provides evidence for role-specific access and user-centered flood interfaces, but its small participant group and specific forecasting architecture limit generalization. FloodSense shares the concern for understandable interfaces and permissions but does not provide real-time forecasting or warnings.",
    ),
    (
        "Developing a Climate-Smart Web GIS App for Multi-Hazard Early Warning against Climate-Based Disaster Risks: A Tanzanian Web GIS for Flood and Drought Early Warning",
        "Msusa et al. (2025) developed a Web GIS application for flood and drought early warning in Tanzania. The study used Height Above Nearest Drainage analysis and GloFAS forecast data for flooding in Dar es Salaam, a ConvLSTM model and geospatial weather data for drought in Dodoma, participatory mapping with vulnerable communities for flood validation, and Rapid Application Development for the application.",
        "Community-mapped flood points showed strong spatial agreement with the flood-model output, while the drought model reported a median ROC-AUC of 0.87 and median F1 score of 0.81. The Web GIS integrated the resulting information for dissemination. Its dynamic forecast inputs, multi-hazard models, Tanzanian setting, and RAD process differ from FloodSense's deterministic hypothetical-scenario assessment. Its model performance and mapped areas cannot be generalized to Bacoor City.",
    ),
]
for heading, first, second in foreign_entries:
    clone_before(local_studies, subheading_template, heading)
    clone_before(local_studies, body_template, first)
    clone_before(local_studies, body_template, second)

synthesis_heading = find_paragraph(doc, "Synthesis of the Reviewed Literature")
clear_between(local_studies, synthesis_heading)
local_entries = [
    (
        "Web-Based Solution for Flood Warning Decision Support in the Province of Leyte, Philippines: An IoT Rainfall- and Water-Level Decision-Support Prototype",
        "Bentoso et al. (2021) developed an automated flood-warning prototype for the Province of Leyte using rainfall and water-level sensors, data transmission, a web interface, graphical reporting, threshold-based alerts, and decision-support outputs. The documented technical procedure focused on acquiring environmental readings, converting them into understandable web reports, and delivering alert messages to responsible authorities.",
        "The authors reported that the system displayed the flooding situation graphically, delivered alert-level messages, and supported monitoring and flood-control planning. The publication establishes a functional research prototype but does not establish province-wide deployment or transferability of its sensor thresholds. It is relevant to digital decision support, but FloodSense does not conduct continuous sensor monitoring or issue alerts, and Leyte-specific procedures cannot validate a Bacoor assessment.",
    ),
    (
        "Isabela City Flood Monitoring System Using Arduino and GSM Module: An Arduino and GSM Flood-Monitoring and SMS-Warning System",
        "Carpio, Asnaul, Said, and Amerin (2022) developed an Arduino Uno and GSM flood-monitoring system for Riverside Barangay in Isabela City, Basilan. The authors expressly used an Agile Development Model for incremental system development and a quantitative survey design for evaluation. Twenty City Disaster Risk Reduction and Management Office personnel participated in the system demonstration and initial device test covering functionality, design, security, efficiency, and reliability.",
        "The study reported an overall mean of 3.70, interpreted by the authors as Acceptable, and described the device as capable of providing warning signals under the tested conditions. It also identified a battery operating limit of approximately ten hours. The system shows a Philippine implementation of sensor monitoring and SMS warning, but its hardware, respondents, local thresholds, and evaluation cannot be transferred to Bacoor City. FloodSense instead uses a user-confirmed scenario and does not issue automatic warnings.",
    ),
    (
        "IOT-Based Flood Alert System for Narvacan, Sto. Tomas, La Union: A Water-Level Monitoring, Web, and SMS Alert Prototype",
        "Dela Cruz et al. (2024) used descriptive and developmental research designs and Agile methodology to develop an IoT water-level monitoring and alert prototype for Barangay Narvacan, Sto. Tomas, La Union. Interviews and observations informed development, while purposive sampling selected 30 evaluators: nine barangay council officials and 21 residents from seven sectors. The User Technology Acceptability Questionnaire assessed usefulness, information quality, and interface quality after system testing.",
        "The system obtained a grand mean of 6.87 on the study's seven-point scale, interpreted as extremely high acceptance; reported category means included 6.87 for usefulness, 6.83 for information quality, and 6.82 for interface quality. The results apply to the selected community and prototype, not to Bacoor City. Its continuous sensor readings, web monitoring, and SMS alerts differ from FloodSense's pre-event scenario assessment and separate preparedness guidance.",
    ),
    (
        "BantayBaha: Flood Monitoring System with SMS Notification: An IoT Water-Level Monitoring, Web-Logging, and SMS Notification System",
        "Delos Reyes, Cabreros, Paulino, and Garcia (2020) developed BantayBaha for Barangay Virgen Delas Flores in Baliwag, Bulacan after examining the barangay's manual measurement and megaphone-based warning practice. The project used IoT water-level sensing, SMS notification, and website logging, followed by verification, validation, repeated testing, and an evaluation involving 70 Baliuag University students.",
        "The respondents rated system functionality at 4.70, user interface at 4.50, and reliability at 4.57, each interpreted by the authors as Very Satisfactory. The study was limited to a single barangay and identified further hardware improvements, including alternative power and additional sensing. BantayBaha is relevant as a local digital warning system, but its evaluators, sensor conditions, and alert functions do not establish Bacoor-specific susceptibility or FloodSense effectiveness.",
    ),
]
for heading, first, second in local_entries:
    clone_before(synthesis_heading, subheading_template, heading)
    clone_before(synthesis_heading, body_template, first)
    clone_before(synthesis_heading, body_template, second)

gap_heading = find_paragraph(doc, "Research Gap")
clear_between(synthesis_heading, gap_heading)
for text in [
    "The reviewed literature and systems demonstrate several established approaches to flood information: sensor-based monitoring, forecast dissemination, Web GIS, mobile warnings, explainable interfaces, preparedness checklists, and Philippine SMS alert prototypes. The foreign systems show that usefulness depends on clear location context, role-appropriate interfaces, understandable warnings, and evaluation with users or geographic data. The Philippine systems show that local implementations commonly combine sensors, web or mobile presentation, and alert messages, although their thresholds, hardware, samples, and findings remain tied to their respective communities.",
    "These systems agree that flood information must be understandable and connected to action, but they differ in purpose, authority, data, method, geographic scale, and disaster-management phase. Wisudawan (2021), Bentoso et al. (2021), Carpio et al. (2022), Dela Cruz et al. (2024), and Delos Reyes et al. (2020) rely on sensed water or rainfall conditions. A4alerts disseminates warnings from an operational warning system, the HCI-centered prototype addresses real-time forecasting, and the Tanzanian Web GIS integrates forecast and predictive models. FloodSense does not adopt these monitoring, forecasting, or automatic-warning functions.",
    "The literature on susceptibility, GIS, deterministic rules, decision support, provenance, privacy, and software quality supports the verified FloodSense design. FloodSense combines a resident-confirmed hypothetical rainfall scenario and location with deterministic rule matching, an explained susceptibility output, and preparedness guidance that remains separate from assessment. Authentication, roles, source records, review status, and administrative controls support responsible use. These design choices do not eliminate the need for locally appropriate knowledge, geographic validation, expert review, and institutional authorization.",
    "The methods and results of the reviewed systems cannot be transferred directly to Bacoor City. Sensor thresholds, hydrological models, participant ratings, map classes, forecast inputs, and reported accuracy were produced in different environments and for different users. Their contributions therefore guide concepts and methods rather than supplying Bacoor-specific susceptibility facts or proving that FloodSense is effective.",
]:
    clone_before(gap_heading, body_template, text)

comparison_heading = find_paragraph(doc, "Comparative Synthesis of Related Studies")
clear_between(gap_heading, comparison_heading)
for text in [
    "Existing studies already provide evaluated flood-monitoring devices, mobile warning applications, Web GIS platforms, and decision-support prototypes. What remains insufficiently addressed in the inspected evidence is a Bacoor-focused, public-facing flow in which a resident deliberately selects a hypothetical rainfall scenario and location, receives an explainable deterministic susceptibility result, and then accesses preparedness guidance as a separate form of support.",
    "The gap also concerns provenance and local validity. None of the reviewed systems supplies an approved Bacoor City susceptibility knowledge base, locally validated rule thresholds, authorized geographic inputs, or evidence that external results can be generalized to the city's 47 barangays. FloodSense addresses the combination of scenario input, explainable rule-based assessment, location confirmation, separate preparedness support, authentication, and administrative governance as a research system. It remains distinct from official forecasting, warning, monitoring, and emergency-command systems.",
]:
    clone_before(comparison_heading, body_template, text)

methodology_heading = find_paragraph(doc, "METHODOLOGY")
clear_between(comparison_heading, methodology_heading)
clone_before(methodology_heading, body_template, "The following tables compare the verified methodologies, objectives, and contributions of the foreign and Philippine system studies retained in this chapter. Each entry is based on an inspectable primary publication or institutional full text that reports a development or research procedure and documented results.")

foreign_table_rows = [
    ["SYSTEM", "AUTHOR(S)/YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"],
    [foreign_entries[0][0], "Wisudawan (2021)", "Arduino/Blynk/GSM system design and 30-day field observation of rain, temperature, humidity, and water level.", "The study aimed to build and implement a real-time IoT flood early-warning system with mobile access.", "It produced four warning states and reported 30-day mean sensor values; no flood occurred during observation."],
    [foreign_entries[1][0], "Meléndez-Landaverde and Sempere-Torres (2023)", "Online scenario-based user experiment with 76 participants; Likert evaluation, ranking, Friedman test, and Wilcoxon post hoc analysis.", "The study aimed to evaluate the usefulness, functional appropriateness, and adoption of A4alerts.", "It found 97% perceived usefulness and ranked site-specific warnings, official warnings, and self-protection guidance highest."],
    [foreign_entries[2][0], "Alsabhan and Dudin (2023)", "Activity theory and scenario-based design; nominal groups of five users and five experts; prototype task and satisfaction evaluation.", "The study aimed to derive HCI requirements and evaluate a mobile real-time flood-forecasting and warning prototype.", "It identified 93 design requirements, three user groups, and mean satisfaction scores from 3.93 to 4.08."],
    [foreign_entries[3][0], "Msusa et al. (2025)", "HAND/GloFAS flood analysis, participatory mapping, ConvLSTM drought modeling, and RAD Web GIS development.", "The study aimed to develop a localized multi-hazard Web GIS early-warning application for Tanzania.", "It reported strong spatial agreement for flood validation and drought-model median ROC-AUC 0.87 and F1 0.81."],
]
local_table_rows = [
    ["SYSTEM", "AUTHOR(S)/YEAR", "METHODOLOGY", "OBJECTIVES", "CONTRIBUTIONS"],
    [local_entries[0][0], "Bentoso et al. (2021)", "IoT rainfall/water-level sensing, web visualization, threshold alerts, decision-support output, and functional demonstration.", "The study aimed to develop a web-based aid for monitoring flood conditions and supporting warning decisions in Leyte.", "It reported graphical monitoring, alert messages to authorities, and decision-support information for flood-control planning."],
    [local_entries[1][0], "Carpio et al. (2022)", "Agile Development Model; quantitative survey and initial device test with 20 CDRRMO personnel.", "The study aimed to develop an Arduino/GSM system that monitors water level and sends SMS warnings.", "It obtained an overall mean of 3.70, interpreted as Acceptable, and documented a ten-hour battery limit."],
    [local_entries[2][0], "Dela Cruz et al. (2024)", "Descriptive and developmental research; Agile; purposive UTAQ evaluation with 30 officials and residents.", "The study aimed to develop an IoT flood-alert prototype and assess its acceptability in Barangay Narvacan.", "It reported a grand mean of 6.87, interpreted as extremely high acceptance, across usefulness, information, and interface quality."],
    [local_entries[3][0], "Delos Reyes et al. (2020)", "Barangay interviews, IoT prototype development, verification/validation, repeated tests, and evaluation with 70 students.", "The study aimed to monitor water level, log readings on a website, and send SMS notifications in Virgen Delas Flores.", "It reported means of 4.70 for functionality, 4.50 for interface, and 4.57 for reliability."],
]

add_caption_before(methodology_heading, "Table 1. Comparison of related foreign studies")
add_table_before(methodology_heading, foreign_table_rows[:3], [1.42, 1.02, 1.42, 1.32, 1.32], font_size=8.2)
add_caption_before(methodology_heading, "Table 1. Continued", page_break_before=True)
add_table_before(methodology_heading, [foreign_table_rows[0]] + foreign_table_rows[3:], [1.42, 1.02, 1.42, 1.32, 1.32], font_size=8.2)
add_caption_before(methodology_heading, "Table 2. Comparison of related local studies", page_break_before=True)
add_table_before(methodology_heading, local_table_rows[:3], [1.42, 1.02, 1.42, 1.32, 1.32], font_size=8.2)
add_caption_before(methodology_heading, "Table 2. Continued", page_break_before=True)
add_table_before(methodology_heading, [local_table_rows[0]] + local_table_rows[3:], [1.42, 1.02, 1.42, 1.32, 1.32], font_size=8.2)

clone_before(methodology_heading, body_template, "Table 3 shows that the reviewed systems primarily emphasize sensor-based monitoring, forecasting, warning dissemination, and web or mobile presentation. A4alerts adds site-specific communication and self-protection guidance, while the HCI-centered and Tanzanian systems demonstrate role-aware design and geographic analysis. FloodSense differs by combining a resident-confirmed hypothetical rainfall scenario, location confirmation, deterministic rule-based susceptibility assessment, an explained result, and separate preparedness guidance. The matrix is not a performance ranking; a blank cell means that the feature was not established by the inspected source.")
add_caption_before(methodology_heading, "Table 3. Comparison of developed study to the other existing systems.")
matrix = [
    ["FEATURES", "EXISTING SYSTEM", "EXISTING SYSTEM", "EXISTING SYSTEM", "EXISTING SYSTEM", "EXISTING SYSTEM", "EXISTING SYSTEM", "EXISTING SYSTEM", "EXISTING SYSTEM", "DEVELOPED SYSTEM"],
    ["", "S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8", "FloodSense"],
    ["User-facing mobile or web interface", "✔", "✔", "✔", "✔", "✔", "", "✔", "✔", "✔"],
    ["Geographic or map-based presentation", "", "✔", "✔", "✔", "", "", "", "", "✔"],
    ["Sensor- or forecast-based monitoring", "✔", "", "✔", "✔", "✔", "✔", "✔", "✔", ""],
    ["Automatic warning or alert dissemination", "✔", "✔", "✔", "✔", "✔", "✔", "✔", "✔", ""],
    ["Location-specific information", "", "✔", "✔", "✔", "", "", "✔", "", "✔"],
    ["User-selected hypothetical rainfall scenario", "", "", "", "", "", "", "", "", "✔"],
    ["Deterministic rule-based susceptibility assessment", "", "", "", "", "", "", "", "", "✔"],
    ["Human-readable result context or explanation", "✔", "✔", "", "✔", "✔", "✔", "✔", "✔", "✔"],
    ["Preparedness or self-protection guidance", "", "✔", "", "", "", "", "", "", "✔"],
    ["Administrative management interface", "", "", "✔", "", "", "", "✔", "✔", "✔"],
]
add_table_before(methodology_heading, matrix, [1.75] + [0.53] * 9, font_size=7.0, two_header_rows=True)
for text in [
    "Legend:",
    "✔ - included in the documented system; blank - not established by the inspected source.",
    "S1 - Design and Implementation of Real-Time Flood Early Warning System (FEWS) Based on IoT Blynk Application",
    "S2 - A User Experience Evaluation of a Mobile Application for Disseminating Site-Specific Impact-Based Flood Warnings: The A4alerts App",
    "S3 - Real-Time Flood Forecasting and Warning: A Comprehensive Approach toward HCI-Centric Mobile App Development",
    "S4 - Developing a Climate-Smart Web GIS App for Multi-Hazard Early Warning against Climate-Based Disaster Risks",
    "S5 - Web-Based Solution for Flood Warning Decision Support in the Province of Leyte, Philippines",
    "S6 - Isabela City Flood Monitoring System Using Arduino and GSM Module",
    "S7 - IOT-Based Flood Alert System for Narvacan, Sto. Tomas, La Union",
    "S8 - BantayBaha: Flood Monitoring System with SMS Notification",
]:
    clone_before(methodology_heading, body_template, text)

# CHAPTER 3: SURECUT-style Materials -> Method; Agile is discussed inside Method; feasibility is kept separate.
references_heading = find_paragraph(doc, "REFERENCES")
clear_between(methodology_heading, references_heading)
clone_before(references_heading, body_template, "This chapter presents the materials and methods used to identify resident and institutional requirements, develop FloodSense, analyze the initial needs-assessment survey, verify the software, and prepare the product-quality evaluation.")
clone_before(references_heading, subheading_template, "Materials")
for text in [
    "Research materials included the bilingual resident questionnaire and consent notice, Google Forms and the exported response dataset, interview guides and recorded institutional interviews, request letters and approved school documents, scholarly and official references, the development-manuscript alignment records, requirements documents, diagrams, interface designs, test cases, and verification records.",
    "System-development materials included Windows computers, Android test devices, Git version control, Flutter and Dart for the resident application, Python with Django and Django REST Framework for the backend and administration system, PostgreSQL/PostGIS for relational and spatial storage, Mapbox and OpenStreetMap-compatible mapping components, and browser-based administration interfaces. These materials supported development and testing; their presence did not by itself establish the scientific or institutional validity of geographic or susceptibility information.",
]:
    clone_before(references_heading, body_template, text)
clone_before(references_heading, subheading_template, "Method")
for text in [
    "The study used a developmental-descriptive research design. The developmental component covered requirements analysis, design, iterative construction, integration, and technical verification of the resident application, backend services, geographic functions, deterministic Expert System, Decision Support System, and administration system. The descriptive component summarized the experiences, information problems, preparedness needs, and technology conditions reported in the initial resident survey. Institutional interviews and document review supplied qualitative evidence concerning current practices, data ownership, provenance, validation, and operational boundaries.",
    "FloodSense followed an Agile Development Method. This method was selected because requirements, interfaces, backend services, geographic functions, and tests were refined through repeated cycles of planning, design, implementation, verification, review, and correction. The study does not claim Scrum or Kanban ceremonies that are not established by the project records. Agile activities were organized into the following recurring phases:",
]:
    clone_before(references_heading, body_template, text)
for text in [
    "Requirements and planning. Resident survey evidence, institutional interviews, documents, verified literature, and system constraints were translated into traceable resident and administrator requirements.",
    "Design. The researchers prepared functional modules, interfaces, geographic flows, data structures, permissions, and test conditions before or alongside implementation.",
    "Iterative development. Mobile, backend, geographic, Expert System, Decision Support System, account, and administrative functions were implemented and integrated in manageable increments.",
    "Testing and review. Automated and manual checks were used to compare actual behavior with defined requirements, identify defects or contradictions, and return necessary corrections to the next development cycle.",
    "Release preparation. Stable functions, source and validation boundaries, privacy controls, and documentation were reviewed before a build could proceed to formal user evaluation or deployment preparation.",
]:
    add_bullet(references_heading, body_template, text)

clone_before(references_heading, subheading_template, "Requirements Gathering")
for text in [
    "Resident requirements were gathered through the initial questionnaire on flood experience, information sources, information problems, effects on decisions, preparedness needs, evacuation-center information, Android access, connectivity, language preference, and location-sharing comfort. BDRRMO and CPDCO interviews and documents were examined for institutional practices, data ownership, geographic records, preparedness procedures, evacuation information, and validation restrictions. Literature and official publications were used to define concepts and safety boundaries.",
    "Survey and interview evidence was not converted directly into a feature list. Each candidate requirement was checked for relevance to the research questions, consistency with the established pre-event scope, technical feasibility, privacy and safety implications, source authority, and compatibility with the current system architecture. Requests for capabilities outside the approved scope were retained as research context rather than represented as FloodSense functions.",
]:
    clone_before(references_heading, body_template, text)

clone_before(references_heading, subheading_template, "Participants and Sampling")
for text in [
    "Purposive sampling was used for the initial needs assessment. The final dataset contained 40 qualified respondents who consented, were at least 18 years old, lived in Bacoor City, and had experienced flooding or serious heavy-rain effects. Respondents came from 14 of the city's 47 barangays. Niog contributed 13 respondents, Zapote I to III contributed 15, and ten other barangays contributed 12. Because 70.0% of the sample came from Niog and Zapote I to III, the results describe the participating respondents and are not generalized to the entire Bacoor City population. A conventional response rate was not computed because no fixed sampling frame or distribution denominator was established.",
    "Thirty respondents (75.0%) were 18 to 34 years old. Key informants from BDRRMO and CPDCO were included for institutional requirements and data-governance context. Their statements were used to explain practices, sources, and limitations rather than as evidence of resident acceptance or system effectiveness.",
]:
    clone_before(references_heading, body_template, text)
add_caption_before(references_heading, "Table 4. Distribution of initial resident-survey respondents")
add_table_before(references_heading, [
    ["RESPONDENT AREA", "FREQUENCY", "PERCENTAGE"],
    ["Niog", "13", "32.5%"],
    ["Zapote I, II, and III", "15", "37.5%"],
    ["Other ten represented barangays", "12", "30.0%"],
    ["Total", "40", "100.0%"],
], [3.3, 1.5, 1.7], font_size=10.0)

clone_before(references_heading, subheading_template, "Research Instruments")
for text in [
    "The primary needs-assessment instrument was the bilingual FloodSense Resident Survey Questionnaire. It contained a Data Privacy Act notice and consent item, respondent eligibility and demographic questions, flood-experience items, current-information and preparedness items, evacuation-center information needs, Android and internet access, location-sharing preferences, and optional comments. Institutional interview guides and document requests were used for BDRRMO and CPDCO evidence.",
    "Technical verification used unit, integration, API, spatial, permission, and end-to-end test cases. The planned product-quality evaluation will use an approved instrument based on relevant ISO/IEC 25010:2023 characteristics. The initial resident questionnaire is a needs-assessment instrument and is not treated as a post-use evaluation of FloodSense.",
]:
    clone_before(references_heading, body_template, text)

clone_before(references_heading, subheading_template, "Data Gathering Procedure")
for text in [
    "The initial survey was distributed to eligible Bacoor City residents and closed on October 11, 2026. Responses were exported, screened for consent and eligibility, normalized for barangay reporting, and summarized without direct personal identifiers. The screening retained all 40 qualified submissions while documenting quality concerns, including one identical response pair, five unusually fast submissions, repeated long comments in six records, and some mixed selections in items that were intended to be mutually exclusive. These records were not silently removed because the available evidence did not establish that they were invalid; the limitations are disclosed when interpreting the results.",
    "Institutional evidence was collected through interviews, written requests, and review of available records. Research sources were checked against their primary publications or official documentation. Development requirements and test conditions were then revised through successive Agile cycles, with contradictions resolved according to source authority, research scope, safety, privacy, and implementation evidence.",
]:
    clone_before(references_heading, body_template, text)

clone_before(references_heading, subheading_template, "Statistical Treatment and Qualitative Analysis")
for text in [
    "Survey responses were analyzed using frequencies and percentages for categorical and multiple-response items. Ordered five-point items were summarized through their distributions and arithmetic means using the scoring defined in the approved survey summary. The eight information-problem items all fell within the survey's Often interpretation; the highest means concerned information that was too general for the respondent's location (4.16), difficulty accessing information when internet service or electricity was unreliable (4.14), preparedness advice that was too general (4.11), and unclear information source or date (4.08).",
    "Open comments were coded thematically for recurring information needs, map or location concerns, preparedness, warnings, severity, and evacuation-center information. Themes were interpreted as needs-assessment evidence and were not treated as proof that a requested capability belonged within the approved FloodSense scope. Institutional interviews and documents were coded for current practices, data ownership, provenance, validation, permissions, and system boundaries.",
]:
    clone_before(references_heading, body_template, text)

clone_before(references_heading, subheading_template, "System Development and Technical Verification")
for text in [
    "The resident application was developed in Flutter, while the backend, REST API, geographic services, and custom Web Administration System were developed with Django, Django REST Framework, GeoDjango, and PostgreSQL/PostGIS. The assessment service applies deterministic rule matching and conflict handling to the confirmed scenario and geographic facts. Preparedness guidance is delivered through a separate structured Decision Support System and does not alter the susceptibility class.",
    "Technical verification covers authentication and permissions, account flows, scenario and location confirmation, input validation, rule matching and conflict handling, API responses, spatial queries, map behavior, preparedness flows, evacuation-center lookup, administrative workflows, and mobile-to-service integration. Test records identify the precondition, input, expected result, actual result, status, environment, and defect reference where applicable. Technical test success establishes behavior under the tested conditions but does not establish scientific validity, institutional approval, usability, or production readiness.",
]:
    clone_before(references_heading, body_template, text)

clone_before(references_heading, subheading_template, "System Evaluation")
clone_before(references_heading, body_template, "Formal product-quality evaluation will be conducted on a stable FloodSense build using an approved instrument based on the relevant ISO/IEC 25010:2023 characteristics. The procedure will identify the evaluator groups, tasks, response scale, scoring method, interpretation, valid responses, and comments. Automated technical results and user evaluation results will be reported separately because they answer different questions.")
note = clone_before(references_heading, body_template, "[AUTHOR NOTE: Complete this section after the validated post-use system-evaluation responses and approved analysis are available.]")
note.runs[0].italic = True

clone_before(references_heading, subheading_template, "Ethical and Data-Governance Procedures")
clone_before(references_heading, body_template, "Participation in the initial survey was voluntary and governed by the approved consent and privacy notice. Direct identifiers were excluded from the analysis dataset, small geographic cells were interpreted cautiously, and access to source records was limited to authorized researchers and advisers. Resident account and optional location data are collected only for stated functions and remain subject to purpose limitation, access control, and the Data Privacy Act of 2012. Geographic, susceptibility, preparedness, and evacuation information must retain its source, date, validation status, and publication authority.")

# REFERENCES: retain current entries that remain relevant, remove the displaced ML study, and add the eight verified system studies.
appendices_heading = find_paragraph(doc, "PROVISIONAL APPENDICES (ORDER AND LETTERS TO BE CONFIRMED)")
# Collect reference paragraphs by traversing the XML between REFERENCES and APPENDICES.
refs = []
node = references_heading._p.getnext()
while node is not None and node is not appendices_heading._p:
    if node.tag == qn("w:p"):
        p = Paragraph(node, references_heading._parent)
        if p.text.strip():
            refs.append(p.text.strip())
    node = node.getnext()
refs = [r for r in refs if not r.startswith("Wang, H., Sanderson")]
refs.extend([
    "Alsabhan, W., & Dudin, B. (2023). Real-time flood forecasting and warning: A comprehensive approach toward HCI-centric mobile app development. Multimodal Technologies and Interaction, 7(5), 44. https://doi.org/10.3390/mti7050044",
    "Carpio, C. O., Asnaul, M. U., Said, A. H., & Amerin, B. Y. (2022). Isabela City flood monitoring system using Arduino and GSM module. International Journal of Multidisciplinary Research and Publications, 5(6), 13-17. https://ijmrap.com/wp-content/uploads/2022/11/IJMRAP-V5N5P134Y22.pdf",
    "Dela Cruz, C. A. V., Dela Cruz, J. R. L., Estolero, J. O., Lachica, B. J., Pineda, R. J. E., & Viduya, A. P. (2024). IOT-based flood alert system for Narvacan, Sto. Tomas, La Union [Unpublished undergraduate thesis]. Don Mariano Marcos Memorial State University-South La Union Campus. https://lakasa.dmmmsu.edu.ph/handle/123456789/117",
    "Delos Reyes, D. C., Cabreros, R., Paulino, J. L., & Garcia, P. J. (2020). BantayBaha: Flood monitoring system with SMS notification. Tekne Portal, 7(1), 58-73. https://buresearch.baliuagu.edu.ph/storage/01K6Q692A27DH2FQZTDQ4M5Y45.pdf",
    "Meléndez-Landaverde, E. R., & Sempere-Torres, D. (2023). A user experience evaluation of a mobile application for disseminating site-specific impact-based flood warnings: The A4alerts app. Journal of Flood Risk Management, e12951. https://doi.org/10.1111/jfr3.12951",
    "Msusa, A. D., Deus, D., Nyalali, O. G., Gidudu, A., Wangabo, E., Makuti, S., & Marandu, G. (2025). Developing a climate-smart Web GIS app for multi-hazard early warning against climate-based disaster risks. The International Archives of the Photogrammetry, Remote Sensing and Spatial Information Sciences, XLVIII-4/W16-2025, 59-66. https://doi.org/10.5194/isprs-archives-XLVIII-4-W16-2025-59-2025",
    "Wisudawan, H. N. P. (2021). Design and implementation of real-time Flood Early Warning System (FEWS) based on IoT Blynk application. ELKHA: Jurnal Teknik Elektro, 13(2), 113-121. https://doi.org/10.26418/elkha.v13i2.49003",
])
refs = sorted(dict.fromkeys(refs), key=lambda value: value.casefold())
clear_between(references_heading, appendices_heading)
for ref in refs:
    p = clone_before(appendices_heading, reference_template, ref)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)

# APPENDICES: replace provisional/conflicting material with compatible supplied diagrams.
replace_text(appendices_heading, "APPENDICES")
clear_after(appendices_heading)
appendix_items = [
    ("Appendix A. FloodSense Fishbone Diagram", DIAGRAM_DIR / "Fishbone_p1.png", 6.2),
    ("Appendix B. Resident Flood-Susceptibility Assessment Flowchart", DIAGRAM_DIR / "Flowchart_p1.png", 6.2),
    ("Appendix C. Resident Preparedness and Evacuation-Support Flowchart", DIAGRAM_DIR / "Flowchart_p2.png", 6.2),
]
for idx, (title, path, width) in enumerate(appendix_items):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    doc.add_picture(str(path), width=Inches(width))
    pic_p = doc.paragraphs[-1]
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

preliminary_hash_after = sha256("\n".join(p.text for p in doc.paragraphs[:65]).encode("utf-8")).hexdigest()
if preliminary_hash_before != preliminary_hash_after:
    raise RuntimeError("Preliminary-page text changed unexpectedly")

doc.save(DOCX)
print(DOCX)
print(f"references={len(refs)} tables={len(doc.tables)} paragraphs={len(doc.paragraphs)} shapes={len(doc.inline_shapes)}")
