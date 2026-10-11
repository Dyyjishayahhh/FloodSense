from __future__ import annotations

import shutil
import textwrap
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(r"E:\School Downloads-Jiraah\00 FILES IN SCHOOL\3RD YEAR 2ND SEM\DCIT 60\FLOODSENSE\1.FLOODSENSE-200A\FloodSense")
SOURCE = ROOT / "1. FLOODSENSE-MAIN" / "MAIN FILESS" / "FloodSense_Chapters_1_and_2_and_3_MAIN.docx"
OUT = ROOT / "1. FLOODSENSE-MAIN" / "WORKING_REVISION_2026-10-11"
DOCX_OUT = OUT / "FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx"
OUT.mkdir(parents=True, exist_ok=True)


def write(name: str, content: str) -> None:
    (OUT / name).write_text(textwrap.dedent(content).strip() + "\n", encoding="utf-8")


def font(size: int, bold: bool = False):
    path = Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size=size)


def wrapped(draw: ImageDraw.ImageDraw, xy, text, width, fnt, fill="#132238", spacing=8, anchor=None, align="center"):
    avg = max(1, int(width / max(1, fnt.getlength("M"))))
    lines = textwrap.wrap(text, width=avg)
    draw.multiline_text(xy, "\n".join(lines), font=fnt, fill=fill, spacing=spacing, anchor=anchor, align=align)


def box(draw, rect, text, fill="#EAF3F8", outline="#1B4965", title=False):
    draw.rounded_rectangle(rect, radius=24, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = rect
    wrapped(draw, ((x1+x2)//2, (y1+y2)//2), text, x2-x1-50, font(34 if title else 30, title), anchor="mm")


def arrow(draw, start, end, fill="#315B72", width=7):
    draw.line([start, end], fill=fill, width=width)
    import math
    ang = math.atan2(end[1]-start[1], end[0]-start[0])
    length = 26
    for off in (2.55, -2.55):
        p = (end[0] + length*math.cos(ang+off), end[1] + length*math.sin(ang+off))
        draw.line([end, p], fill=fill, width=width)


def canvas(title):
    im = Image.new("RGB", (2400, 1500), "white")
    d = ImageDraw.Draw(im)
    d.text((1200, 70), title, font=font(48, True), fill="#0B3954", anchor="mm")
    return im, d


def save_diagrams():
    # System architecture
    im, d = canvas("Provisional FloodSense System Architecture")
    box(d, (80, 420, 420, 720), "Resident\n(Android user)", "#E7F5EA")
    box(d, (1980, 420, 2320, 720), "Authorized\nAdministrator", "#FFF3D6")
    box(d, (530, 340, 920, 800), "Resident Mobile Application\n\nScenario and location input\nAssessment result and explanation\nPreparedness and center information", "#EAF3F8")
    box(d, (1010, 270, 1510, 870), "Django REST Services\n\nAuthentication and permissions\nDeterministic inference\nGeographic services\nStructured DSS\nEvacuation-center service", "#E8E7FA")
    box(d, (1600, 340, 1880, 800), "Custom Web Administration System\n\nGoverned content and record workflows", "#FFF3D6")
    box(d, (1010, 1040, 1510, 1320), "PostgreSQL / PostGIS\nControlled relational and spatial records", "#F0F0F0")
    box(d, (850, 115, 1670, 230), "External reference sources (not system actors): approved rainfall, geographic, preparedness, and center information", "#FCE8E6")
    arrow(d, (420, 570), (530, 570)); arrow(d, (920, 570), (1010, 570)); arrow(d, (1510, 570), (1600, 570)); arrow(d, (1980, 570), (1880, 570))
    arrow(d, (1260, 870), (1260, 1040)); arrow(d, (1260, 230), (1260, 270))
    d.text((1200, 1420), "Reference sources require provenance, permission, and validation; they are not users or actors.", font=font(25), fill="#8C2F39", anchor="mm")
    im.save(OUT / "Diagram_System_Architecture_Provisional.png")

    # Conceptual framework
    im, d = canvas("Provisional FloodSense Conceptual Framework")
    box(d, (80, 280, 650, 1190), "INPUTS\n\nUser-confirmed rainfall scenario\nManual or foreground-confirmed location\nApproved geographic and reference records\nValidated rules and provenance\nPublished preparedness and center information", "#EAF3F8", title=True)
    box(d, (815, 280, 1585, 1190), "FUNCTIONAL PROCESSES\n\n1. Scenario and Location Input\n\n2. Susceptibility Assessment and Explanation\n\n3. Preparedness and Evacuation Information\n\n4. Governance and Data Stewardship", "#E8E7FA", title=True)
    box(d, (1750, 280, 2320, 1190), "OUTPUTS\n\nScenario-specific susceptibility class\nHuman-readable explanation\nInteractive map presentation\nContextual preparedness guidance\nNearest-center information when verified records are available\nSource and limitation notices", "#E7F5EA", title=True)
    arrow(d, (650, 735), (815, 735)); arrow(d, (1585, 735), (1750, 735))
    im.save(OUT / "Diagram_Conceptual_Framework_Provisional.png")

    # Use case
    im, d = canvas("Provisional FloodSense Use Case Diagram")
    box(d, (80, 520, 410, 790), "Resident")
    box(d, (1990, 520, 2320, 790), "Authorized Administrator")
    cases = [
        (650, 200, "Register, verify, and manage account"),
        (650, 440, "Select or confirm scenario and location"),
        (650, 680, "Request assessment and view explanation"),
        (650, 920, "Use structured preparedness guidance"),
        (650, 1160, "View nearest-center information*"),
        (1450, 260, "Review dashboard and governed records"),
        (1450, 550, "Manage DSS and evacuation-center content"),
        (1450, 840, "Review sources, imports, and audit history"),
        (1450, 1130, "Review parameter governance (no raw rule editing)"),
    ]
    for x, y, txt in cases:
        d.ellipse((x-260, y-85, x+260, y+85), fill="#F7FBFD", outline="#1B4965", width=4)
        wrapped(d, (x, y), txt, 470, font(25), anchor="mm")
        if x < 1000:
            arrow(d, (410, 655), (x-260, y), width=4)
        else:
            arrow(d, (1990, 655), (x+260, y), width=4)
    d.text((1200, 1415), "*Public usefulness depends on authorized, current evacuation-center records; no live capacity or safe-route guarantee.", font=font(23), fill="#8C2F39", anchor="mm")
    im.save(OUT / "Diagram_Use_Case_Provisional.png")

    # DFD
    im, d = canvas("Provisional FloodSense Level-0 Data Flow Diagram")
    box(d, (70, 570, 390, 820), "Resident")
    box(d, (2010, 570, 2330, 820), "Administrator")
    box(d, (880, 320, 1520, 980), "FLOODSENSE\n\nAuthenticate and authorize\nValidate scenario and location\nApply deterministic rules\nProvide explanation and DSS guidance\nServe governed records", "#E8E7FA", title=True)
    box(d, (840, 1120, 1560, 1360), "D1: Controlled system data\naccounts • geography • rules • DSS • centers • provenance • audit", "#F0F0F0")
    box(d, (810, 120, 1590, 250), "D2: Approved external reference information", "#FCE8E6")
    arrow(d, (390, 640), (880, 550)); arrow(d, (880, 720), (390, 760))
    arrow(d, (2010, 640), (1520, 550)); arrow(d, (1520, 720), (2010, 760))
    arrow(d, (1200, 980), (1200, 1120)); arrow(d, (1200, 250), (1200, 320))
    im.save(OUT / "Diagram_DFD_Level_0_Provisional.png")

    # HIPO
    im, d = canvas("Provisional FloodSense HIPO Hierarchy")
    box(d, (840, 150, 1560, 330), "FloodSense")
    tops = [(130, 570, "Resident Application"), (680, 570, "Assessment Services"), (1230, 570, "Preparedness Services"), (1780, 570, "Administration and Governance")]
    for x, y, txt in tops:
        box(d, (x, y, x+480, y+180), txt, "#EAF3F8", title=True)
        arrow(d, (1200, 330), (x+240, y), width=5)
    subs = [
        (130, ["Account and onboarding", "Scenario/location input", "Map and results"]),
        (680, ["Input validation", "Rule matching/conflict handling", "Explanation generation"]),
        (1230, ["Structured DSS flow", "Published guidance", "Nearest-center lookup"]),
        (1780, ["Content/record workflows", "Provenance and audit", "Permissions and technical admin separation"]),
    ]
    for x, items in subs:
        for i, txt in enumerate(items):
            y = 860 + i*170
            box(d, (x, y, x+480, y+120), txt, "#F7FBFD")
            arrow(d, (x+240, 750), (x+240, y), width=4)
    im.save(OUT / "Diagram_HIPO_Provisional.png")


save_diagrams()


write("00_README_DELIVERABLE_INDEX.md", r'''
# FloodSense Phase 1 working-revision package

Audit date: 2026-10-11 (Asia/Manila)  
Repository commit: `9d32d815e0c54e081e4b2ca617bbe939d1df6b7c`  
Commit date/message: 2026-10-11 — `Replace Project FIles`  
Survey status: ongoing; no final count, percentages, response rate, or conclusions are reported.

## Outputs

1. `FloodSense_Chapters_1_to_3_Survey_Ongoing_Working_Revision_2026-10-11.docx`
2. `01_Evidence_and_Consistency_Audit.md`
3. `02_Problem_Objective_Method_Feature_Traceability.md`
4. `03_Literature_and_Related_Study_Matrix.md`
5. `04_Survey_Data_Dictionary_and_Analysis_Plan.md`
6. `05_Feasibility_Study_Materials.md`
7. `06_Provisional_Diagrams.md` and five PNG diagrams
8. `07_Appendix_Inventory.md`
9. `08_Survey_Results_Insertion_Guide.md`
10. `09_Revision_Log.md`
11. `10_Missing_Information_and_Pending_Decisions.md`

The original manuscript and the active survey questionnaire were not overwritten or edited. Appendix letters, the specific SDLC, final sampling details, final survey statistics, final ISO/IEC 25010 evaluation, and official local-data validation remain pending.
''')


write("01_Evidence_and_Consistency_Audit.md", r'''
# Evidence and consistency audit

## Repository state

- Current development commit: `9d32d815e0c54e081e4b2ca617bbe939d1df6b7c`
- Commit date/message: 2026-10-11 — `Replace Project FIles`
- Branch: `main`
- Working tree: not clean. Several tracked interview/questionnaire-path copies are deleted while equivalent or newer interview files are untracked under `interviews/`; `tmp/` contains audit working files. These user changes were preserved.
- Current mobile verification: `flutter test` passed 389 tests; `flutter analyze` reported no issues.
- Backend verification: not rerun in the available bundled Python because Django dependencies were unavailable. Repository documentation records a historical 2026-10-05 run of 891 backend tests, 327 Flutter tests, Flutter analysis, and a debug APK build. This is historical evidence, not a current rerun.

## Evidence hierarchy applied

Integrated code, routes, services, models, migrations, interfaces, and tests were used together. A model, proposal, mockup, script, or screenshot alone was not treated as a complete function. The 47-barangay layer was treated as an administrative reference only. Provisional MGB and other research data were not described as official or operational.

## Current development status

| Feature or module | Status | Current evidence and interface | Research/data qualification | Manuscript action |
|---|---|---|---|---|
| User-confirmed rainfall intensity and duration | IMPLEMENTED AND CURRENTLY VERIFIED | Mobile assessment flow, API options/evaluate routes, current mobile tests | Scenario input; not live weather | Retain and state explicitly |
| Scenario confirmation and explicit assessment request | IMPLEMENTED AND CURRENTLY VERIFIED | Mobile interaction and tests | No automatic polling | Retain |
| Manual area selection and map pin | IMPLEMENTED AND CURRENTLY VERIFIED | Resident map and assessment widgets | Limited to supported data | Retain |
| One-time foreground GPS and barangay confirmation/correction | IMPLEMENTED AND CURRENTLY VERIFIED | Location flow, geographic endpoints, current mobile tests | Not background tracking | Retain with privacy limit |
| 47-barangay reference display | IMPLEMENTED AND CURRENTLY VERIFIED | Boundary endpoint and map display | Derived/technically checked; not City-issued and not susceptibility data | Keep qualification |
| Deterministic rule matching, priorities, conflicts, and explanation | IMPLEMENTED AND CURRENTLY VERIFIED | Expert-system services/endpoints and current mobile integration tests | Operational local validity blocked by data/rule approval | Describe software function, not official result |
| Susceptibility output (Low–Very High) | IMPLEMENTED AND CURRENTLY VERIFIED | Assessment response and result UI | Demonstration/provisional inputs only | Retain qualification |
| Structured multi-question DSS | IMPLEMENTED AND CURRENTLY VERIFIED | Structured-flow start/answer endpoints, resident Prepare UI, current mobile tests | Content still requires source/approval governance | Correct manuscript’s “unverified/incomplete” wording |
| Published DSS guidance separate from assessment | IMPLEMENTED AND CURRENTLY VERIFIED | Guidance endpoints and UI | Cannot change susceptibility class or issue orders | Retain separation |
| Resident registration, login, verification, legal acceptance, onboarding, preferences, password and deletion controls | IMPLEMENTED AND CURRENTLY VERIFIED | Auth/account routes, resident shell, screens, current mobile tests | Production identity/deployment not established | Correct manuscript’s “unimplemented” wording |
| Assessment history | NOT IMPLEMENTED | No persistent resident assessment history; DSS history flag is false | — | Keep as limitation |
| Notifications/live alerts | OUTSIDE THE CURRENT SCOPE | Mobile copy explicitly states live alerts are not enabled | — | Do not claim |
| Provider-aware Mapbox/OSM map | IMPLEMENTED AND CURRENTLY VERIFIED | Mapbox and flutter_map/OSM paths, provider configuration, tests | Token/provider configuration required | Update technology description |
| Evacuation-center Admin workflow | IMPLEMENTED AND CURRENTLY VERIFIED | Models, routes, custom Admin forms/import review | Data approval/currentness separate | Retain |
| Resident nearest-center lookup and display | IMPLEMENTED AND CURRENTLY VERIFIED | Map/nearest endpoints, resident UI, current mobile tests | Public usefulness blocked by authorized current center data; no live capacity/safe routing | Correct manuscript’s “not exposed” wording |
| Evacuation-center CSV/XLSX staged import | IMPLEMENTED AND CURRENTLY VERIFIED | Staging/review/import workflow and tests in code | Applies to center records, not Expert System parameters | Distinguish from parameter CSV requirement |
| Fixed-schema Expert System parameter CSV export/import | CONFIRMED REQUIREMENT — NOT YET IMPLEMENTED | No complete export/import routes, schema validation, ordinary Admin controls, and end-to-end tests | Confirmed governance requirement | State as unimplemented |
| Custom Admin authentication/dashboard | IMPLEMENTED AND CURRENTLY VERIFIED | Authenticated custom portal routes/templates/tests | Technical Django Admin is separate | Retain distinction |
| Map/rainfall reference review, parameter governance, DSS content, sources, audit history | IMPLEMENTED AND CURRENTLY VERIFIED | Custom Admin routes, services, interfaces | Publication/approval status still governs content | Update scope |
| Raw rule editing by ordinary administrators | OUTSIDE THE CURRENT SCOPE | Not exposed in custom Admin | Technical Django Admin reserved for research/developer maintenance | Preserve prohibition |
| MGB susceptibility import/data | IMPLEMENTED BUT NOT EXPOSED TO THE INTENDED USER | Models/import processing exist | Data are provisional/unapproved; no proof of operational integration | Do not present as official/current public data |
| Approved operational Bacoor susceptibility dataset | BLOCKED OR AWAITING DATA | No approved dataset in `research_data/approved` | Requires provenance, permission, method, and institutional/expert validation | Keep as primary limitation |
| Real-time rainfall/water-level monitoring, official forecast/warning | OUTSIDE THE CURRENT SCOPE | No operational sensors/feed/polling; explicitly excluded | — | Exclude claims |
| Offline operation | NOT IMPLEMENTED | Integrated functions depend on server/network | Error/offline states are not offline capability | Keep limitation |
| Production deployment | NOT IMPLEMENTED | No verified production release/operations evidence | — | Do not claim |
| Needs survey | ACTIVE WORK IN PROGRESS | Exact approved questionnaire exists and is currently being answered | Final valid N, percentages, and response rate unavailable | Use placeholders only |
| ISO/IEC 25010 evaluation | PLANNED | Quality model/reference present; no completed validated evaluation dataset | Evaluator groups/instrument/scoring pending | Do not report results |

## Contradictions and corrections

1. The manuscript said resident accounts and onboarding were conceptual or unavailable; the current mobile/backend integration exposes them. The working copy corrects this.
2. The manuscript said the structured multi-question DSS was not verified; it is now integrated and covered by current mobile tests. The working copy corrects this while preserving content-validation limits.
3. The manuscript said residents could not discover nearby centers; current endpoints and resident interfaces provide nearest-center lookup. The correction does not imply official center data, live capacity, or safe routing.
4. The manuscript described `flutter_map`/Leaflet as the sole map approach; the repository now includes a provider-aware Mapbox path with OSM/flutter_map fallback and Mapbox-enabled Admin maps.
5. The manuscript named Design Science Research and a generic iterative/incremental SDLC. Current records prove artifact development, surveys, interviews, testing, and revisions, but do not prove that the group deliberately followed the Peffers DSR method or a specific Scrum/Kanban/RAD process. The working copy uses a provisional developmental-descriptive research classification and marks specific SDLC selection pending.
6. Older feasibility material describes Firebase, an AI assistant, offline Firestore behavior, and other superseded concepts. These were not carried into the working copy because current integrated code uses Django/PostgreSQL/PostGIS and deterministic/structured DSS flows.
7. The BDRRMO consultation suggested a Low-yellow, Moderate-violet, High-orange, Very High-red scheme, whereas the manuscript/code use green/yellow/orange/red. The interview suggestion is not recorded as a formally approved standard; the matter remains pending.

## Evidence ready now versus provisional

Ready now: problem context from official literature and interviews; system scope boundaries; current implemented software functions; current mobile test count; research instruments; planned statistical procedures; provisional feasibility analysis; data-governance limitations.

Provisional: exact SOP wording affected by final survey evidence; final respondent profile and needs findings; final sampling/reporting; selected specific SDLC; approved susceptibility parameters and data; BDRRMO-approved DSS content and center dataset; final diagrams/appendix letters; ISO/IEC 25010 evaluation and conclusions; production claims.
''')


write("02_Problem_Objective_Method_Feature_Traceability.md", r'''
# Problem–objective–method–feature traceability matrix

All questions and objectives below are provisional until the ongoing survey is closed and analyzed.

| Evidence or identified problem | Source | Provisional SOP question | Aligned objective | Research method | Intended finding or system response | Validation status | Manuscript location/action |
|---|---|---|---|---|---|---|---|
| Residents use several information channels, but final evidence on access, timing, clarity, specificity, and consistency is still pending | Active survey Q4–Q14; earlier surveys kept separate; BDRRMO/CPDCO interviews | 1. What flood-information practices, access conditions, and problems do eligible Bacoor respondents report? | Identify resident and stakeholder flood-information needs | Descriptive frequencies/percentages; multiple-response analysis; qualitative coding of interviews | Evidence-based problem profile; decide which proposed functions remain justified | Survey ongoing | SOP, Objective 1, Chapter 3 participants/analysis; insert final results later |
| Geographic, rainfall, preparedness, and center information require provenance, permission, and local validation | BDRRMO/CPDCO interviews; repository data status; related literature | 2. What verified knowledge, parameters, geographic records, and governance controls are required? | Determine validated facts, rules, sources, and controls | Document/source review; key-informant thematic analysis; provenance/validation checklist | Controlled knowledge requirements; limitations and approval workflow | Partially validated; operational dataset absent | SOP, Objective 2, framework, scope, requirements gathering |
| Need to turn validated requirements into a usable resident/admin artifact without implying live forecasting | Repository and confirmed scope | 3. How can the verified requirements be implemented as a scenario-based Android and Web Administration system? | Develop the verified major functions | Developmental system work; requirements traceability; interface and API implementation | Resident scenario/location flow, deterministic explanation, separate DSS, map, accounts, governed Admin workflows | Software implemented in major areas; data validity pending | Objective 3, framework, architecture, Chapter 3 |
| Software functions require technical verification distinct from user evaluation | Current tests and historical backend verification | 4. To what extent do implemented functions operate correctly and consistently under defined test cases? | Test the core modules and integrations | Unit, integration, API, spatial, permission, map, and end-to-end test cases; defect logs | Verified behavior and documented defects | Current mobile: 389 tests pass; backend current rerun unavailable | Objective 4, Chapter 3 testing |
| Quality and usability cannot be inferred from code tests | ISO/IEC 25010:2023; evaluation instrument pending | 5. How do the defined evaluator groups assess relevant FloodSense quality characteristics? | Evaluate using an approved ISO/IEC 25010-based instrument | Task-based evaluation plus descriptive item/subscale statistics; open-comment coding | Quality profile, limitations, and prioritized revisions | Not yet conducted | Objective 5, Chapter 3 evaluation; results placeholders |

## Feature-to-evidence guardrails

- Scenario/location assessment responds to validated needs only after Q5–Q14 results are available; code completion alone does not prove public need.
- Nearest-center lookup is implemented, but retention as a validated thesis requirement still depends on survey evidence, stakeholder confirmation, and an authorized current center dataset.
- Resident accounts are implemented for access/legal/onboarding controls; no claim is made that respondents requested persistent histories.
- Structured DSS guidance is implemented as a separate support flow. It does not alter the Expert System result or issue an evacuation order.
- Live monitoring, official warnings, safe routing, occupancy tracking, background GPS, offline operation, machine learning, and self-modifying rules are not responses within the current scope.
''')


write("03_Literature_and_Related_Study_Matrix.md", r'''
# Literature and related-study matrix

The matrix separates peer-reviewed research from official platforms. An official product page may establish purpose and technical behavior, but it is not converted into an IMRaD study.

| Source | Type and method/data | Main verified contribution | Relevance to FloodSense | Important difference/limitation | Manuscript use |
|---|---|---|---|---|---|
| Membele, Naidu, & Mutanga (2022), *Examining flood vulnerability mapping approaches in developing countries* | Peer-reviewed scoping review of mapping approaches | Synthesizes data, indicator, weighting, and validation problems in developing-country mapping | Supports careful distinction among hazard, vulnerability, risk, and susceptibility | Does not validate a Bacoor model or thresholds | Foreign literature: susceptibility/data limitations |
| Mostafiz et al. (2022), *Actionable information in flood risk communications...* | Peer-reviewed analysis/review of flood-risk communication and web tools | Emphasizes actionable individual/community information and limitations of technical-only tools | Supports separate understandable preparedness guidance | Long-term risk communication differs from FloodSense’s user-confirmed scenario | Foreign literature: risk communication/DSS |
| Fakhruddin et al. (2022), *Harnessing risk-informed data...* | Peer-reviewed synthesis of risk-informed data for resilience | Supports responsible integration and governance of risk data | Supports provenance and institutional responsibility | Does not supply local Bacoor data | Foreign literature: data governance |
| Papadopoulos et al. (2022), rule-based CDSS review | Systematic review of rule-based decision-support technologies and standards | Distinguishes knowledge representation, rule engines, and case-specific evaluation | Supports controlled knowledge base/inference distinction | Clinical domain; cannot transfer domain rules | Foreign literature: Expert Systems |
| Kostopoulos et al. (2024), explainable DSS review | Recent review of explainable AI/DSS | Highlights transparency, interpretability, and human-understandable justification | Supports explanations of inputs, rule basis, sources, and limitations | Mainly AI-based; FloodSense is deterministic | Foreign literature: explainability |
| Tullis & Kar (2021), GIS provenance | Peer-reviewed conceptual/ethical GIS analysis | Shows reproducibility and ethical importance of provenance | Supports source/version/validation records | Does not validate specific Philippine layers | Foreign literature: provenance |
| Deltares (2025), FloodAdapt | Official software documentation; SFINCS physics-based flood simulation plus Delft-FIAT impact assessment | Provides scenario and adaptation comparison maps/metrics | Demonstrates explicit scenario exploration | Model-based flood/impact simulation; FloodSense does not implement SFINCS/FIAT | Foreign system discussion; label as platform, not study |
| Copernicus EMS (2023/2026), GloFAS | Official operational service; LISFLOOD hydrological modeling, gauge calibration, parameter regionalization | Global river-discharge forecasting/early awareness products | Illustrates operational forecasting and scale | FloodSense is not a forecast/monitoring service | Foreign system discussion; label as platform |
| Iowa Flood Center, IFIS | Official platform and older peer-reviewed system paper; radar/sensor/hydrologic forecasting and public map | Integrates Iowa observations, forecasts, alerts, and inundation maps | Illustrates public geographic communication and disclaimers | Iowa-specific operational inputs; older foundational paper | Foreign system discussion; publication date caveat |
| FEMA (2025), RAPT | Official user guide; GIS layer comparison for resilience planning | Combines hazard, infrastructure, and community layers for planners | Supports layer-based administrative decision context | Planning tool, not resident susceptibility inference | Foreign platform/literature |
| GOV.UK/Environment Agency (2024), long-term flood-risk service | Official public service and national assessment; modeled area risk | Communicates multiple flood sources and limits of property-level inference | Supports careful authority/limitation statements | England-specific long-term risk, not Bacoor scenario input | Foreign platform/literature |
| UP NOAH Center, Know Your Hazards/NOAH Studio | Official Philippine platforms using simulated, satellite, historical, and mapped hazard layers | Provides location lookup and Philippine flood/hazard visualization | Relevant local map/risk communication comparison | Scenario definitions, data production, and authority differ; no direct data transfer | Local platform discussion |
| GeoRisk Philippines, HazardHunterPH | Official multi-agency platform; point lookup/spatial intersection of agency layers | Provides indicative site-specific multi-hazard assessment with source agencies | Supports source-aware location lookup | Indicative official platform, not FloodSense’s rule-based rainfall scenario | Local platform discussion |
| DOST-PHIVOLCS (2022a), GeoMapperPH training | Official institutional activity report; web/mobile collection to GeoRiskPH | Documents governed LGU data collection and account/dashboard use | Supports administrative data governance | Not a peer-reviewed evaluation or Bacoor implementation | Local literature/platform, not IMRaD study |
| DOST-PHIVOLCS (2022b), PlanSmartPH training | Official institutional activity report; baseline and post-disaster data entry for recovery planning | Documents structured rehabilitation/recovery planning workflow | Illustrates governed planning inputs | Post-disaster LGU planning differs from resident pre-event assessment | Local literature/platform, not IMRaD study |
| Bentoso et al. (2021), Leyte flood-warning DSS | Peer-reviewed conference paper; rain/water-level sensors, web visualization, and alert messaging in Leyte | Reports graphical flood-condition presentation and alert messages for authorities | Philippine comparison for DSS and warning communication | Sensor-based monitoring/warning; not transferable to Bacoor and outside FloodSense scope | Related local study; present method/results/limits |
| ISO/IEC 25010:2023 | International standard product-quality model | Defines product-quality characteristics | Basis for planned evaluation instrument | Not a research design, SDLC, or completed evaluation | Chapter 2 quality literature; Chapter 3 plan |

## Verification notes

- Bentoso et al. (2021) is the strongest current local IMRaD-style study in the manuscript; DOI: `10.1109/ICRACOS53680.2021.9701990`.
- FloodAdapt, GloFAS, IFIS, RAPT, the England service, NOAH, HazardHunterPH, GeoMapperPH, and PlanSmartPH are chiefly official platforms or institutional documentation in the present evidence set. Their “methodology” cells should describe technical procedures, not invented research designs or SDLC models.
- A recent Philippine peer-reviewed needs/risk-communication study should be added only after its full text, journal status, participants, method, results, and bibliographic details are independently verified.
- No reviewed foreign formula, weight, threshold, color convention, or map classification is transferable to Bacoor without local source documentation and validation.
''')


write("04_Survey_Data_Dictionary_and_Analysis_Plan.md", r'''
# Survey data dictionary and analysis plan

Source instrument: `revised-questionnaire-the-approved-survey.docx`. This is the active version and was not modified. Close the survey and archive a dated export before cleaning or analysis.

## Eligibility and record handling

- Include only respondents who consented, answered Q1 as 18 or older, and answered Q2 as living in an included Bacoor community.
- Q2 “I am not sure” requires a documented inclusion decision; do not silently count it as Yes.
- Retain “Prefer not to answer,” “I am not sure,” “I cannot recall,” “not applicable,” and skipped items as distinct codes where analytically useful; do not convert them to favorable or unfavorable responses.
- Report the denominator for every item. Multiple-response items may exceed 100% when percentages are based on valid respondents selecting each option.
- Remove names from the analysis dataset and keep any linkage file separately with restricted access.

| Item | Variable | Response/coding | Valid/missing treatment | Planned analysis/output | SOP/requirement link |
|---|---|---|---|---|---|
| Consent | consent | 1 agree; 0 decline | Decline = terminate/exclude | Flow count only | Ethics |
| D1 | name_raw | Text; direct identifier | Never tabulate publicly; blank allowed | Remove from analysis copy | Privacy |
| D2 | age | Integer; special code for prefer not to answer | Range-check; report missing | Median/range or approved age groups after collection | Participant description |
| D3 | sex_gender | 1 male; 2 female; 3 self-describe; 9 prefer not | Self-description protected; suppress small cells if needed | Frequency/percentage | Participant description |
| D4 | residence_raw; barangay_std | Free text normalized to approved barangay list; retain original separately | Invalid/non-Bacoor flagged; ambiguous entries reviewed | Frequency/map only when confidentiality permits; `[Insert final distribution by barangay.]` | Local coverage/eligibility |
| D5 | residence_length | 1 <1 year; 2 1–5; 3 6–10; 4 >10; 9 prefer not | 9 separate from missing | Frequency/percentage | Context |
| Q1 | age_eligible | 1 yes; 0 no; 9 prefer not | 0/9 terminate/exclude | Eligibility flow | Sampling |
| Q2 | bacoor_eligible | 1 yes; 0 no; 8 unsure | 0 exclude; 8 pending documented rule | Eligibility flow | Sampling |
| Q3 | flood_effect_frequency | 5 very often; 4 often; 3 sometimes; 2 rarely; 1 never; 8 cannot recall | 8 separate/not in ordered summaries | Frequency/percentage; optional median, no unsupported mean | Problem context |
| Q4 | first_information_source | Categories as printed (official social media; barangay PA; SMS; TV/radio; government/weather site; family/community; personal observation; none) | Single response; blank missing | Bar chart/table | Current practices |
| Q5 | obtain_needed_information | 5 always to 1 never; 8 usually no need | 8 separate | Frequency/percentage; collapse only with predeclared rule | Access problem |
| Q6 | information_timing | Early enough; shortly before action; after action started; after affected; none; cannot recall | Keep all categories | Frequency/percentage | Timeliness problem |
| Q7 | understandability | Very easy; easy; neutral; difficult; very difficult; receives none | “Receives none” separate | Frequency/percentage; favorable/neutral/unfavorable only if declared before analysis | Clarity problem |
| Q8 | location_specificity | All details; some; too general; not important; receives none; unsure | Keep separate | Frequency/percentage | Location-specific information |
| Q9 | conflicting_information | Very often; often; sometimes; rarely; never; one source; unsure | One-source/unsure separate | Frequency/percentage | Consistency/provenance |
| Q10 | decision_effect | Helps in time; helps but late; available/not useful; cannot obtain; no need; unsure | Keep separate | Frequency/percentage | Decision support need |
| Q11 | information_problems | One binary variable per option | Multi-select; check contradictions with “no problem” | Count and % selecting each; ranked bar chart | Problem evidence |
| Q12 | useful_lead_time | One binary variable per option | Multi-select; “depends/no need” separate | Count and % selecting each | Preparedness timing; not evidence of live forecast capability |
| Q13 | helpful_information_types | One binary variable per option; maximum 2 selections | Flag >2 selections for review; do not auto-delete without rule | Count and % selecting each; feature-retention evidence with interview/literature triangulation | Requirements |
| Q14 | needs_met | 5 completely to 1 not at all; 8 no need | 8 separate | Frequency/percentage; optional median | Main information gap |
| Q15 | mobile_access | Own smartphone; shared smartphone; basic phone; no regular phone; prefer not | Keep separate | Frequency/percentage; `[Insert final device-ownership result.]` | Android operational feasibility |
| Q16 | helpful_delivery_methods | One binary variable per option | Multi-select; contradictory selections reviewed | Count and % selecting each; `[Insert final connectivity result.]` only if connectivity was separately measured—this item measures method preference, not internet quality | Delivery feasibility |

## Open-text and validation procedures

The active instrument contains direct identifiers and free-text residence/self-description fields, but no general open-comment question. Two researchers should independently normalize barangay text using an approved codebook, compare disagreements, retain an audit column, and resolve uncertain locations without guessing. Any later open comments should be coded through an explicit codebook, double-coding sample, discrepancy resolution, and frequency reporting with de-identified excerpts only when consent permits.

## Statistical treatment

Use counts and valid percentages for nominal and multiple-response items. Report ordered response distributions and medians where useful. Do not use a weighted mean merely because choices are numbered; it is defensible only if the team predefines a genuinely ordered equal-interval interpretation and the adviser approves it. No inferential test is currently required by the SOP. If subgroup comparisons are later proposed, document hypotheses, sample adequacy, cell-size protection, and multiplicity before running them.
''')


write("05_Feasibility_Study_Materials.md", r'''
# Provisional feasibility-study materials

## Schedule feasibility

The title page identifies a September 2026–May 2027 study period. The repository shows three broad development checkpoints (2026-09-20, 2026-10-09, and 2026-10-11) and staged implementation guides. The following is a planning schedule, not a record that every listed task has been completed.

| Activity | Sep–Oct 2026 | Nov–Dec 2026 | Jan–Feb 2027 | Mar–Apr 2027 | May 2027 | Status/evidence |
|---|---|---|---|---|---|---|
| Requirements and consultation review | ███ | ░ |  |  |  | Ongoing; survey still open |
| Core backend/mobile/Admin development | ███ | █ | ░ |  |  | Major integrated functions present |
| Local data acquisition and validation | ░ | ███ | ███ | ░ |  | Blocked/awaiting authorized data and expert validation |
| Specific SDLC documentation | ░ | █ |  |  |  | Pending group decision/evidence |
| Functional/integration testing | ███ | ███ | █ | ░ |  | Mobile current; backend current rerun pending |
| Needs-survey closure and analysis | ░ | ███ |  |  |  | `[Insert cutoff date and valid N.]` |
| ISO/IEC 25010 evaluation |  | ░ | ███ | █ |  | Planned; evaluator/instrument details pending |
| Manuscript consolidation and defense preparation | ░ | █ | ███ | ███ | ███ | Working revision prepared |

Legend: █ planned/active concentration; ░ preparatory or contingency period. Confirm dates with the approved school calendar.

## Cost feasibility

No reliable actual-expense ledger was found. Older feasibility material refers to superseded Firebase/AI concepts and must not be used as the current budget. Enter documented receipts or approved estimates only.

| Line item | Basis | Quantity/unit | Actual or approved estimate (PHP) | Evidence required |
|---|---|---|---|---|
| Flutter/Dart, Python/Django, PostgreSQL/PostGIS | Open-source development stack | — | 0 licensing fee; exclude hosting/usage costs | Official license/project records |
| Map provider usage | Mapbox/OSM provider-aware configuration | `[Enter tier/usage]` | `[Enter verified amount.]` | Account billing/tier and attribution terms |
| Hosting/database/deployment | Not yet verified as production deployed | `[Enter plan.]` | `[Enter verified amount.]` | Quotation or invoice |
| Internet/data | Survey, development, testing | `[Enter months/GB.]` | `[Enter actual amount.]` | Receipts/plan statement |
| Transportation | BDRRMO, CPDCO, school, survey work | `[Enter trips.]` | `[Enter actual amount.]` | Trip log/receipts |
| Printing/reproduction | Letters, consent, instruments, manuscript | `[Enter pages/copies.]` | `[Enter actual amount.]` | Receipts |
| Device testing | Existing or borrowed Android devices | `[List devices.]` | `[Enter incremental cost.]` | Ownership/borrow record |
| Contingency | Only if adviser approves | `[Enter basis.]` | `[Enter approved amount.]` | Approved budget rule |
| Total | Sum only verified entries | — | `[Compute after entries are supplied.]` | Complete ledger |

## Operational feasibility

The workflow is understandable as a resident-initiated process: authenticate/complete onboarding, confirm a hypothetical rainfall scenario and location, request an assessment, review an explanation and map, and optionally use structured preparedness or center information. Authorized administrators maintain governed records through a separate portal. Operational adoption is not yet established because the needs survey and system evaluation are incomplete, public center and susceptibility data remain unapproved, and production support arrangements are absent.

Operational risks include users mistaking a scenario result for a forecast, outdated center records, inaccessible network service, unclear responsibility for publication approval, and incorrect extrapolation from provisional maps. Mitigations include prominent limitation/source notices, publication and review states, audit history, one-time foreground location, no raw-rule editing in the ordinary Admin portal, and an escalation statement directing users to official warnings and authorities.

## Technical feasibility

The integrated stack demonstrates technical feasibility for a prototype: Flutter/Dart mobile UI; Django/DRF/GeoDjango services; PostgreSQL/PostGIS records; deterministic inference; structured DSS; provider-aware Mapbox/OSM mapping; and custom Admin workflows. Current mobile verification passed 389 tests and static analysis. A current backend rerun, production-like environment test, device/browser matrix, load/security testing, backup/restore procedure, monitoring, and deployment evidence remain pending.

The largest technical constraint is not the ability to compute or display results; it is the absence of an approved operational Bacoor susceptibility dataset, validated rule parameters, and authorized current evacuation-center data. Software readiness must not be substituted for scientific or institutional validity.
''')


write("06_Provisional_Diagrams.md", r'''
# Provisional diagram register

The diagrams are based on the current integrated repository and are subject to adviser/TA approval, final actor confirmation, approved datasets, and final appendix ordering.

1. `Diagram_Conceptual_Framework_Provisional.png` — input–functional process–output model with four function-named modules; the database is not a conceptual module.
2. `Diagram_System_Architecture_Provisional.png` — Resident on the left, authorized administrator on the right, application/services in the center, storage below, and external reference information above as data sources rather than actors.
3. `Diagram_Use_Case_Provisional.png` — resident and authorized-administrator cases. Technical Django Admin is not presented as an ordinary user workflow.
4. `Diagram_DFD_Level_0_Provisional.png` — high-level data flows for resident input/results, administrative governed records, controlled storage, and approved external references.
5. `Diagram_HIPO_Provisional.png` — hierarchy of Resident Application, Assessment Services, Preparedness Services, and Administration/Governance.

Important limits shown or implied: no real-time monitoring/forecasting, autonomous alerts, background GPS, safe-road routing, live center occupancy, self-modifying rules, or official warning/evacuation authority. Nearest-center information is a software function whose public validity depends on authorized, current records.
''')


write("07_Appendix_Inventory.md", r'''
# Provisional appendix inventory

Letters must be assigned only after the department confirms final ordering.

| Provisional appendix | Contents | Current status | Required action/cross-reference |
|---|---|---|---|
| Appendix [A] | Approved active resident questionnaire and privacy/consent notice | Available; survey ongoing | Preserve exact version; cite in Participants/Instruments |
| Appendix [B] | Survey distribution record, cutoff notice, cleaning log, and final export metadata | Missing/ongoing | Add after closure; never publish direct identifiers |
| Appendix [C] | BDRRMO and CPDCO interview guides/permission letters | Partially available | Confirm interview guide versions and permissions |
| Appendix [D] | Consultation and interview transcripts | Available in multiple versions/locations | Identify canonical copies; redact personal data if required |
| Appendix [E] | Problem–objective–method–feature traceability matrix | Prepared | Cross-reference in Requirements Gathering |
| Appendix [F] | Conceptual Framework and System Architecture | Provisional diagrams prepared | Obtain approval before final insertion |
| Appendix [G] | Use Case, DFD, and HIPO diagrams | Provisional diagrams prepared | Validate actors/data flows; finalize captions |
| Appendix [H] | Gantt chart/schedule feasibility | Provisional planning table prepared | Confirm dates against school calendar |
| Appendix [I] | Budget/cost feasibility | Template prepared; actual costs missing | Add receipts/approved estimates and total |
| Appendix [J] | Functional/integration test plan, cases, and defect log | Code tests exist; manuscript appendix missing | Export representative cases/results; rerun backend |
| Appendix [K] | ISO/IEC 25010 evaluation instrument, evaluator instructions, consent, scoring plan | Planned | Validate instrument and evaluator groups before use |
| Appendix [L] | Data provenance, permissions, validation records, and approved local datasets | Missing/incomplete | Required before operational local claims |
| Appendix [M] | Interface screenshots/build information | Available in development outputs, status varies | Select only current integrated screens and caption by date/build |

Suggested in-text wording: “The current questionnaire is reproduced in Appendix [letter to confirm].” “The provisional system diagrams are provided in Appendix [letter to confirm] and will be finalized after actor and data-source validation.”
''')


write("08_Survey_Results_Insertion_Guide.md", r'''
# Survey-results insertion guide

| Manuscript location | Required final evidence | Survey item(s) | Planned display | Claim/decision supported |
|---|---|---|---|---|
| SOP opening | Final valid respondent count and eligible flow | Consent, Q1, Q2 | Flow table | Defines evidence base, not Bacoor-wide prevalence |
| SOP problem paragraph | Frequency of direct flood/heavy-rain effects | Q3 | Frequency table/bar | Establishes respondent experience context |
| SOP information-access paragraph | First source and ability to obtain information | Q4, Q5 | Two tables/bars | Describes current practice/access problem |
| SOP timing/clarity/specificity paragraph | Timing, understandability, location specificity | Q6–Q8 | Distribution table | Supports or weakens timing/clarity/location problem claims |
| SOP consistency/decision paragraph | Conflicting sources and decision effect | Q9, Q10 | Distribution table | Supports provenance/decision-support rationale |
| SOP problem summary | Problems selected | Q11 | Multiple-response ranked bar | Prioritizes verified problems |
| Requirements discussion | Useful lead time and information types | Q12, Q13 | Multiple-response tables | `[Confirm whether the survey evidence supports retaining this feature.]` |
| Research gap/needs statement | Extent current sources meet needs | Q14 | Distribution table | Supports remaining information gap |
| Operational feasibility | Mobile-phone access | Q15 | Frequency table | `[Insert final device-ownership result.]` |
| Delivery feasibility | Helpful delivery methods | Q16 | Multiple-response table | Supports delivery-channel discussion; does not measure internet connectivity |
| Participants | Valid N, excluded/invalid submissions, response rate, barangay distribution | Consent, Q1, Q2, D4 plus distribution log | Participant flow/profile | `[Insert final number of valid respondents.]` and `[Insert final distribution by barangay.]` |
| Limitations | Missing items, ambiguous locations, duplicates, coverage | All + metadata | Data-quality table | Bounds generalization |

## Decision rules

1. Do not retain a feature solely because it already exists in code. Triangulate relevant survey selections with interviews, literature, safety, and data availability.
2. A survey preference does not authorize an official forecast, warning, evacuation order, safe route, live occupancy, or unvalidated local susceptibility classification.
3. For Q13 and Q16, rank selections using valid-respondent denominators and report that they are multiple-response items.
4. If smartphone access or mobile-application preference is low, revise operational claims and recommend complementary channels; do not alter the responses.
5. Q16 does not directly measure internet quality. Use `[Insert final connectivity result.]` only if a separate documented connectivity measure exists; otherwise state that connectivity was not measured.
6. Before analysis, document the survey cutoff, export timestamp, duplicate rule, eligibility rule, missing-data treatment, and any exclusions.

## Mandatory placeholders to locate and replace

- `[Insert final number of valid respondents.]`
- `[Insert final percentage of respondents who reported this problem.]`
- `[Insert final distribution by barangay.]`
- `[Insert final device-ownership result.]`
- `[Insert final connectivity result.]`
- `[Confirm whether the survey evidence supports retaining this feature.]`
''')


write("09_Revision_Log.md", r'''
# Revision log — working copy dated 2026-10-11

| Section | Change | Evidence/reason | Status |
|---|---|---|---|
| Introduction | Removed incomplete-survey percentages as final evidence; strengthened general-to-local lead-in | Survey ongoing; official/peer-reviewed sources already cited | Applied |
| Statement of the Problem | Replaced solution-first opening with background, evidence gaps, placeholders, and five aligned provisional questions | Oct. 7 consultation source; active survey; interviews | Applied/provisional |
| Objectives | Aligned five objectives with five SOP questions; shortened testing/evaluation; listed major implemented functions | Current repository and traceability audit | Applied/provisional |
| Framework | Consolidated technology-named components into four function-named modules; removed database as conceptual module | Consultation source and current architecture | Applied; diagram provisional |
| Architecture | Clarified Resident/Admin actors and external reference data as sources; updated Mapbox/OSM and current service functions | Current repository | Applied; image provisional |
| Significance | Updated resident account, structured DSS, and nearest-center descriptions without claiming official data | Current repository/data limits | Applied |
| Time and Place | Removed methodology/procedure discussion and retained period/locations concisely | Consultation source | Applied |
| Scope and Limitations | Separated scope from limitations and corrected current functions | Current code/routes/interfaces/tests | Applied |
| Definitions | Updated Android/Admin/DSS/nearest-center/account/map terms | Current repository | Applied selectively |
| Chapter 2 | Preserved substantive review and comparison tables; added no unverified study claims in this phase | Existing sources and verification matrix | Preserved; further full-text local-study expansion pending |
| Chapter 3 research design | Removed unsupported Design Science Research claim; used a provisional developmental-descriptive classification | Actual surveys, interviews, artifact development | Applied/provisional |
| Chapter 3 SDLC | Declined to label Scrum/Kanban/RAD/Waterfall without records; listed missing evidence | Git history, day guides, backlog plan | Applied/pending decision |
| Chapter 3 feasibility | Added schedule, cost, operational, and technical feasibility summaries with placeholders | Current code and project records | Applied/provisional |
| Participants/instruments/analysis | Described ongoing survey and exact active instrument; added placeholders and categorical analysis plan | Approved questionnaire | Applied/provisional |
| Testing/evaluation | Distinguished current mobile verification, historical backend evidence, and planned ISO evaluation | Test run/repository docs/ISO standard | Applied |
| Appendices | Added provisional diagram/appendix register and cross-reference language | User request | Applied/provisional |
| References | Removed Peffers et al. (2007) because DSR is no longer claimed; otherwise preserved | Directly affected citation only | Applied |

No source code, dataset, consultation file, survey instrument, approved school document, original manuscript, or separate baseline was modified.
''')


write("10_Missing_Information_and_Pending_Decisions.md", r'''
# Missing information and pending decisions

## Survey and research records

- Final survey cutoff date/time, total submissions, valid N, response rate, exclusions, duplicate rule, and final barangay distribution.
- Whether Q2 “I am not sure” is excluded or manually verified.
- Survey distribution log/denominator needed to calculate a response rate.
- Separate connectivity measure; Q16 is preference, not connectivity quality.
- Approved sampling description and target population frame.
- Canonical interview/consultation file list and any required redactions/consent records.

## Research design and SDLC

- Adviser/TA approval of the provisional developmental-descriptive research classification.
- A specific SDLC decision supported by actual evidence. Scrum would require documented roles, backlog ownership, sprint planning/review/retrospective and timeboxes; Kanban would require a maintained board, explicit workflow/WIP practice, and flow records; RAD would require documented prototyping/workshop/time-boxed construction evidence. These records were not established.
- Final methodology diagram after the SDLC decision.

## Data, rules, and governance

- Approved operational Bacoor City susceptibility dataset.
- Provenance, permission, derivation method, validation, resolution, and update process for every public map/data layer.
- Expert/institutional validation of rainfall definitions, duration combinations, rule conditions, priorities, conflict handling, class labels, explanation text, and color convention.
- Resolution of BDRRMO’s suggested yellow/violet/orange/red convention versus current green/yellow/orange/red.
- Authorized current evacuation-center records, verification date, contact/operating-status governance, and update owner.
- Approved preparedness/DSS content and source/version review workflow.
- Fixed-schema Expert System parameter CSV specification and implementation; current center import is not this requirement.

## Testing, evaluation, and deployment

- Current backend dependency setup and full test rerun.
- Production-like integration, device/browser coverage, accessibility, security, performance, backup/restore, and deployment tests.
- Final functional-test appendix and defect log.
- Approved ISO/IEC 25010 evaluator groups, instrument, scale, scoring/interpretation, sample, and completed results.
- Production hosting, support owner, monitoring, update, incident, and retention arrangements.

## Manuscript and appendices

- Final SOP/objective wording after survey integration.
- Final appendix letters and school-approved diagram notation/captions.
- Full-text verification and additional recent Philippine peer-reviewed studies if Chapter 2 depth requires them.
- Final references cross-check after all survey/evaluation sections are inserted.
- Final diagram approval; provisional images must not be treated as approved architecture or theory until reviewed.
''')


# Build the revised manuscript by copying the current manuscript and replacing only
# evidence-sensitive portions. The source file is never overwritten.
shutil.copy2(SOURCE, DOCX_OUT)
doc = Document(DOCX_OUT)


HEADING_INDICES = {320, 323, 326, 328, 330, 332, 334, 336, 338, 340, 342, 345, 347}
LABEL_INDICES = {88, 89, 90, 91, 110, 111, 112, 113, 119, 120, 121, 122, 123, 124, 125, 126, 127, 128, 131, 135, 136, 149}


def set_p(idx: int, text: str) -> None:
    p = doc.paragraphs[idx]
    for r in list(p.runs):
        p._p.remove(r._r)
    if text:
        if idx in LABEL_INDICES and ". " in text:
            label, rest = text.split(". ", 1)
            p.add_run(label + ".").bold = True
            p.add_run(" " + rest)
        else:
            r = p.add_run(text)
            if idx in HEADING_INDICES:
                r.bold = True


replacements = {
67: "Flooding remains one of the most persistent hazards affecting Philippine communities. Its consequences extend beyond inundation and property damage to interrupted travel, disrupted livelihoods, displacement, and difficult household decisions before and during severe weather. National disaster statistics continue to document the effects of extreme events on Filipino families and communities (Philippine Statistics Authority [PSA], 2026). These conditions make accurate, understandable, and appropriately qualified risk information an essential part of disaster preparedness.",
68: "Bacoor City, Cavite, includes low-lying and river-connected areas exposed to recurring flood conditions. Flood-mitigation works in the Imus River Basin were established to reduce peak floodwater volume in Cavite, including low portions of Bacoor and its tributaries (Japan International Cooperation Agency [JICA], 2021). Local interviews also identified historical flood records, geographic materials, preparedness procedures, and evacuation information held by city offices, together with restrictions on the availability and public use of machine-readable spatial data.",
69: "Residents may obtain flood information through official social-media pages, barangay announcements, text messages, broadcast media, websites, community networks, and personal observation. However, the ongoing resident survey must still establish how eligible respondents assess the accessibility, timing, clarity, location specificity, consistency, and usefulness of these sources. [Insert final number of valid respondents.] [Insert final percentage of respondents who reported this problem.] Preliminary survey reports retained in the project are treated only as earlier needs-assessment records and are not pooled with the active survey or presented as final findings.",
70: "Research on flood-risk communication emphasizes that technical information is most useful when users can understand its geographic meaning, source, limitation, and connection to appropriate preparedness action (Mostafiz et al., 2022). Geographic provenance is also necessary because a technically processed boundary or map does not automatically become institutionally verified hazard information (Tullis & Kar, 2021). In Bacoor City, the absence of an approved operational susceptibility dataset and the need for local expert validation remain central constraints.",
71: "Against this context, the study examines a scenario-based approach to pre-event flood-susceptibility assessment and preparedness support. The approach keeps a deterministic assessment separate from decision-support guidance, requires the user to confirm a hypothetical rainfall scenario and location, and discloses that results are not live observations, official forecasts, warnings, or evacuation orders.",
72: "FloodSense is developed as an Android resident application and a separate Web Administration System that organize these verified concepts. The current prototype demonstrates rule-based assessment, explanation, interactive geographic presentation, structured preparedness guidance, resident account controls, and governed administrative workflows. Its public operational validity remains dependent on approved local data, documented provenance, expert and institutional validation, completed evaluation, and production arrangements.",
74: "Flood preparedness in Bacoor City depends not only on the availability of information but also on whether residents can obtain it in time, understand it, relate it to a location, and use it in an appropriate decision. City offices hold relevant historical, geographic, preparedness, and evacuation information, yet access restrictions, data provenance, validation, and public communication responsibilities limit what can be translated directly into a public assessment. The active resident survey is still collecting evidence about current sources, timing, clarity, location specificity, consistency, information needs, and technology access. [Insert final number of valid respondents.] [Insert final distribution by barangay.]",
75: "The local research problem therefore has two connected dimensions. The first is an evidence problem: the study must establish, without treating incomplete responses as final, which information difficulties and preparedness needs eligible respondents actually report. The second is a governance and design problem: any scenario-based susceptibility assessment must use traceable and locally appropriate knowledge, distinguish administrative reference geography from flood-susceptibility data, explain its basis and limitations, and remain separate from official forecasting and warning authority.",
76: "The main problem of the study is how verified resident needs, local institutional requirements, controlled geographic and rainfall information, deterministic rule-based inference, and source-backed preparedness guidance can be organized into an understandable pre-event support system for Bacoor City without overstating the authority or completeness of its data.",
77: "Specifically, the study seeks to answer the following provisional questions: (1) What flood-information practices, access conditions, and problems do eligible Bacoor City respondents report? (2) What verified knowledge, rainfall parameters, geographic records, preparedness content, and governance controls are required for a locally appropriate scenario-based assessment? (3) How can the verified requirements be implemented through an Android resident application and a separate Web Administration System? (4) To what extent do the implemented functions operate correctly and consistently under defined technical test cases? (5) How do the defined evaluator groups assess the relevant FloodSense product-quality characteristics using an approved ISO/IEC 25010-based instrument? Final wording will be reviewed after the ongoing survey is closed and analyzed.",
79: "Generally, the study aimed to develop and evaluate FloodSense as a scenario-based Android system for explainable flood-susceptibility assessment and separate pre-event preparedness support for Bacoor City, subject to validated local knowledge, data provenance, and institutional limitations.",
81: "identify the flood-information practices, access conditions, problems, and preparedness needs reported by eligible Bacoor City respondents and relevant local offices through the active resident survey, interviews, and document review;",
82: "determine the locally appropriate facts, rainfall parameters, geographic records, rule conditions, preparedness content, evacuation-center records, sources, validation states, and governance controls required for the knowledge base and supporting services;",
83: "develop the Flutter-based Android resident application and Django-based Web Administration System with: (a) resident authentication, legal acceptance, onboarding, preferences, and account controls; (b) user-confirmed rainfall and location scenarios; (c) deterministic rule matching, conflict handling, susceptibility output, and explanation; (d) provider-aware interactive mapping, manual selection, map pin, one-time foreground location, and barangay confirmation; (e) structured Decision Support System guidance kept separate from the assessment; (f) nearest-center information subject to authorized current records; and (g) governed administrative workflows that do not expose raw rule editing to ordinary administrators;",
84: "test the implemented mobile, backend, geographic, Expert System, Decision Support System, account, map, permission, and administrative functions through defined unit, integration, API, spatial, and end-to-end test cases; and",
85: "evaluate the relevant product-quality characteristics of FloodSense using an approved ISO/IEC 25010-based instrument and defined evaluator groups. [Insert final evaluation participants, scores, and interpretation after data collection.]",
87: "The theoretical and conceptual framework of FloodSense relates user-confirmed inputs, controlled knowledge, deterministic inference, geographic presentation, and separate preparedness support. The framework uses four function-named modules so that technologies and database storage remain implementation mechanisms rather than conceptual modules.",
88: "Scenario and Location Input Module. The resident selects or confirms a hypothetical rainfall intensity and duration, chooses a supported location manually or through an optional one-time foreground location request, and confirms or corrects the proposed barangay before explicitly requesting an assessment. These inputs are not live measurements or forecasts.",
89: "Susceptibility Assessment and Explanation Module. The deterministic Rule-Based Expert System validates the confirmed inputs, matches applicable rules, resolves priorities or conflicts, and returns a Low, Moderate, High, or Very High class with a human-readable explanation. Repeatability does not establish scientific validity; the rules and local data still require provenance, expert review, and institutional authorization.",
90: "Preparedness and Evacuation Information Module. A separate structured Decision Support System uses the assessment only as context for published guidance and does not change its class or issue an evacuation order. The mobile application can request nearest-center information, but distance does not establish route safety, opening status, accessibility, or live capacity, and public use depends on authorized current records.",
91: "Governance and Data Stewardship Module. Authorized administrators use the custom Web Administration System for dashboards and controlled workflows involving map and rainfall references, parameters, DSS content and structured flows, evacuation centers and staged center imports, data sources, and audit history. Ordinary administrators do not edit raw rules, rule conditions, priorities, conflict-resolution logic, or algorithm code. Technical Django Admin remains separate. The confirmed fixed-schema Expert System parameter CSV exchange is not yet implemented.",
92: "The four modules produce an explainable, scenario-specific result, geographic presentation, and preparedness information while maintaining source, permission, validation, and limitation boundaries. The 47-barangay layer is an administrative reference rather than a susceptibility dataset, and the framework does not include real-time monitoring, autonomous alerts, official forecasts or warnings, background GPS, automatic evacuation orders, safe-route calculation, machine-learning prediction, or self-modifying rules.",
93: "[AUTHOR NOTE: Insert the approved Conceptual Framework diagram here after the provisional module names and relationships are reviewed.]",
94: "",
95: "",
96: "",
97: "",
98: "",
99: "",
100: "The provisional System Architecture places the resident on the left, the authorized administrator on the right, and the resident application, custom Web Administration System, Django REST services, deterministic inference, geographic services, structured DSS, and evacuation-center service between them. PostgreSQL/PostGIS provides controlled relational and spatial storage below these services. Approved rainfall, geographic, preparedness, and center information enters as external reference data rather than as users or actors. Provider-aware Mapbox/OSM interfaces present maps when configured. [AUTHOR NOTE: Insert the approved System Architecture diagram here after actor and data-source validation.]",
109: "FloodSense investigates how an explainable, scenario-based assessment and separate preparedness guidance may support pre-event understanding in Bacoor City. Any benefit described below remains subject to approved data, completed evaluation, and the system's role as a complementary information tool.",
110: "For the Bacoor Disaster Risk Reduction and Management Office (BDRRMO). The governed prototype may provide an additional channel for publishing validated preparedness and evacuation-center information while preserving source and audit records. It does not replace official warnings, operations, or emergency directives.",
111: "For the City Planning and Development Coordinator Office (CPDCO). The study demonstrates how appropriately authorized planning and geographic information may be referenced in a location-specific application. City-held data cannot be assumed public or operational without permission, provenance, and validation.",
112: "For Bacoor City residents. The resident application provides a user-initiated flow for confirming a hypothetical rainfall scenario and location, requesting an explained susceptibility result, viewing a map, completing structured preparedness guidance, and viewing nearest-center information when verified records are available. It does not provide live monitoring, official forecasts or warnings, safe-route advice, live occupancy, or safety guarantees. [Confirm whether the survey evidence supports retaining this feature.]",
113: "For future researchers. The study documents the separation of deterministic assessment from preparedness guidance, the role of geographic provenance and validation, and the difference between implemented software and authorized operational data. Its methods and findings should not be transferred to another locality without renewed validation.",
115: "The study is scheduled from September 2026 to May 2027. Development, testing, consultation review, and manuscript preparation are conducted primarily at Cavite State University–Bacoor City Campus, Soldiers Hills IV, Molino, Bacoor City, Cavite.",
116: "Resident survey activities are conducted in included Bacoor City communities. Recorded interviews and document requests involve the BDRRMO at the City of Bacoor Legislative and Disaster Resilience Building and the CPDCO at the Bacoor Government Center, Barangay Bayanan. [Insert final survey distribution locations and dates.]",
117: "",
119: "Scope. FloodSense is a scenario-based, pre-event flood-susceptibility assessment and preparedness-support system for Bacoor City. A resident authenticates, completes required legal/onboarding steps, selects or confirms a hypothetical rainfall intensity and duration, selects or confirms a supported location, and explicitly requests an assessment. The deterministic Rule-Based Expert System returns a Low, Moderate, High, or Very High demonstration class and explanation, while a separate structured DSS presents published preparedness guidance without changing the class.",
120: "Scope. The resident application includes manual area selection, map-pin adjustment, optional one-time foreground location, barangay confirmation or correction, provider-aware interactive maps, account and preference controls, and nearest-center lookup. The custom Web Administration System includes authenticated dashboards and governed workflows for map and rainfall references, parameter review, DSS guidance and structured flows, evacuation-center records and staged center import, data sources, and audit history.",
121: "Scope. Ordinary administrators do not edit raw Expert System rules, rule conditions, priorities, conflict-resolution logic, or algorithm code. Technical Django Admin is separate and reserved for developer or research maintenance. The confirmed fixed-schema CSV export/import requirement for permitted Expert System parameter values remains unimplemented and is distinct from the available evacuation-center import workflow.",
122: "Limitations. FloodSense is not a real-time monitoring, forecasting, or warning system. It does not perform continuous rainfall or water-level monitoring, automatic background polling, background GPS tracking, autonomous alerts, official warnings, automatic evacuation orders, safe-road routing, live center occupancy, machine-learning prediction, self-modifying rules, or fully offline operation.",
123: "Limitations. The 47-barangay display is an administrative reference layer derived and technically checked for software use; it is not automatically City-issued, City-verified, or a flood-susceptibility dataset. No approved operational Bacoor City susceptibility dataset is currently established in the repository. Provisional MGB or other materials, processing scripts, mockups, and screenshots do not establish operational integration or institutional approval.",
124: "Limitations. Nearest-center discovery is implemented as a software flow, but its public usefulness depends on authorized, current center records. A displayed distance does not verify road passability, accessibility, current opening, available capacity, or safety. Users must confirm emergency information through responsible authorities.",
125: "Limitations. The active survey is ongoing. [Insert final number of valid respondents.] [Insert final distribution by barangay.] [Insert final device-ownership result.] [Insert final connectivity result.] No final needs findings, response rate, user-satisfaction result, ISO/IEC 25010 score, or Bacoor-wide generalization is made in this working version.",
126: "Limitations. Current mobile verification passed 389 automated tests and static analysis, while the most recent full backend result available in project documentation is historical rather than a current rerun. Automated tests establish software behavior under their cases; they do not establish scientific validity, institutional approval, usability, or production readiness.",
127: "Limitations. Integrated functions currently require networked backend services. Assessment history, live notifications, approved operational datasets, fixed-schema Expert System parameter CSV exchange, final evaluation, and verified production deployment are not established.",
128: "Limitations. Personal data and one-time foreground location must be handled under Republic Act No. 10173 and the approved consent/privacy process. Direct identifiers from the active survey will not be included in the analysis dataset or public results, and location use does not authorize continuous tracking.",
131: "Android Application. The resident-facing FloodSense software developed with Flutter. It includes authentication, legal acceptance, onboarding, scenario and location confirmation, assessment requests, map/results presentation, structured preparedness guidance, nearest-center information, preferences, and account controls. It does not provide live alerts, assessment history, fully offline operation, or official warnings.",
135: "Decision Support System (DSS). The component separate from the Rule-Based Expert System that presents published preparedness guidance and structured questions using the assessment context. It does not change the susceptibility result, issue an official warning or evacuation order, or make the user's final decision.",
136: "Dynamic Map. The provider-aware interactive map that supports panning, zooming, scenario-dependent demonstration coloring, manual selection, map-pin placement, foreground location, reference boundaries, and center markers when data are available. Its display does not convert provisional data into official hazard information.",
149: "Web Administration System. The custom browser-based interface used by authorized administrators for dashboards and governed workflows involving map and rainfall references, parameters, DSS content and structured flows, evacuation centers and staged center imports, data sources, and audit history. It does not expose raw rule editing. Technical Django Admin is separate. Fixed-schema Expert System parameter CSV export/import is confirmed but not implemented.",
316: "This chapter describes the research design, materials, requirements gathering, feasibility, participants, instruments, analysis plan, development evidence, testing, and planned evaluation for FloodSense. Completed activities, ongoing data collection, and planned work are distinguished throughout.",
318: "Research materials include the active approved resident questionnaire and privacy/consent notice, preliminary needs-assessment reports kept as separate historical records, BDRRMO and CPDCO interviews and letters, consultation records, official and scholarly sources, requirements documents, diagrams, interface evidence, and the 47-barangay administrative reference layer. Direct identifiers are excluded from public analysis.",
319: "Development and verification materials include Windows computers and available Android devices; Git version control; Flutter and Dart; Python, Django, Django REST Framework, and GeoDjango; PostgreSQL/PostGIS; Mapbox and OpenStreetMap-compatible mapping components; browser-based custom Admin templates; database migrations; API documentation; automated tests; and defect/verification records. Survey distribution and analysis use the exact approved instrument and an exported tabular dataset after the survey closes. Tools not supported by project records are not listed as methods.",
320: "Research Design",
321: "The study provisionally follows a developmental-descriptive research design. The developmental component covers the requirements-based creation and technical verification of the mobile, backend, geographic, Expert System, DSS, and administrative artifact. The descriptive quantitative component covers the resident questionnaire through frequencies and percentages, while qualitative key-informant interviews and document review explain institutional practice, data access, provenance, and validation constraints. The final designation will be confirmed with the adviser before submission.",
322: "",
323: "System Development Life Cycle",
324: "The software has been developed through successive requirements, design, implementation, testing, feedback, and correction increments. The available project records do not establish the roles and ceremonies required for Scrum, the maintained workflow and work-in-progress controls required for Kanban, or a documented RAD or Waterfall process. [AUTHOR NOTE: Select and describe a specific SDLC only after the group supplies the corresponding workflow evidence and obtains approval. Insert the final SDLC diagram here.]",
325: "Until that decision is made, Chapter 3 reports the actual activities rather than assigning an unsupported method name: requirements were elicited from surveys, interviews, documents, and consultations; functions were designed and implemented in backend, mobile, geographic, and Admin components; automated and manual checks informed corrections; and data-dependent or evaluation-dependent work remains pending.",
326: "Schedule Feasibility",
327: "The documented study period is September 2026 to May 2027. Major software functions are integrated, but authorized local-data acquisition, rule/content validation, survey closure and analysis, current backend verification, formal ISO/IEC 25010 evaluation, and manuscript completion remain on the critical path. A provisional Gantt chart is provided in Appendix [letter to confirm]. Its dates must be reconciled with the approved school calendar.",
328: "Cost Feasibility",
329: "The core development stack is open source, but map-provider usage, hosting, connectivity, transportation, printing, device testing, and deployment may create costs. No reliable complete expense ledger was found. The line-item budget in Appendix [letter to confirm] therefore retains placeholders for receipts or approved estimates and excludes superseded Firebase or AI-assistant assumptions.",
330: "Operational Feasibility",
331: "The resident workflow is user initiated and separates assessment from guidance, while the custom Admin portal provides governed publication and record workflows. Operational adoption remains unconfirmed because the active needs survey and final evaluation are incomplete, approved susceptibility and center datasets are unavailable, and production ownership and support arrangements have not been established.",
332: "Technical Feasibility",
333: "The current repository establishes prototype-level technical feasibility for authentication, scenario/location input, deterministic assessment and explanation, geographic display, structured DSS guidance, nearest-center lookup, and governed Admin workflows. Current mobile verification passed 389 tests and static analysis. A current backend rerun, production-like deployment, security/performance/accessibility testing, and authorized operational data remain pending.",
334: "Requirements Gathering",
335: "Resident requirements are gathered through the active questionnaire on flood experience, information practices, timing, clarity, location specificity, consistency, decision effect, preparedness information, mobile access, and delivery preference. BDRRMO and CPDCO interviews inform institutional practices, source ownership, geographic/data access, evacuation information, and validation limitations. Literature and official documents inform concepts and safety boundaries. A proposed function becomes a thesis requirement only after it is traceable to evidence, feasible and safe, consistent with the approved scope, and validated through the appropriate technical or stakeholder method. The traceability matrix is provided in Appendix [letter to confirm].",
336: "Participants and Sampling",
337: "The active survey targets consenting adults aged 18 or older who live in Bacoor City communities included in the study. The final sampling method, distribution denominator, response rate, and handling of Q2 ‘I am not sure’ require documentation. [Insert final number of valid respondents.] [Insert final distribution by barangay.] [Insert final excluded or invalid submissions and response rate.] Key informants include available BDRRMO and CPDCO representatives whose statements are used for requirements and limitations, not as proof that the application is effective.",
338: "Research Instruments",
339: "The active instrument is the bilingual FloodSense Resident Survey Questionnaire with a Data Privacy Act notice and consent, demographics, eligibility/flood experience, current flood-information experience, and technology access/preferences. It is preserved unchanged while responses are being collected. Interview guides, permission letters, consultation records, functional test cases, defect records, and the planned ISO/IEC 25010-based evaluation instrument support other evidence stages. Instrument validation records and the final evaluation form remain to be supplied.",
340: "Data Gathering and Analysis",
341: "After closure, the survey export will be archived with its cutoff and cleaned under documented consent, eligibility, duplicate, range, missing-data, and multiple-response rules. Categorical items will be summarized by counts and valid percentages; ordered items will be shown through distributions and medians where useful. Weighted means will not be used unless the actual scale and adviser-approved interpretation justify them. Barangay text will be normalized through a documented codebook. Interview and document evidence will be coded thematically for information practices, data ownership, provenance, validation, governance, and scope. The item-level plan is provided in Appendix [letter to confirm].",
342: "Testing and Evaluation",
343: "Functional verification covers rule matching and conflict handling, validation, API responses, spatial queries, authentication and permissions, account flows, scenario/location confirmation, map behavior, structured DSS flow, nearest-center behavior, center import, and Admin workflows. Integration testing checks mobile-to-API and service-to-database behavior. Test records must identify case, precondition, input, expected result, actual result, status, environment, date, and defect link. Current mobile tests pass; the backend result available from project documents is historical until rerun.",
344: "Formal product-quality evaluation is planned using only the relevant ISO/IEC 25010:2023 characteristics and an approved instrument. The evaluator groups, sample, tasks, response scale, scoring, interpretation bands, consent, and analysis must be defined before administration. [Insert final evaluator count, scores, comments, and interpretation.] Automated test success and user ratings will be reported separately because neither substitutes for the other.",
345: "Ethical and Data-Governance Procedures",
346: "Participation is voluntary and governed by the active consent/privacy notice. Names and other direct identifiers will be removed from the analysis copy, small geographic cells will be reviewed for disclosure risk, and access will be limited to authorized researchers and academic advisers. Geographic, susceptibility, preparedness, and evacuation information will be labeled by source, validation, and publication status; technical processing alone will not be described as institutional approval.",
347: "Research Limitations",
348: "The active resident survey, final system evaluation, and several validation records are incomplete. The absence of an approved operational susceptibility dataset prevents official Bacoor classification claims. Nearest-center and DSS software functions do not establish that every underlying record is authorized or current. Assessment history, live alerts, fully offline use, fixed-schema Expert System parameter CSV exchange, official forecasting/warnings, and production deployment are not established.",
349: "[AUTHOR NOTE: Complete the survey-dependent findings and final evaluation sections only after the validated exports, cleaning log, approved analysis, and interpretation are available.]",
350: "The provisional appendices contain the active questionnaire reference, traceability materials, feasibility tables, diagrams, and testing/evaluation templates. Appendix letters will be finalized after the complete order is approved.",
}

for idx, value in replacements.items():
    set_p(idx, value)

# Remove the superseded landscape architecture-diagram section, including its
# temporary landscape/portrait section breaks. The replacement working diagram
# is appended in the provisional-appendix pages and referenced in text.
for remove_idx in range(107, 100, -1):
    p_remove = doc.paragraphs[remove_idx]
    p_remove._element.getparent().remove(p_remove._element)

# Remove the now-uncited Peffers DSR reference while preserving all unrelated entries.
for p in list(doc.paragraphs):
    if p.text.startswith("Peffers, K."):
        p._element.getparent().remove(p._element)

# Add provisional appendix materials after the references.
doc.add_page_break()
p = doc.add_paragraph("PROVISIONAL APPENDICES (ORDER AND LETTERS TO BE CONFIRMED)")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
if p.runs:
    p.runs[0].bold = True

for title, img in [
    ("Appendix [Letter to Confirm]. Provisional Conceptual Framework", "Diagram_Conceptual_Framework_Provisional.png"),
    ("Appendix [Letter to Confirm]. Provisional System Architecture", "Diagram_System_Architecture_Provisional.png"),
    ("Appendix [Letter to Confirm]. Provisional Use Case Diagram", "Diagram_Use_Case_Provisional.png"),
    ("Appendix [Letter to Confirm]. Provisional Level-0 Data Flow Diagram", "Diagram_DFD_Level_0_Provisional.png"),
    ("Appendix [Letter to Confirm]. Provisional HIPO Diagram", "Diagram_HIPO_Provisional.png"),
]:
    doc.add_page_break()
    cap = doc.add_paragraph(title)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap.runs:
        cap.runs[0].bold = True
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(OUT / img), width=Inches(6.5))
    note = doc.add_paragraph("Provisional: validate actors, data sources, relationships, notation, and appendix numbering before final submission.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()
doc.add_paragraph("Appendix [Letter to Confirm]. Schedule Feasibility and Gantt Chart")
doc.add_paragraph("See the phase schedule in 05_Feasibility_Study_Materials.md. Replace provisional periods with the approved academic calendar and actual completion records.")
doc.add_page_break()
doc.add_paragraph("Appendix [Letter to Confirm]. Cost Feasibility and Line-Item Budget")
doc.add_paragraph("See the cost template in 05_Feasibility_Study_Materials.md. Enter only documented receipts or approved estimates before calculating the total.")

doc.core_properties.title = "FloodSense Chapters 1–3 Survey-Ongoing Working Revision"
doc.core_properties.subject = "Evidence-aligned working revision; final survey results pending"
doc.core_properties.comments = "Original manuscript preserved; working copy generated 2026-10-11."
doc.save(DOCX_OUT)

print(DOCX_OUT)
